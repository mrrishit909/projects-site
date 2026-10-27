"""Build artifacts/changeworks-analyst-workbook.xlsx with openpyxl (P09-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: the published snapshot ``web/public/data`` (typed into the data sheets through the independent Python side: the 700 benefit rows, the 900 cost rows, the 30 scenarios and, for each
of their 210 scenario x initiative pairs, the assumption versions the scenario uses: the adoption ramp and the capacity conversion rate), the Python domain (``Domain reference``: the
figures it publishes, pasted for the check columns only), tests/fixtures/golden/expected.json and artifacts/analysis-results.json (data/sql/analysis/run_all.py). openpyxl cannot calculate,
so the file it writes holds formulas only. With --recalc, LibreOffice (soffice) opens it headless and saves a recalculated copy over the same path, which then also carries cached values.

What the formulas do, live, for the scenario chosen at ``Controls!B3`` (any of the 30; scn-025 as delivered): the effective period rate (the declared rate when its basis is the period basis,
else the recorded conversion; an optional override at ``Controls!B4``), the horizon and the members; each benefit row's ramp factor (indexed from its own starts_period, last value held), its
class (cash, revenue, avoided cost realised only with cash_realisation and an evidence reference, capacity converted only with a conversion rate, otherwise beside the flow or hours), the
expansion of every cost over its half-open range, the dated cash flow of each initiative (costs, the four classes, net, cumulative, discount factor 1/(1+r)^t, discounted, cumulative
discounted), NPV, simple and discounted payback by the last crossing with their reasons, the double-count rule (claims sharing a source key in one period across initiatives count once at
the largest; ties by entering the flow, then initiative and benefit id), its present value and the attributable portfolio NPV, and the hours residual. What it does not do: choose which
assumption version a scenario uses (pinned, else the latest created by the scenario's created_at); that choice is TYPED from the independent Python side (limit W-02). Excel is never run.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from fractions import Fraction as F
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "domain"), str(ROOT), str(ROOT / "data" / "sql" / "analysis")]

import python_side as ps  # noqa: E402
from changeworks import run_pipeline  # noqa: E402
from prepare_snapshot import snapshot_inputs  # noqa: E402

OUT = ROOT / "artifacts" / "changeworks-analyst-workbook.xlsx"
RESULTS_PATH = ROOT / "artifacts" / "analysis-results.json"
EXPECTED = json.loads((ROOT / "tests/fixtures/golden/expected.json").read_text())
CODE_REVISION = "29948f9"
DATASET_HASH = "6bf7ebaad6bf75084c82c41de1031d6d585af730c8a22979ed7c77df2a38e3a9"
DEFAULT_SCENARIO = "scn-025"
HDR = PatternFill("solid", fgColor="1F3A33")
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
EXCEL_EPOCH = datetime(1899, 12, 30, tzinfo=timezone.utc)
SHEET_ORDER = ["Read me", "Controls", "Checks", "Portfolio", "Initiatives", "Cash flow", "Golden case", "Assumptions", "Benefits", "Costs", "Scenario inputs", "Scenarios", "Periods",
               "Domain reference", "SQL results (pasted)"]
MAX_RAMP = 8
NB, NC, NSI, NS, NI, NP = 700, 900, 210, 30, 24, 36          # rows of the data sheets (asserted when built)
BL, CL, SIL, SL = 4 + NB - 1, 4 + NC - 1, 4 + NSI - 1, 4 + NS - 1
CFL = 4 + NI * NP - 1
IL = 4 + NI - 1


def serial(s: str | None):
    return None if s is None else (ps.ts(s) - EXCEL_EPOCH).total_seconds() / 86400


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def put(ws, rows: list[list], first_row: int):
    for r, row in enumerate(rows, start=first_row):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if v is not None and not (isinstance(v, str) and v.startswith("=")):
                cell.font = BLUE


def num(x):
    """A decimal string or Fraction as a float for a typed cell (inputs carry at most 6 decimal places, so the float is exact to the cent)."""
    return None if x is None else float(F(x) if not isinstance(x, str) else ps.dec(x))


def idnum(s: str) -> int:
    return int(re.sub(r"\D", "", s))


class Data:
    def __init__(self) -> None:
        self.c = ps.load_snapshot()
        inputs, manifest = snapshot_inputs()
        self.manifest = manifest
        self.dom = run_pipeline(inputs)
        self.M = self.dom.metrics


# ---- data sheets -------------------------------------------------------------------------------------------------------------------------------------------------------------
def build_scenarios(wb, d: Data):
    ws = wb.create_sheet("Scenarios")
    title(ws, "The 30 scenario versions (immutable; an edit is a new scenario_id)", "conversion_rate is the declared 6-place monthly conversion of an annual rate (an assumption version, D-P9-06); effective_rate is a formula.")
    header(ws, 3, ["scenario_id", "baseline_version", "rate_basis", "declared_rate", "conversion_rate", "horizon_periods", "created_at (serial)", "initiatives", "effective_rate", "rate_source"])
    rows = []
    for sid in sorted(d.c.scn):
        s = d.c.scn[sid]
        conv = d.c.assump[s["rate_conversion_assumption_id"]]["value"] if s.get("rate_conversion_assumption_id") else None
        rows.append([sid, s["baseline_version"], s["rate_basis"], num(s["discount_rate"]), num(conv), s["horizon_periods"], serial(s["created_at"]), len(set(s["initiative_ids"]))])
    assert len(rows) == NS
    put(ws, rows, 4)
    for i in range(NS):
        r = 4 + i
        ws.cell(row=r, column=7).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=r, column=9, value=f"=IF(C{r}=Controls!$B$5,D{r},E{r})").number_format = "0.000000"
        ws.cell(row=r, column=10, value=f'=IF(C{r}=Controls!$B$5,"declared","conversion")')
    widths(ws, {"A": 11, "G": 18})


def build_scenario_inputs(wb, d: Data):
    ws = wb.create_sheet("Scenario inputs")
    title(ws, "The assumption versions each scenario uses for each of its initiatives (TYPED: the pinned version, else the latest created by the scenario's created_at; W-02)",
          "ramp_1..ramp_8 = the adoption ramp (fractions per period from each benefit row's starts_period, the last value held); capacity_rate = USD per hour, only when a staffing_change_evidence_ref exists.")
    header(ws, 3, ["key", "scenario_id", "initiative_id", "ramp_assumption_id", "ramp_len"] + [f"ramp_{k}" for k in range(1, MAX_RAMP + 1)] + ["capacity_rate"])
    rows = []
    for sid in sorted(d.c.scn):
        for iid in sorted(set(d.c.scn[sid]["initiative_ids"])):
            m = d.c.model(sid, iid)
            ramp = m["ramp"] or []
            assert len(ramp) <= MAX_RAMP
            rows.append([f"{sid}|{iid}", sid, iid, m["ramp_id"], len(ramp)] + [num(v) for v in ramp] + [None] * (MAX_RAMP - len(ramp)) + [num(m["rate_cap"])])
    assert len(rows) == NSI
    put(ws, rows, 4)
    widths(ws, {"A": 20, "D": 18})


def build_benefits(wb, d: Data):
    ws = wb.create_sheet("Benefits")
    title(ws, "700 benefit rows (typed) and their treatment in the selected scenario (formulas)",
          "row = the line of Scenario inputs for (Controls!B3, initiative), 0 when the initiative is not in the scenario; class follows D-01; money is USD (hours x rate for converted capacity); the double-count columns rank the claims of one source key in one period.")
    cols = ["benefit_id", "initiative_id", "init_no", "benefit_no", "benefit_type", "period_index", "starts_period", "amount", "cash_realisation", "evidence_ref", "benefit_source_key", "baseline_trend",
            "row", "in_scenario", "ramp_factor", "ramped", "capacity_rate", "class", "money", "enters", "claim", "other_claimants", "claim_rank", "removed", "removed_pv",
            "hour_claim", "hour_other", "hour_rank", "hours_removed"]
    header(ws, 3, cols)
    rows = []
    for b in sorted(d.c.rows["benefit_items"], key=lambda b: b["benefit_id"]):
        rows.append([b["benefit_id"], b["initiative_id"], idnum(b["initiative_id"]), idnum(b["benefit_id"]), b["benefit_type"], b["period_index"], b["starts_period"], num(b["amount"]),
                     b.get("cash_realisation") is True, b.get("evidence_ref") or None, b.get("benefit_source_key") or None, num(b.get("baseline_trend_per_period"))])
    assert len(rows) == NB
    put(ws, rows, 4)
    K, Fp, U, S, T, Z, P, C, D = (f"$K$4:$K${BL}", f"$F$4:$F${BL}", f"$U$4:$U${BL}", f"$S$4:$S${BL}", f"$T$4:$T${BL}", f"$Z$4:$Z${BL}", f"$P$4:$P${BL}", f"$C$4:$C${BL}", f"$D$4:$D${BL}")
    B = f"$B$4:$B${BL}"
    for i in range(NB):
        r = 4 + i
        f = {
            "M": f"=IFERROR(MATCH(Controls!$B$3&\"|\"&B{r},'Scenario inputs'!$A$4:$A${SIL},0),0)",
            "N": f"=IF(M{r}=0,FALSE,F{r}<Controls!$B$12)",
            "O": f"=IF(N{r},IF(INDEX('Scenario inputs'!$E$4:$E${SIL},M{r})=0,1,INDEX('Scenario inputs'!$F$4:$M${SIL},M{r},MIN(MAX(F{r}-G{r},0),INDEX('Scenario inputs'!$E$4:$E${SIL},M{r})-1)+1)),0)",
            "P": f"=H{r}*O{r}",
            "Q": f"=IF(M{r}=0,\"\",IF(ISNUMBER(INDEX('Scenario inputs'!$N$4:$N${SIL},M{r})),INDEX('Scenario inputs'!$N$4:$N${SIL},M{r}),\"\"))",
            "R": (f"=IF(NOT(N{r}),\"\",IF(OR(E{r}=\"cash\",E{r}=\"revenue\"),E{r},IF(E{r}=\"avoided_cost\",IF(AND(I{r},LEN(TRIM(J{r}))>0),\"avoided_cost_realised\",\"avoided_cost_beside\"),"
                  f"IF(ISNUMBER(Q{r}),\"capacity_converted\",\"capacity_hours\"))))"),
            "S": f"=IF(OR(R{r}=\"\",R{r}=\"capacity_hours\"),0,IF(R{r}=\"capacity_converted\",P{r}*Q{r},P{r}))",
            "T": f"=IF(OR(R{r}=\"cash\",R{r}=\"revenue\",R{r}=\"avoided_cost_realised\",R{r}=\"capacity_converted\"),1,0)",
            "U": f"=IF(AND(N{r},R{r}<>\"capacity_hours\",LEN(K{r})>0),1,0)",
            "V": f"=IF(U{r}=1,SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({U}=1)*({B}<>B{r})),0)",
            "W": (f"=IF(AND(U{r}=1,V{r}>0),1+SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({U}=1)*({S}>S{r}))+SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({U}=1)*({S}=S{r})*({T}>T{r}))"
                  f"+SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({U}=1)*({S}=S{r})*({T}=T{r})*({C}<C{r}))+SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({U}=1)*({S}=S{r})*({T}=T{r})*({C}=C{r})*({D}<D{r})),0)"),
            "X": f"=IF(AND(U{r}=1,V{r}>0,W{r}>1),S{r},0)",
            "Y": f"=IF(T{r}=1,X{r}/(1+Controls!$B$11)^F{r},0)",
            "Z": f"=IF(AND(N{r},E{r}=\"capacity\",LEN(K{r})>0),1,0)",
            "AA": f"=IF(Z{r}=1,SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({Z}=1)*({B}<>B{r})),0)",
            "AB": (f"=IF(AND(Z{r}=1,AA{r}>0),1+SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({Z}=1)*({P}>P{r}))+SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({Z}=1)*({P}=P{r})*({C}<C{r}))"
                   f"+SUMPRODUCT(({K}=K{r})*({Fp}=F{r})*({Z}=1)*({P}=P{r})*({C}=C{r})*({D}<D{r})),0)"),
            "AC": f"=IF(AND(Z{r}=1,AA{r}>0,AB{r}>1),P{r},0)",
        }
        for col, v in f.items():
            ws[f"{col}{r}"] = v
    widths(ws, {"A": 9, "B": 10, "J": 22, "K": 24, "R": 22})


def build_costs(wb, d: Data):
    ws = wb.create_sheet("Costs")
    title(ws, "900 cost rows (typed) and their half-open period range (formulas)",
          "first = period_index for a one-time cost, start_period for a recurring one; last_exclusive = period_index + 1 or end_period: a recurring cost is one row per period of [first, last_exclusive).")
    header(ws, 3, ["cost_id", "initiative_id", "cost_type", "cost_basis", "amount", "baseline_version", "period_index", "start_period", "end_period", "first", "last_exclusive", "periods"])
    rows = [[c["cost_id"], c["initiative_id"], c["cost_type"], c["cost_basis"], num(c["amount"]), c.get("baseline_version"), c.get("period_index"), c.get("start_period"), c.get("end_period")]
            for c in sorted(d.c.rows["cost_items"], key=lambda c: c["cost_id"])]
    assert len(rows) == NC
    put(ws, rows, 4)
    for i in range(NC):
        r = 4 + i
        ws[f"J{r}"] = f'=IF(C{r}="one_time",G{r},H{r})'
        ws[f"K{r}"] = f'=IF(C{r}="one_time",G{r}+1,I{r})'
        ws[f"L{r}"] = f"=K{r}-J{r}"
    widths(ws, {"A": 9, "B": 10, "C": 10})


def build_periods(wb, d: Data):
    ws = wb.create_sheet("Periods")
    title(ws, "The 36 monthly periods from 2026-01-01; a period is matured when its end (the start of the next one) is on or before the as-of (D-P9-22)")
    header(ws, 3, ["period_index", "period_start", "period_end", "matured"])
    for t in range(NP):
        r = 4 + t
        ws.cell(row=r, column=1, value=t).font = BLUE
        ws[f"B{r}"] = f"=DATE(YEAR(Controls!$B$7),MONTH(Controls!$B$7)+A{r},DAY(Controls!$B$7))"
        ws[f"C{r}"] = f"=DATE(YEAR(Controls!$B$7),MONTH(Controls!$B$7)+A{r}+1,DAY(Controls!$B$7))"
        ws[f"D{r}"] = f"=IF(C{r}<=Controls!$B$6,1,0)"
        ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = "yyyy-mm-dd"
    widths(ws, {"B": 12, "C": 12})


def build_cash_flow(wb):
    ws = wb.create_sheet("Cash flow")
    title(ws, "The dated cash flow of every initiative in the selected scenario (formulas only)",
          "Rows past the horizon or for initiatives outside the scenario are 0. costs = planned rows whose half-open range holds the period; the four classes enter the flow; capacity hours and the avoided cost beside never do.")
    header(ws, 3, ["initiative_id", "period_index", "member", "active", "costs", "cash", "revenue", "avoided_cost_realised", "capacity_converted", "benefits", "net_cash_flow", "cumulative",
                   "discount_factor", "discounted", "cumulative_discounted", "negative", "negative_discounted"])
    for k in range(NI):
        iid = f"init-{k + 1:03d}"
        for t in range(NP):
            r = 4 + k * NP + t
            ws.cell(row=r, column=1, value=iid).font = BLUE
            ws.cell(row=r, column=2, value=t).font = BLUE
            ws[f"C{r}"] = f"=IF(COUNTIF('Scenario inputs'!$A$4:$A${SIL},Controls!$B$3&\"|\"&A{r})>0,1,0)"
            ws[f"D{r}"] = f"=IF(AND(C{r}=1,B{r}<Controls!$B$12),1,0)"
            ws[f"E{r}"] = (f"=IF(D{r}=1,SUMPRODUCT((Costs!$B$4:$B${CL}=A{r})*(Costs!$D$4:$D${CL}=\"planned\")*(Costs!$J$4:$J${CL}<=B{r})*(Costs!$K$4:$K${CL}>B{r})*Costs!$E$4:$E${CL}),0)")
            for col, cls in (("F", "cash"), ("G", "revenue"), ("H", "avoided_cost_realised"), ("I", "capacity_converted")):
                ws[f"{col}{r}"] = f"=IF(D{r}=1,SUMIFS(Benefits!$S$4:$S${BL},Benefits!$B$4:$B${BL},A{r},Benefits!$F$4:$F${BL},B{r},Benefits!$R$4:$R${BL},\"{cls}\"),0)"
            ws[f"J{r}"] = f"=SUM(F{r}:I{r})"
            ws[f"K{r}"] = f"=J{r}-E{r}"
            ws[f"L{r}"] = f"=IF(D{r}=1,{'K' + str(r) if t == 0 else f'L{r - 1}+K{r}'},0)"
            ws[f"M{r}"] = f"=IF(D{r}=1,1/(1+Controls!$B$11)^B{r},0)"
            ws[f"N{r}"] = f"=K{r}*M{r}"
            ws[f"O{r}"] = f"=IF(D{r}=1,{'N' + str(r) if t == 0 else f'O{r - 1}+N{r}'},0)"
            ws[f"P{r}"] = f"=IF(AND(D{r}=1,L{r}<0),1,0)"
            ws[f"Q{r}"] = f"=IF(AND(D{r}=1,O{r}<0),1,0)"
    widths(ws, {"A": 10})


def build_initiatives(wb):
    ws = wb.create_sheet("Initiatives")
    title(ws, "Each initiative in the selected scenario: costs, benefits entering, NPV, both paybacks by the last crossing, checked against the domain",
          "A payback is the first period after the last negative cumulative; never negative = NO_INVESTMENT; ending negative = NOT_RECOVERED; no rows at all = NO_FLOWS (R9-12). Checks compare with `Domain reference` for the same scenario.")
    header(ws, 3, ["initiative_id", "member", "has_flows", "planned_costs", "benefits_entering", "net_total", "npv", "negatives", "last_negative_plus_1", "simple_payback", "simple_reason",
                   "simple_interpolated", "negatives_discounted", "last_negative_discounted_plus_1", "discounted_payback", "discounted_reason", "discounted_interpolated",
                   "domain npv", "domain simple", "domain discounted", "check"])
    CFA, CFB = f"'Cash flow'!$A$4:$A${CFL}", f"'Cash flow'!$B$4:$B${CFL}"
    for k in range(NI):
        r = 4 + k
        ws.cell(row=r, column=1, value=f"init-{k + 1:03d}").font = BLUE
        h = "Controls!$B$12"
        f = {
            "B": f"=IF(COUNTIF('Scenario inputs'!$A$4:$A${SIL},Controls!$B$3&\"|\"&A{r})>0,1,0)",
            "C": (f"=IF(B{r}=1,OR(SUMPRODUCT((Costs!$B$4:$B${CL}=A{r})*(Costs!$D$4:$D${CL}=\"planned\")*(Costs!$J$4:$J${CL}<{h})*(Costs!$K$4:$K${CL}>Costs!$J$4:$J${CL}))>0,"
                  f"SUMPRODUCT((Benefits!$B$4:$B${BL}=A{r})*(Benefits!$T$4:$T${BL}=1))>0),FALSE)"),
            "D": f"=SUMIFS('Cash flow'!$E$4:$E${CFL},{CFA},A{r})",
            "E": f"=SUMIFS('Cash flow'!$J$4:$J${CFL},{CFA},A{r})",
            "F": f"=SUMIFS('Cash flow'!$K$4:$K${CFL},{CFA},A{r})",
            "G": f"=IF(C{r},SUMIFS('Cash flow'!$N$4:$N${CFL},{CFA},A{r}),\"\")",
            "H": f"=COUNTIFS({CFA},A{r},'Cash flow'!$P$4:$P${CFL},1)",
            "I": f"=SUMPRODUCT(MAX(({CFA}=A{r})*('Cash flow'!$P$4:$P${CFL}=1)*({CFB}+1)))",
            "J": f"=IF(OR(NOT(C{r}),H{r}=0,I{r}={h}),\"\",I{r})",
            "K": f"=IF(NOT(C{r}),\"NO_FLOWS\",IF(H{r}=0,\"NO_INVESTMENT\",IF(I{r}={h},\"NOT_RECOVERED\",\"\")))",
            "L": (f"=IF(J{r}=\"\",\"\",(J{r}-1)+(-SUMIFS('Cash flow'!$L$4:$L${CFL},{CFA},A{r},{CFB},J{r}-1))/(SUMIFS('Cash flow'!$L$4:$L${CFL},{CFA},A{r},{CFB},J{r})"
                  f"-SUMIFS('Cash flow'!$L$4:$L${CFL},{CFA},A{r},{CFB},J{r}-1)))"),
            "M": f"=COUNTIFS({CFA},A{r},'Cash flow'!$Q$4:$Q${CFL},1)",
            "N": f"=SUMPRODUCT(MAX(({CFA}=A{r})*('Cash flow'!$Q$4:$Q${CFL}=1)*({CFB}+1)))",
            "O": f"=IF(OR(NOT(C{r}),M{r}=0,N{r}={h}),\"\",N{r})",
            "P": f"=IF(NOT(C{r}),\"NO_FLOWS\",IF(M{r}=0,\"NO_INVESTMENT\",IF(N{r}={h},\"NOT_RECOVERED\",\"\")))",
            "Q": (f"=IF(O{r}=\"\",\"\",(O{r}-1)+(-SUMIFS('Cash flow'!$O$4:$O${CFL},{CFA},A{r},{CFB},O{r}-1))/(SUMIFS('Cash flow'!$O$4:$O${CFL},{CFA},A{r},{CFB},O{r})"
                  f"-SUMIFS('Cash flow'!$O$4:$O${CFL},{CFA},A{r},{CFB},O{r}-1)))"),
            "R": f"=IFERROR(INDEX('Domain reference'!$C$40:$C$249,MATCH(Controls!$B$3&\"|\"&A{r},'Domain reference'!$A$40:$A$249,0)),\"\")",
            "S": f"=IFERROR(INDEX('Domain reference'!$D$40:$D$249,MATCH(Controls!$B$3&\"|\"&A{r},'Domain reference'!$A$40:$A$249,0)),\"\")",
            "T": f"=IFERROR(INDEX('Domain reference'!$E$40:$E$249,MATCH(Controls!$B$3&\"|\"&A{r},'Domain reference'!$A$40:$A$249,0)),\"\")",
            "U": (f"=IF(LEN(Controls!$B$4)>0,TRUE,IF(B{r}=0,AND(R{r}=\"\",E{r}=0),AND(ABS(N(G{r})-N(R{r}))<0.005,ISNUMBER(G{r})=ISNUMBER(R{r}),IF(S{r}=-1,J{r}=\"\",J{r}=S{r}),"
                  f"IF(T{r}=-1,O{r}=\"\",O{r}=T{r}))))"),
        }
        for col, v in f.items():
            ws[f"{col}{r}"] = v
        for col in "DEFGR":
            ws[f"{col}{r}"].number_format = "#,##0.00"
        for col in "LQ":
            ws[f"{col}{r}"].number_format = "0.0000"
    widths(ws, {"A": 11, "K": 16, "P": 16})


def build_domain_reference(wb, d: Data):
    ws = wb.create_sheet("Domain reference")
    title(ws, "PASTED VALUES: what the Python domain (changeworks.run_pipeline) and the SQL ledger publish, per scenario and per scenario x initiative; used by the check columns only")
    header(ws, 3, ["scenario_id", "sum_of_initiatives", "double_counted_pv", "attributable_npv", "planned_benefits", "residual", "planned_hours", "hours_residual",
                   "planned_costs (SQL 01)", "cash (SQL 01)", "revenue (SQL 01)", "avoided_cost_realised (SQL 01)", "capacity_converted (SQL 01)", "avoided_cost_beside (SQL 01)", "capacity_hours_not_converted (SQL 01)", "attributable_benefits"])
    res = json.loads(RESULTS_PATH.read_text())
    led = {r["scenario_id"]: r for r in res["01_portfolio_ledger__by_scenario"]}
    rows = []
    for sid in sorted(d.M["scenarios"]):
        s = d.M["scenarios"][sid]
        n, dc, l_ = s["npv_total"], s["double_count"], led[sid]
        rows.append([sid, num(n["sum_of_initiatives"]), num(dc["double_counted_pv"]), num(n["value"]), num(dc["planned_benefits"]), num(dc["residual"]), num(dc["planned_hours"]), num(dc["hours_residual"])]
                    + [num(l_[k]) for k in ("planned_costs", "cash", "revenue", "avoided_cost_realised", "capacity_converted", "avoided_cost_beside", "capacity_hours_not_converted")]
                    + [num(dc["attributable_benefits"])])
    put(ws, rows, 4)
    ws["A38"] = "per scenario x initiative (key = scenario|initiative): NPV, simple payback period, discounted payback period (-1 = null: no payback), simple reason, discounted reason"
    ws["A38"].font = BOLD
    header(ws, 39, ["key", "scenario_id", "npv", "simple_payback", "discounted_payback", "simple_reason", "discounted_reason"])
    r2 = []
    for sid in sorted(d.M["scenarios"]):
        for iid, x in sorted(d.M["scenarios"][sid]["initiatives"].items()):
            r2.append([f"{sid}|{iid}", sid, num(x["npv"]["value"]), -1 if x["payback"]["simple"]["period"] is None else x["payback"]["simple"]["period"],
                       -1 if x["payback"]["discounted"]["period"] is None else x["payback"]["discounted"]["period"], x["payback"]["simple"]["reason"], x["payback"]["discounted"]["reason"]])
    assert len(r2) == NSI
    put(ws, r2, 40)
    widths(ws, {"A": 20})


def build_portfolio(wb):
    ws = wb.create_sheet("Portfolio")
    title(ws, "The portfolio ledger of the selected scenario: four benefit classes apart, the attributable totals beside the plain sums, checked against the domain and the SQL ledger",
          "Column B is a formula over the data sheets; C is the reference (pasted, `Domain reference`); D is the tolerance; E is TRUE when they agree (or when the rate override is in use: see Controls!B4).")
    header(ws, 3, ["measure (USD unless hours)", "workbook (formula)", "reference (pasted)", "tolerance", "check"])

    def ref(col):
        return f"=INDEX('Domain reference'!${col}$4:${col}${SL},MATCH(Controls!$B$3,'Domain reference'!$A$4:$A${SL},0))"
    cf = lambda col: f"=SUM('Cash flow'!{col}4:{col}{CFL})"  # noqa: E731
    rows = [
        ("initiatives in the scenario", f"=SUM(Initiatives!B4:B{IL})", f"=INDEX(Scenarios!$H$4:$H${SL},MATCH(Controls!$B$3,Scenarios!$A$4:$A${SL},0))", 0),
        ("planned costs (inside the horizon, recurring rows expanded)", cf("E"), ref("I"), 0.005),
        ("cash benefits (enter the flow)", cf("F"), ref("J"), 0.005),
        ("revenue (enters the flow)", cf("G"), ref("K"), 0.005),
        ("avoided cost realised (cash_realisation and evidence)", cf("H"), ref("L"), 0.005),
        ("capacity converted to money (a conversion rate with staffing evidence)", cf("I"), ref("M"), 0.005),
        ("avoided cost beside the cash flow (never in it)", f"=SUMIFS(Benefits!S4:S{BL},Benefits!R4:R{BL},\"avoided_cost_beside\")", ref("N"), 0.005),
        ("capacity hours not converted (hours, never added to dollars)", f"=SUMIFS(Benefits!P4:P{BL},Benefits!R4:R{BL},\"capacity_hours\")", ref("O"), 0.005),
        ("planned benefits: the plain sum of the money claims", "=SUM(B6:B9)+B10", ref("E"), 0.005),
        ("double-counted amount (the residual): claims beyond the largest of a shared key", f"=SUM(Benefits!X4:X{BL})", ref("F"), 0.005),
        ("attributable benefits", "=B12-B13", ref("P"), 0.005),
        ("capacity hours planned (every capacity row, converted or not)", f"=SUMIFS(Benefits!P4:P{BL},Benefits!E4:E{BL},\"capacity\",Benefits!N4:N{BL},TRUE)", ref("G"), 0.005),
        ("hours residual (shared capacity counted twice)", f"=SUM(Benefits!AC4:AC{BL})", ref("H"), 0.005),
        ("NPV, plain sum of the initiative NPVs", f"=SUM(Initiatives!G4:G{IL})", ref("B"), 0.005),
        ("present value of the claims counted twice (double_counted_pv)", f"=SUM(Benefits!Y4:Y{BL})", ref("C"), 0.005),
        ("NPV, attributable (the headline)", "=B17-B18", ref("D"), 0.005),
        ("initiatives that never pay back (simple)", f"=SUMPRODUCT((Initiatives!B4:B{IL}=1)*(Initiatives!K4:K{IL}<>\"\"))", None, None),
    ]
    for i, (lab, f, refv, tol) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=2, value=f).number_format = "#,##0.00"
        if refv is not None:
            ws.cell(row=r, column=3, value=refv).number_format = "#,##0.00"
            ws.cell(row=r, column=4, value=tol)
            ws.cell(row=r, column=5, value=f"=IF(LEN(Controls!$B$4)>0,TRUE,ABS(B{r}-C{r})<=D{r})")
    widths(ws, {"A": 78, "B": 20, "C": 20, "E": 8})
    return 4, 4 + len(rows) - 2


def build_golden(wb):
    ws = wb.create_sheet("Golden case")
    title(ws, "The known answer (P09-GOLD-01) by hand, with formulas, against tests/fixtures/golden/expected.json",
          "Cost 1000 at period 0, net inflows 600 at periods 1 and 2, 10 % annual on annual periods. Blue cells are typed from the fixture; every figure in column C is a formula.")
    header(ws, 3, ["input or figure", "value (typed)", "formula result", "expected.json / hand answer", "check"])
    e = EXPECTED["metrics"]["scenarios"]["scn-001"]
    i1, i2, i3 = (e["initiatives"][k] for k in ("init-001", "init-002", "init-003"))
    rows = [
        ("rate per period (annual on annual periods)", 0.10, None, None, None),
        ("init-001 net cash flow at period 0", -1000, None, None, None),
        ("init-001 net cash flow at period 1 (600 cash; the 400 capacity hours stay out)", 600, None, None, None),
        ("init-001 net cash flow at period 2", 600, None, None, None),
        ("init-001 NPV = -1000 + 600/1.1 + 600/1.1^2", None, "=B5+B6/(1+B4)+B7/(1+B4)^2", float(i1["npv"]["value"]), "=ROUND(C8,2)=D8"),
        ("init-001 NPV exactly 5000/121", None, "=5000/121", None, "=ABS(C8-C9)<1E-9"),
        ("init-001 simple payback period (first period after the last negative cumulative)", None, "=IF(B5+B6<0,2,1)", i1["payback"]["simple"]["period"], "=C10=D10"),
        ("init-001 simple payback, interpolated (t-1) + (-cum[t-1]) / (cum[t] - cum[t-1])", None, "=1+(-(B5+B6))/((B5+B6+B7)-(B5+B6))", float(i1["payback"]["simple"]["interpolated"]), "=ROUND(C11,4)=D11"),
        ("init-001 discounted payback period", None, "=IF(B5+B6/(1+B4)<0,2,1)", i1["payback"]["discounted"]["period"], "=C12=D12"),
        ("init-001 discounted payback, interpolated", None, "=1+(-(B5+B6/(1+B4)))/((B7/(1+B4)^2))", float(i1["payback"]["discounted"]["interpolated"]), "=ROUND(C13,4)=D13"),
        ("init-002 net cash flow at periods 0, 1, 2 (ramped)", "-500, 100, 425", None, None, None),
        ("init-002 NPV = -500 + 100/1.1 + 425/1.1^2 = -7000/121", None, "=-500+100/(1+B4)+425/(1+B4)^2", float(i2["npv"]["value"]), "=ROUND(C15,2)=D15"),
        ("init-003 NPV = -800 + 150/1.1^2 = -81800/121 (never pays back)", None, "=-800+150/(1+B4)^2", float(i3["npv"]["value"]), "=ROUND(C16,2)=D16"),
        ("sum of the three initiative NPVs = -83800/121", None, "=C8+C15+C16", float(e["npv_total"]["sum_of_initiatives"]), "=ROUND(C17,2)=D17"),
        ("the 75.00 claimed twice (cost-line:CC-4410, period 2), present value 75/1.1^2 = 7500/121", 75, "=B18/(1+B4)^2", float(e["double_count"]["double_counted_pv"]), "=ROUND(C18,2)=D18"),
        ("attributable NPV = plain sum - present value counted twice = -8300/11", None, "=C17-C18", float(e["npv_total"]["value"]), "=ROUND(C19,2)=D19"),
        ("simple and discounted paybacks are labelled separately and both are found for init-001", None, "=AND(C10=2,C12=2,ROUND(C11,4)<>ROUND(C13,4))", True, "=C20=D20"),
    ]
    for i, (lab, typed, formula, exp, chk) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=lab)
        if typed is not None:
            ws.cell(row=r, column=2, value=typed).font = BLUE
        if formula is not None:
            ws.cell(row=r, column=3, value=formula)
        if exp is not None:
            ws.cell(row=r, column=4, value=exp)
        if chk is not None:
            ws.cell(row=r, column=5, value=chk)
    widths(ws, {"A": 84, "B": 16, "C": 18, "D": 26})
    return [8, 9, 10, 11, 12, 13, 15, 16, 17, 18, 19, 20]


def build_checks(wb, p_first, p_last, golden_rows):
    ws = wb.create_sheet("Checks")
    title(ws, "Checks: TRUE when every comparison with the Python domain, the SQL ledger and the golden hand answer holds")
    ws["A2"] = "All checks pass"
    ws["A2"].font = BOLD
    g = ",".join(f"'Golden case'!E{r}" for r in golden_rows)
    ws["B2"] = f"=AND(Portfolio!E{p_first}:E{p_last},Initiatives!U4:U{IL},{g},B5,B6,B7)"
    ws["A4"], ws["B4"] = "check", "result"
    ws["A5"] = "Every benefit row has a class when its initiative is in the scenario and the row is inside the horizon, and none otherwise"
    ws["B5"] = f"=SUMPRODUCT((Benefits!N4:N{BL}=TRUE)*(Benefits!R4:R{BL}=\"\"))+SUMPRODUCT((Benefits!N4:N{BL}=FALSE)*(Benefits!R4:R{BL}<>\"\"))=0"
    ws["A6"] = "Every recurring cost row expands to end - start periods and every one-time row to exactly one (Costs!L)"
    ws["B6"] = f"=SUMPRODUCT((Costs!C4:C{CL}=\"one_time\")*(Costs!L4:L{CL}<>1))=0"
    ws["A7"] = "The scenario switch names a stored scenario"
    ws["B7"] = f"=COUNTIF(Scenarios!A4:A{SL},Controls!B3)=1"
    for c in ("A4", "B4"):
        ws[c].font = BOLD
    widths(ws, {"A": 110, "B": 12})


def build_controls(wb, d: Data):
    ws = wb.create_sheet("Controls")
    title(ws, "Controls", "Change B3 to any stored scenario: every sheet follows it. B4 overrides the period rate (blank = the scenario's own rate; the checks against the domain stand down while it is set). Blue cells are inputs.")
    ws["A3"], ws["B3"] = "Scenario (switch)", DEFAULT_SCENARIO
    dv = DataValidation(type="list", formula1=f"Scenarios!$A$4:$A${SL}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    ws["A4"], ws["B4"] = "Period-rate override (blank = the scenario's rate)", None
    ws["A5"], ws["B5"] = "Period basis of the dataset", d.c.basis
    ws["A6"], ws["B6"] = "As-of (UTC)", serial("2026-10-05T00:00:00Z")
    ws["B6"].number_format = "yyyy-mm-dd hh:mm"
    ws["A7"], ws["B7"] = "First period starts", serial("2026-01-01T00:00:00Z")
    ws["B7"].number_format = "yyyy-mm-dd"
    for c in ("B3", "B4", "B5", "B6", "B7"):
        ws[c].font = BLUE
    ws["A9"], ws["B9"] = "Selected scenario's declared rate and basis", f"=INDEX(Scenarios!$D$4:$D${SL},MATCH($B$3,Scenarios!$A$4:$A${SL},0))&\" \"&INDEX(Scenarios!$C$4:$C${SL},MATCH($B$3,Scenarios!$A$4:$A${SL},0))"
    ws["A10"], ws["B10"] = "Rate source", f"=INDEX(Scenarios!$J$4:$J${SL},MATCH($B$3,Scenarios!$A$4:$A${SL},0))"
    ws["A11"], ws["B11"] = "Effective period rate used", f"=IF(LEN(B4)>0,B4,INDEX(Scenarios!$I$4:$I${SL},MATCH($B$3,Scenarios!$A$4:$A${SL},0)))"
    ws["B11"].number_format = "0.000000"
    ws["A12"], ws["B12"] = "Horizon (periods)", f"=INDEX(Scenarios!$F$4:$F${SL},MATCH($B$3,Scenarios!$A$4:$A${SL},0))"
    ws["A13"], ws["B13"] = "Matured periods at the as-of", "=SUM(Periods!D4:D39)"
    ws["A15"], ws["B15"] = "All checks pass (Checks!B2)", "=Checks!B2"
    ws["A15"].font = BOLD
    ws["A16"], ws["B16"] = "Attributable NPV of the selected scenario (Portfolio!B19)", "=Portfolio!B19"
    ws["B16"].number_format = "#,##0.00"
    widths(ws, {"A": 58, "B": 26})


def build_readme(wb):
    ws = wb.create_sheet("Read me")
    title(ws, "ChangeWorks analyst workbook (P09-ART-05)")
    rows = [
        ("What it is", f"A formula-driven recomputation of the ChangeWorks scenario model over the published synthetic snapshot cw-20261026-20261026 (dataset hash {DATASET_HASH}, code revision {CODE_REVISION}, "
                       "as-of 2026-10-05T00:00:00Z): 24 initiatives, 30 scenarios, 900 cost rows, 700 benefit rows, 36 monthly periods from 2026-01."),
        ("The decision", "Which transformation initiative should be funded and how will success be established? The workbook shows each initiative's dated cash flow, NPV and both paybacks in the "
                         "selected scenario, the four benefit classes kept apart and the attributable portfolio NPV beside the plain sum. Nothing here funds anything: a selected scenario is a proposal."),
        ("Synthetic data", "Everything is synthetic (a seeded generator). Observed = what this one draw shows; constructed = what the generator built in on purpose (its designed relationships DR-01..DR-07 "
                           "and injected defects INJ-01..INJ-12); every NPV is a plan estimate; no benefit to a real organisation is claimed."),
        ("How to use it", "Open the .xlsx in any spreadsheet application; nothing needs installing to read the cached values. Controls!B3 switches the scenario; Controls!B4 overrides the period rate; "
                          "Checks!B2 is TRUE when every figure agrees with the Python domain, the SQL ledger and the golden hand answer."),
        ("Recalculation status", "openpyxl wrote formulas only; LibreOffice (headless) recalculated the delivered file. Microsoft Excel was NOT run: nothing here was opened or recalculated by Excel (docs/workbook.md)."),
        ("Sheets", "Controls, Checks, Portfolio, Initiatives, Cash flow, Golden case, Assumptions (read-this-first limits), data sheets (Benefits, Costs, Scenario inputs, Scenarios, Periods), "
                   "Domain reference (pasted values), SQL results (pasted)."),
        ("Colour", "Blue = typed input or pasted value; black = formula."),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 24, "B": 150})


ASSUMPTIONS = [
    ("A-01", "Four benefit classes, never summed into one number (D-01): cash and revenue enter the net cash flow; avoided cost enters only with cash_realisation and a non-blank evidence_ref, "
             "otherwise it is reported beside the flow; capacity hours never enter unless the scenario's capacity_hourly_rate carries a staffing_change_evidence_ref, when each row becomes a separate "
             "derived cash item (hours x rate)."),
    ("A-02", "A recurring cost is one row per period of its half-open range [start_period, end_period); a one-time cost sits in exactly its period_index (D-01, D-P9-16). Only planned rows enter the flow."),
    ("A-03", "NPV = sum over t < horizon of net_t / (1 + r)^t with r the period rate: the declared rate when its basis is the period basis, else the recorded 6-place monthly conversion (D-P9-06)."),
    ("A-04", "Payback (simple and discounted, labelled separately) = the first period after the last negative cumulative, interpolated (t-1) + (-cum[t-1]) / (cum[t] - cum[t-1]); null with a reason "
             "(NO_INVESTMENT, NOT_RECOVERED, NO_FLOWS), never 0 (D-09 R9-12)."),
    ("A-05", "Double count (D-01, R9-02, R9-03): money claims sharing a benefit_source_key in one period across two or more initiatives count once, at the largest (ties: the claim that enters the flow, "
             "then the smaller initiative and benefit id). The attributable NPV is the plain sum less the present value of the claims counted twice. Hours follow the same rule separately."),
    ("A-06", "The adoption ramp multiplies each benefit row from its own starts_period, the last value held (D-P9-02); the unramped plan is not shown here (it is in the domain's output)."),
    ("W-01", "Observed, constructed and hypothetical results are kept apart: the NPVs and paybacks are plan estimates of one synthetic draw; the designed relationships are constructed; no real benefit is claimed."),
    ("W-02", "TYPED from the independent Python side, not replayed by formulas: which assumption version each scenario uses for each initiative (the pinned one, else the latest created by the "
             "scenario's created_at) and therefore its ramp and conversion rate (`Scenario inputs`). Every result is compared with the Python domain after recalculation."),
    ("W-03", "Recalculation: openpyxl writes formulas only; the delivered file was recalculated by LibreOffice headless and carries its cached values; Excel was NOT run. fullCalcOnLoad is set so "
             "Excel recalculates on open. Functions used: SUMIFS, COUNTIFS, COUNTIF, SUMPRODUCT, INDEX, MATCH, IF, IFERROR, ISNUMBER, AND, OR, NOT, MIN, MAX, SUM, ABS, ROUND, LEN, TRIM, DATE, YEAR, MONTH, DAY, N (no _xlfn functions)."),
    ("W-04", "Spreadsheet arithmetic is binary floating point; the domain's figures are exact rationals rounded once. Money checks use a tolerance of 0.005 USD; the exact figures are in the SQL analysis and the domain."),
    ("W-05", "The rate override (Controls!B4) is an analyst's what-if: while it is set the checks against the domain stand down (they would compare different rates). The workbook has no macros, no external links and no credentials."),
    ("W-06", "Rows recorded after a scenario's created_at would not belong to it (R9-01); no row of the published world carries a recorded_at, so every row is used."),
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions and limits")
    header(ws, 3, ["id", "statement"])
    for i, (k, v) in enumerate(ASSUMPTIONS, start=4):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 8, "B": 150})


def build_sql_results(wb):
    ws = wb.create_sheet("SQL results (pasted)")
    title(ws, "PASTED VALUES from data/sql/analysis/run_all.py (artifacts/analysis-results.json)", "The headline result sets; reconciliation: every row equal across the SQL, the reference SQL, the Python domain, the manifests and the independent Python side.")
    res = json.loads(RESULTS_PATH.read_text())
    meta = res["meta"][0]
    r = 4
    ws.cell(row=r, column=1, value=f"reconciliation rows {meta['reconciliation_rows']}, equal {meta['reconciliation_equal']}; result sets {meta['result_sets']}, rows {meta['result_rows']}; "
                                     f"dataset {meta['dataset_version']} / {meta['dataset_hash']}; code {meta['code_revision']}").font = BOLD
    r += 2
    for key in ("02_initiative_npv__portfolio_scn025", "03_double_count__by_source_key", "04_payback__init015_by_scenario", "07_adoption_training__pooled", "08_forecast_revisions__init003_versions",
                "09_roadmap_capacity__deployment_before_feasible_start", "11_proposal_screen__options", "10_golden_known_answer__init001_answer"):
        rows = res[key]
        ws.cell(row=r, column=1, value=key).font = BOLD
        r += 1
        cols = list(rows[0])
        header(ws, r, cols)
        r += 1
        for row in rows:
            for j, c in enumerate(cols, start=1):
                v = row[c]
                if isinstance(v, str) and re.fullmatch(r"-?\d+(\.\d+)?", v):
                    v = float(v) if "." in v else int(v)
                ws.cell(row=r, column=j, value=v).font = BLUE
            r += 1
        r += 1
    widths(ws, {"A": 34})


def recalc(path: Path) -> None:
    """Open with LibreOffice headless and save a recalculated copy over ``path`` (cached values included)."""
    if shutil.which("soffice") is None:
        raise SystemExit("soffice not installed; the file has formulas but no cached values")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in" / path.name
        src.parent.mkdir()
        shutil.copy(path, src)
        out = Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=1800)
        shutil.copy(out / path.name, path)
    set_full_calc_on_load(path)


def set_full_calc_on_load(path: Path) -> None:
    """LibreOffice drops fullCalcOnLoad; put it back (workbook.xml only) so Excel recalculates when it opens the file."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                t = data.decode("utf-8")
                if "fullCalcOnLoad" not in t:
                    assert "<calcPr " in t
                    t = t.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ', 1)
                data = t.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build(path: Path, scenario: str = DEFAULT_SCENARIO) -> None:
    data = Data()
    wb = Workbook()
    build_readme(wb)
    build_controls(wb, data)
    wb["Controls"]["B3"] = scenario
    build_scenarios(wb, data)
    build_scenario_inputs(wb, data)
    build_benefits(wb, data)
    build_costs(wb, data)
    build_periods(wb, data)
    build_cash_flow(wb)
    build_initiatives(wb)
    build_domain_reference(wb, data)
    p_first, p_last = build_portfolio(wb)
    golden_rows = build_golden(wb)
    build_checks(wb, p_first, p_last, golden_rows)
    build_assumptions(wb)
    build_sql_results(wb)
    del wb["Sheet"]
    wb._sheets = [wb[n] for n in SHEET_ORDER]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main(argv: list[str]) -> int:
    build(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    if "--recalc" in argv:
        recalc(OUT)
        print("recalculated with LibreOffice headless; cached values written")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
