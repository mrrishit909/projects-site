"""Build artifacts/invoicebridge-analyst-workbook.xlsx with openpyxl (P05-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: the published snapshot ``web/public/data`` read through the Python domain (the book is derived once, in Python, at the three
snapshot cutoffs: the per-invoice MATCH OUTCOME at each cutoff is typed into ``Invoices``), tests/fixtures/golden/expected.json,
artifacts/analysis-results.json (written by data/sql/analysis/run_all.py) and web/public/data/manifest.json. openpyxl cannot calculate, so
the file it writes holds formulas only. With --recalc, LibreOffice (soffice) opens it headless and saves a recalculated copy over the same
path, which then also carries cached values. Everything on ``Metrics`` is a live formula over the data sheets and the ``Cutoff`` switch
(march / june / asof): the line amounts, subtotal, tax and total of every invoice, the policy-valid approvals (separation of duties, role
and limit by the policy in force at the decision), the export rows (idempotent batches, a row needs a prior approval), credits, the
consumption of each purchase-order line, the backlog and the rates, per currency. The one thing the spreadsheet does not do is replay the
three-way match chain (it needs the cumulative consumption in a particular order): that outcome is typed per invoice and cutoff from the
domain (limit W-02). Money is per currency and is never summed across currencies. Excel is never run.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "domain"), str(ROOT), str(ROOT / "data" / "sql" / "analysis")]

import python_side as ps  # noqa: E402
from invoicebridge import compute_metrics, run_pipeline  # noqa: E402
from invoicebridge.metrics import window_metrics  # noqa: E402
from invoicebridge.match import quantities  # noqa: E402
from invoicebridge.timeutil import parse_ts  # noqa: E402

OUT = ROOT / "artifacts" / "invoicebridge-analyst-workbook.xlsx"
RESULTS_PATH = ROOT / "artifacts" / "analysis-results.json"
MANIFEST = json.loads((ROOT / "web/public/data/manifest.json").read_text())
EXPECTED = json.loads((ROOT / "tests/fixtures/golden/expected.json").read_text())
CODE_REVISION = "2dcc460"
DATASET_HASH = "1f001a43a7031d4260f5abb3b9f7a83f6e1b48397e85cf8b04e0b4d9e052162e"
NEVER = 1e15            # "no such instant" (an approval or export that does not exist), so a numeric comparison with a cutoff is always false
CUTS = (("march", ps.C1), ("june", ps.C2), ("asof", ps.ASOF))
CURRENCIES = ("USD", "EUR", "GBP")

HDR = PatternFill("solid", fgColor="23303A")
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
EPOCH = datetime(1970, 1, 1, tzinfo=timezone.utc)


def epoch(dt: datetime | None) -> int | None:
    return None if dt is None else int((dt - EPOCH).total_seconds())


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


class Cols:
    """Column letters by name for a data sheet, so formulas never hardcode a position."""

    def __init__(self, names: list[str]) -> None:
        self.names = names
        self.letter = {n: get_column_letter(i + 1) for i, n in enumerate(names)}

    def __getitem__(self, name: str) -> str:
        return self.letter[name]

    def rng(self, sheet: str, name: str, last: int) -> str:
        c = self.letter[name]
        return f"'{sheet}'!${c}$2:${c}${last}"


# ---- the derived data (Python domain, once) ------------------------------------------------------------------------------------
def outcome(ctx: ps.Ctx, iid: str, cut: str) -> tuple[str, str]:
    b = ctx.books[cut]
    if iid in b.merged_into:
        return "merged", ""
    if iid in b.review:
        return "needs_review", ""
    if iid in b.rejected:
        return "rejected", ""
    r = b.run.invoices.get(iid)
    if r is None:
        return "received", ""
    return ("matched" if r.matched else "exception"), "|".join(r.reasons)


def domain_reference(ctx: ps.Ctx) -> dict[tuple[str, str], dict[str, float | None]]:
    """The domain's figures for each cutoff and scope (USD, EUR, GBP, ALL) as numbers (pasted into 'Domain reference'; read only by the checks)."""
    out: dict[tuple[str, str], dict[str, float | None]] = {}
    cuts = [parse_ts(c) for _, c in CUTS]
    for name, c in CUTS:
        when = parse_ts(c)
        dm = compute_metrics(ctx.pop, when, [])
        groups = {g["key"]: g["metrics"] for g in window_metrics(ctx.pop, (None, when + timedelta(microseconds=1)), [], "currency")["groups"]}
        run = ctx.books[c].run
        hist_open = [e for e in ctx.exc if ps.is_open(e, when)]
        first_asg: dict[str, list] = {}
        for a in ctx.asg:
            first_asg.setdefault(a["exception_id"], []).append(a["at"])
        unowned = sum(1 for e in hist_open if not any(t <= when for t in first_asg.get(e["exception_id"], [])))

        def pick(m: dict, k: str, cur_pos: dict | None) -> dict[str, float | None]:
            r = m["rates"]
            ex = m["exceptions"]["by_reason"]
            q = m["quantities"]
            v: dict[str, float | None] = {
                "eligible_invoices": r["eligible_invoices"], "matched_invoices": m["matched"]["invoice_count"],
                "matched_amount_before_tax": m["matched"]["amount_before_tax"], "matched_tax": m["matched"]["tax"], "matched_total": m["matched"]["total"],
                "obligations_count": m["obligations"]["count"], "obligations_total": m["obligations"]["total"],
                "credit_total": m["payable"]["credit_total"], "approved_payable": m["payable"]["approved_payable"],
                "exported_invoices": len(m["export"]["rows"]), "accepted_export_total": m["export"]["accepted_export_total"],
                "export_difference": m["export"]["difference"],
                "open_exceptions": m["exceptions"]["open_count"], "open_QUANTITY_BLOCKED": ex["QUANTITY_BLOCKED"], "open_PRICE_VARIANCE": ex["PRICE_VARIANCE"],
                "open_NO_PO_LINE": ex["NO_PO_LINE"], "open_ROUNDING_EXCEPTION": ex["ROUNDING_EXCEPTION"], "open_CURRENCY_MISMATCH": ex["CURRENCY_MISMATCH"],
                "straight_through_rate": r["straight_through_rate"], "exception_rate": r["exception_rate"],
                "ever_excepted_rate": r.get("ever_excepted_rate"), "never_excepted_rate": r.get("straight_through_never_excepted_rate"),
                "mean_minutes_to_approval": m["processing"]["mean_minutes"], "pending_invoices": m["processing"]["pending_count"],
                "duplicates_merged": m["duplicates"]["rows_merged"],
                "ordered_quantity": q["ordered_quantity"], "net_received_quantity": q["net_received_quantity"],
                "accepted_invoiced_quantity": q["accepted_invoiced_quantity"], "open_to_receive_quantity": q["open_to_receive_quantity"],
                "open_to_invoice_quantity": q["open_to_invoice_quantity"], "received_not_invoiced_quantity": q["received_not_invoiced_quantity"],
                "received_not_invoiced_value": q["received_not_invoiced_value"]}
            return {k2: (None if x is None else float(x)) for k2, x in v.items()}

        for cur in CURRENCIES:
            g = groups[cur]
            po = {l.po_line_id for l in ctx.pop.po_lines if ctx.po_cur[l.po_id] == cur}
            g = {**g, "quantities": quantities(run, po)}
            out[(name, cur)] = pick(g, name, None)
            out[(name, cur)]["duplicates_merged"] = None          # a merge is one row of the invoice's own currency: compared on the whole only
        allv = pick(dm, name, None)
        allv["received_not_invoiced_value"] = None                 # not summed across currencies
        for k in ("matched_amount_before_tax", "matched_tax", "matched_total", "obligations_total", "credit_total", "approved_payable", "accepted_export_total", "export_difference"):
            allv[k] = None                                          # money is never summed across currencies
        allv["unowned_open_exceptions_history"] = float(unowned)
        allv["open_exceptions_history"] = float(len(hist_open))
        out[(name, "ALL")] = allv
    return out


def invoice_rows(ctx: ps.Ctx):
    rows = []
    for inv in sorted(ctx.pop.invoices, key=lambda i: i.invoice_id):
        st = [outcome(ctx, inv.invoice_id, c) for _, c in CUTS]
        rows.append([inv.invoice_id, inv.vendor_id, inv.currency, epoch(inv.imported_at), float(inv.tax_rate), float(inv.freight), float(inv.invoice_discount),
                     float(inv.declared_total), inv.created_by, st[0][0], st[0][1], st[1][0], st[1][1], st[2][0], st[2][1]])
    return rows


# ---- sheets ---------------------------------------------------------------------------------------------------------------------
def build_readme(wb):
    ws = wb.active
    ws.title = "Read me"
    title(ws, "InvoiceBridge analyst workbook (P05-ART-05)",
          "SYNTHETIC DATA: every vendor, purchase order and invoice is generated (seed 20261026); nothing describes a real company, supplier or payment.")
    rows = [
        ("What it is", "A formula-driven workbook over the published synthetic snapshot of InvoiceBridge (40 vendors, 3,000 purchase-order lines, 2,800 imported invoices, 12 months to 2026-09-30). It recomputes the headline metrics of the contract per currency with live formulas, so an analyst can follow a number cell by cell, and it can be switched between three cutoffs: the March month-end, the June month-end and the as-of 2026-10-01T02:00:00Z."),
        ("The decision it supports", "Which invoices may advance and which require exception review? Metrics shows, per currency, what matched, what is an obligation (approved), what is blocked and why, what was exported and what is payable once credits are applied."),
        ("How to read it without installing anything", "Open the .xlsx in any spreadsheet application, or read docs/workbook.md (it lists every sheet and the checked values). The delivered file already carries values calculated by LibreOffice."),
        ("How to use the switch", "Cutoff!B3 takes march, june or asof. Every figure on Metrics follows it. The check column compares each figure with the Python domain's value for the same cutoff and currency (pasted on 'Domain reference')."),
        ("Money is per currency", "The contract has no dated conversion rate, so no figure adds one currency to another: Metrics has one column per currency (USD, EUR, GBP) and the all-currency column carries counts and units only, never money."),
        ("Observed, simulated, hypothetical", "Everything here is OBSERVED in a synthetic history that was generated with designed relationships (docs/executive-memo.md). Nothing is a simulation and no benefit to a real accounts-payable team is claimed. The generator approves only invoices that matched without interruption, rejects only invoices that never matched and closes each period with a batch that exports everything approved, so the export reconciliation is zero by construction."),
        ("Inputs (blue) and formulas (black)", "Blue cells are typed inputs: the facts (vendors, lines, receipts, decisions, export requests, credits, exceptions, assignments), the policies, the tolerance and the per-invoice MATCH OUTCOME at each cutoff (matched, exception with its reasons, merged, rejected). Black cells are formulas."),
        ("What the formulas do", "Line amounts, subtotal, tax (once per invoice on the unrounded subtotal, half up), total; the policy in force at each decision and whether the decision is valid (approver is not the creator, role and limit); the first valid approval; the export rows (a repeated key or batch id returns the prior batch; a row needs an approval at or before its request); credits per approved invoice; the net received and consumed quantity of every purchase-order line; the backlog and the rates."),
        ("What the formulas do NOT do", "They do not replay the three-way match chain: whether an invoice matched at a cutoff depends on the cumulative consumption of earlier invoices in a particular order and is typed from the domain (limit W-02). Hence a wrong outcome would not be caught by this file alone; data/sql/analysis recomputes the same book from facts in SQL and reconciles it (docs/sql-analysis.md)."),
        ("Provenance", f"Dataset hash {DATASET_HASH} (generator, seed 20261026, dataset ib-20261026-20261026), contract v1, code revision {CODE_REVISION}. Built by artifacts/build_workbook.py with openpyxl 3.1.5."),
        ("Recalculation status", "openpyxl wrote formulas only; LibreOffice (headless) recalculated the delivered file. Microsoft Excel was NOT run: nothing here was opened or recalculated by Excel (docs/workbook.md)."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=k).font = BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 38, "B": 140})


def build_cutoff(wb):
    ws = wb.create_sheet("Cutoff")
    title(ws, "The cutoff switch", "Type march, june or asof in B3 (a list). Every figure on Metrics, Invoices, PO lines, Lines and Exceptions that depends on a cutoff follows it.")
    ws["A3"], ws["B3"] = "selected cutoff", "asof"
    ws["B3"].font = BLUE
    dv = DataValidation(type="list", formula1='"march,june,asof"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    header(ws, 5, ["name", "meaning", "instant (UTC, inclusive)", "epoch seconds (UTC)", "position"])
    meaning = {"march": "month-end 2026-03-31", "june": "month-end 2026-06-30", "asof": "the as-of of the snapshot (the last nightly matching run)"}
    for r, (name, c) in enumerate(CUTS, start=6):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=meaning[name])
        ws.cell(row=r, column=3, value=c)
        ws.cell(row=r, column=4, value=epoch(parse_ts(c))).font = BLUE
        ws.cell(row=r, column=5, value=r - 5)
    ws["A10"], ws["B10"] = "selected epoch", "=INDEX($D$6:$D$8,MATCH($B$3,$A$6:$A$8,0))"
    ws["A11"], ws["B11"] = "selected position", "=MATCH($B$3,$A$6:$A$8,0)"
    ws["A12"], ws["B12"] = "rounding tolerance per invoice (D-05)", 0.02
    ws["B12"].font = BLUE
    widths(ws, {"A": 34, "B": 52, "C": 28, "D": 22, "E": 10})


def build_policies(wb, ctx):
    ws = wb.create_sheet("Policies")
    title(ws, "Approval policies (typed from approval_policies.json)")
    header(ws, 4, ["effective_from (epoch, ascending)", "amount_limit", "reviewer_role", "escalation_role (none when absent)", "policy_id"])
    for r, p in enumerate(sorted(ctx.pop.approval_policies, key=lambda p: p.effective_from), start=5):
        for c, v in enumerate([epoch(p.effective_from), float(p.amount_limit), p.reviewer_role, p.escalation_role or "none", p.policy_id], start=1):
            ws.cell(row=r, column=c, value=v).font = BLUE
    ws["A9"] = "A limit is compared with the invoice total in the invoice's OWN currency, with no conversion (the contract has no dated rate)."
    widths(ws, {"A": 32, "B": 14, "C": 16, "D": 34, "E": 12})
    return 5, 4 + len(ctx.pop.approval_policies)


def put(ws, rows: list[list], first_row: int = 2, blue: bool = True):
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=first_row + i, column=1 + j, value=v)
            if blue:
                c.font = BLUE


def build_data(wb, ctx, pol_rng):
    """Data sheets. Returns the column maps and last rows used by the formulas of Metrics."""
    inv_typed = invoice_rows(ctx)
    n_inv = len(inv_typed)
    lines = sorted((l for i in ctx.pop.invoices for l in i.lines), key=lambda l: l.invoice_line_id)
    n_lines = len(lines)
    recs = sorted(ctx.pop.receipts, key=lambda r: r.receipt_id)
    pol_lines = sorted(ctx.pop.po_lines, key=lambda l: l.po_line_id)
    po_ordered = {p.po_id: p for p in ctx.pop.purchase_orders}
    decisions = sorted(ctx.pop.approvals, key=lambda a: (a.decided_at, a.approval_id))
    credits = sorted(ctx.pop.credit_notes, key=lambda c: c.credit_id)
    reqs = list(ctx.pop.export_requests)
    assert all(reqs[i].exported_at <= reqs[i + 1].exported_at for i in range(len(reqs) - 1)), "export requests must be chronological"
    exc = sorted(ctx.rows["exceptions"], key=lambda e: e["exception_id"])
    asg = sorted(ctx.rows["exception_assignments"], key=lambda a: a["assignment_id"])
    exp_lines = [(seq, iid) for seq, r in enumerate(reqs, start=1) for iid in r.invoice_ids]

    I = Cols(["invoice_id", "vendor_id", "currency", "imported_epoch", "tax_rate", "freight", "invoice_discount", "declared_total", "created_by",
              "state_march", "reasons_march", "state_june", "reasons_june", "state_asof", "reasons_asof",
              "raw_extended", "subtotal", "tax", "total", "commit_epoch", "export_epoch",
              "sel_state", "sel_reasons", "sel_in_book", "sel_approved", "sel_matched", "sel_good", "sel_exported", "sel_flagged",
              "has_QUANTITY_BLOCKED", "has_PRICE_VARIANCE", "has_NO_PO_LINE", "has_ROUNDING_EXCEPTION", "has_CURRENCY_MISMATCH",
              "sel_credit", "sel_payable", "sel_minutes", "sel_ever", "sel_merged"])
    L = Cols(["invoice_line_id", "invoice_id", "po_line_id", "quantity", "unit_price", "extended", "sel_matched"])
    D = Cols(["approval_id", "invoice_id", "actor_id", "actor_role", "decision", "decided_epoch", "created_by", "imported_epoch", "total", "policy_row",
              "reviewer_role", "escalation_role", "amount_limit", "valid", "key"])
    P = Cols(["po_line_id", "vendor_id", "currency", "ordered_epoch", "ordered_quantity", "unit_price", "exists", "net_received", "consumed",
              "ordered_in", "open_to_receive", "open_to_invoice", "received_not_invoiced", "received_not_invoiced_value"])
    R = Cols(["receipt_id", "po_line_id", "received_quantity", "received_epoch", "reversed"])
    C = Cols(["credit_id", "invoice_id", "credit_amount", "credited_epoch"])
    Q = Cols(["seq", "batch_id", "idempotency_key", "exported_epoch", "invoices_named", "repeat"])
    X = Cols(["seq", "invoice_id", "exported_epoch", "repeat", "after_approval", "key"])
    E = Cols(["exception_id", "invoice_id", "exception_type", "created_epoch", "resolved_epoch", "open_sel", "owned_sel"])
    A = Cols(["exception_id", "owner_id", "assigned_epoch"])
    last = {"I": n_inv + 1, "L": n_lines + 1, "D": len(decisions) + 1, "P": len(pol_lines) + 1, "R": len(recs) + 1, "C": len(credits) + 1,
            "Q": len(reqs) + 1, "X": len(exp_lines) + 1, "E": len(exc) + 1, "A": len(asg) + 1}
    pol_a, pol_z = pol_rng
    CUT, POS = "Cutoff!$B$10", "Cutoff!$B$11"

    # --- Invoices
    ws = wb.create_sheet("Invoices")
    header(ws, 1, I.names)
    put(ws, inv_typed)
    li = last["I"]
    for r in range(2, li + 1):
        f = {n: f"{I[n]}{r}" for n in I.names}
        ws[f["raw_extended"]] = f"=SUMIFS({L.rng('Lines', 'extended', last['L'])},{L.rng('Lines', 'invoice_id', last['L'])},{f['invoice_id']})"
        ws[f["subtotal"]] = f"=ROUND({f['raw_extended']}+0.00000000001,2)"
        ws[f["tax"]] = f"=ROUND({f['raw_extended']}*{f['tax_rate']}+0.00000000001,2)"
        ws[f["total"]] = f"=ROUND({f['subtotal']}+{f['tax']}+{f['freight']}-{f['invoice_discount']},2)"
        ws[f["commit_epoch"]] = (f"=IFERROR(INDEX({D.rng('Decisions', 'decided_epoch', last['D'])},MATCH({f['invoice_id']}&\"|1\","
                                 f"{D.rng('Decisions', 'key', last['D'])},0)),{NEVER:.0f})")
        ws[f["export_epoch"]] = (f"=IFERROR(INDEX({X.rng('Export lines', 'exported_epoch', last['X'])},MATCH({f['invoice_id']}&\"|1\","
                                 f"{X.rng('Export lines', 'key', last['X'])},0)),{NEVER:.0f})")
        ws[f["sel_state"]] = f"=CHOOSE({POS},{f['state_march']},{f['state_june']},{f['state_asof']})"
        ws[f["sel_reasons"]] = f"=CHOOSE({POS},{f['reasons_march']},{f['reasons_june']},{f['reasons_asof']})"
        ws[f["sel_in_book"]] = f"=IF(OR({f['sel_state']}=\"matched\",{f['sel_state']}=\"exception\"),1,0)"
        ws[f["sel_approved"]] = (f"=IF(AND({f['sel_in_book']}=1,{f['commit_epoch']}<={CUT},OR({f['sel_state']}=\"matched\",{f['export_epoch']}<={CUT})),1,0)")
        ws[f["sel_matched"]] = f"=IF({f['sel_state']}=\"matched\",1,0)"
        ws[f["sel_good"]] = f"=IF(OR({f['sel_matched']}=1,{f['sel_approved']}=1),1,0)"
        ws[f["sel_exported"]] = f"=IF(AND({f['sel_in_book']}=1,{f['commit_epoch']}<={CUT},{f['export_epoch']}<={CUT}),1,0)"
        ws[f["sel_flagged"]] = f"=IF({f['sel_reasons']}<>\"\",1,0)"
        for code in ps.REASONS:
            ws[f[f"has_{code}"]] = f"=IF(ISNUMBER(SEARCH(\"{code}\",{f['sel_reasons']})),1,0)"
        ws[f["sel_credit"]] = (f"=IF({f['sel_approved']}=1,SUMIFS({C.rng('Credits', 'credit_amount', last['C'])},{C.rng('Credits', 'invoice_id', last['C'])},"
                               f"{f['invoice_id']},{C.rng('Credits', 'credited_epoch', last['C'])},\"<=\"&{CUT}),0)")
        ws[f["sel_payable"]] = f"=IF({f['sel_approved']}=1,ROUND({f['total']}-{f['sel_credit']},2),0)"
        ws[f["sel_minutes"]] = f"=IF({f['sel_approved']}=1,({f['commit_epoch']}-{f['imported_epoch']})/60,0)"
        ws[f["sel_ever"]] = (f"=IF(AND({f['sel_in_book']}=1,OR({f['sel_flagged']}=1,COUNTIFS({E.rng('Exceptions', 'invoice_id', last['E'])},{f['invoice_id']})>0)),1,0)")
        ws[f["sel_merged"]] = f"=IF({f['sel_state']}=\"merged\",1,0)"
    ws.freeze_panes = "B2"

    # --- Lines
    ws = wb.create_sheet("Lines")
    header(ws, 1, L.names)
    put(ws, [[l.invoice_line_id, l.invoice_id, l.po_line_id, float(l.quantity), float(l.unit_price)] for l in lines])
    for r in range(2, last["L"] + 1):
        ws[f"{L['extended']}{r}"] = f"={L['quantity']}{r}*{L['unit_price']}{r}"
        ws[f"{L['sel_matched']}{r}"] = (f"=INDEX({I.rng('Invoices', 'sel_matched', last['I'])},MATCH({L['invoice_id']}{r},{I.rng('Invoices', 'invoice_id', last['I'])},0))")
    ws.freeze_panes = "A2"

    # --- Decisions
    ws = wb.create_sheet("Decisions")
    header(ws, 1, D.names)
    put(ws, [[a.approval_id, a.invoice_id, a.actor_id, a.actor_role, a.decision, epoch(a.decided_at)] for a in decisions])
    inv_ids = I.rng("Invoices", "invoice_id", last["I"])
    for r in range(2, last["D"] + 1):
        g = lambda n: f"{D[n]}{r}"      # noqa: E731
        m = f"MATCH({g('invoice_id')},{inv_ids},0)"
        ws[g("created_by")] = f"=INDEX({I.rng('Invoices', 'created_by', last['I'])},{m})"
        ws[g("imported_epoch")] = f"=INDEX({I.rng('Invoices', 'imported_epoch', last['I'])},{m})"
        ws[g("total")] = f"=INDEX({I.rng('Invoices', 'total', last['I'])},{m})"
        ws[g("policy_row")] = f"=MATCH({g('decided_epoch')},Policies!$A${pol_a}:$A${pol_z},1)"
        ws[g("reviewer_role")] = f"=INDEX(Policies!$C${pol_a}:$C${pol_z},{g('policy_row')})"
        ws[g("escalation_role")] = f"=INDEX(Policies!$D${pol_a}:$D${pol_z},{g('policy_row')})"
        ws[g("amount_limit")] = f"=INDEX(Policies!$B${pol_a}:$B${pol_z},{g('policy_row')})"
        ws[g("valid")] = (f"=IF(AND({g('decision')}=\"approve\",{g('actor_id')}<>{g('created_by')},{g('decided_epoch')}>={g('imported_epoch')},"
                          f"OR(AND({g('escalation_role')}<>\"none\",{g('actor_role')}={g('escalation_role')}),"
                          f"AND({g('actor_role')}={g('reviewer_role')},{g('total')}<={g('amount_limit')}))),1,0)")
        ws[g("key")] = f"={g('invoice_id')}&\"|\"&{g('valid')}"
    ws.freeze_panes = "A2"

    # --- PO lines, Receipts
    ws = wb.create_sheet("PO lines")
    header(ws, 1, P.names)
    put(ws, [[l.po_line_id, po_ordered[l.po_id].vendor_id, po_ordered[l.po_id].currency, epoch(po_ordered[l.po_id].ordered_at), float(l.quantity), float(l.unit_price)]
             for l in pol_lines])
    for r in range(2, last["P"] + 1):
        g = lambda n: f"{P[n]}{r}"      # noqa: E731
        ws[g("exists")] = f"=IF({g('ordered_epoch')}<={CUT},1,0)"
        ws[g("net_received")] = (f"=IF({g('exists')}=1,SUMIFS({R.rng('Receipts', 'received_quantity', last['R'])},{R.rng('Receipts', 'po_line_id', last['R'])},{g('po_line_id')},"
                                 f"{R.rng('Receipts', 'reversed', last['R'])},0,{R.rng('Receipts', 'received_epoch', last['R'])},\"<=\"&{CUT}),0)")
        ws[g("consumed")] = (f"=IF({g('exists')}=1,SUMIFS({L.rng('Lines', 'quantity', last['L'])},{L.rng('Lines', 'po_line_id', last['L'])},{g('po_line_id')},"
                             f"{L.rng('Lines', 'sel_matched', last['L'])},1),0)")
        ws[g("ordered_in")] = f"={g('exists')}*{g('ordered_quantity')}"
        ws[g("open_to_receive")] = f"=IF({g('exists')}=1,MAX({g('ordered_quantity')}-{g('net_received')},0),0)"
        ws[g("open_to_invoice")] = f"=IF({g('exists')}=1,MAX({g('ordered_quantity')}-{g('consumed')},0),0)"
        ws[g("received_not_invoiced")] = f"=IF({g('exists')}=1,MAX({g('net_received')}-{g('consumed')},0),0)"
        ws[g("received_not_invoiced_value")] = f"={g('received_not_invoiced')}*{g('unit_price')}"
    ws.freeze_panes = "A2"
    ws = wb.create_sheet("Receipts")
    header(ws, 1, R.names)
    put(ws, [[r.receipt_id, r.po_line_id, float(r.received_quantity), epoch(r.received_at), 1 if r.reversed else 0] for r in recs])

    # --- Credits, Export requests, Export lines
    ws = wb.create_sheet("Credits")
    header(ws, 1, C.names)
    put(ws, [[c.credit_id, c.invoice_id, float(c.credit_amount), 0 if c.credited_at is None else epoch(c.credited_at)] for c in credits])
    ws = wb.create_sheet("Export requests")
    header(ws, 1, Q.names)
    put(ws, [[i, r.batch_id, r.idempotency_key, epoch(r.exported_at), len(r.invoice_ids)] for i, r in enumerate(reqs, start=1)])
    for r in range(2, last["Q"] + 1):
        ws[f"{Q['repeat']}{r}"] = (f"=IF(OR(COUNTIFS(${Q['idempotency_key']}$1:{Q['idempotency_key']}{r - 1},{Q['idempotency_key']}{r})>0,"
                                   f"COUNTIFS(${Q['batch_id']}$1:{Q['batch_id']}{r - 1},{Q['batch_id']}{r})>0),1,0)")
    ws = wb.create_sheet("Export lines")
    header(ws, 1, X.names)
    put(ws, [[s, i] for s, i in exp_lines])
    q_seq = Q.rng("Export requests", "seq", last["Q"])
    for r in range(2, last["X"] + 1):
        m = f"MATCH({X['seq']}{r},{q_seq},0)"
        ws[f"{X['exported_epoch']}{r}"] = f"=INDEX({Q.rng('Export requests', 'exported_epoch', last['Q'])},{m})"
        ws[f"{X['repeat']}{r}"] = f"=INDEX({Q.rng('Export requests', 'repeat', last['Q'])},{m})"
        ws[f"{X['after_approval']}{r}"] = (f"=IF({X['exported_epoch']}{r}>=INDEX({I.rng('Invoices', 'commit_epoch', last['I'])},MATCH({X['invoice_id']}{r},"
                                           f"{I.rng('Invoices', 'invoice_id', last['I'])},0)),1,0)")
        ws[f"{X['key']}{r}"] = f"={X['invoice_id']}{r}&\"|\"&IF(AND({X['repeat']}{r}=0,{X['after_approval']}{r}=1),1,0)"

    # --- Exceptions, Assignments
    ws = wb.create_sheet("Exceptions")
    header(ws, 1, E.names)
    put(ws, [[e["exception_id"], e["invoice_id"], e["exception_type"], epoch(e["created"]), NEVER if e["resolved"] is None else epoch(e["resolved"])] for e in
             [{**e, "created": parse_ts(e["created_at"]), "resolved": parse_ts(e["resolved_at"]) if e["resolved_at"] else None} for e in exc]])
    for r in range(2, last["E"] + 1):
        ws[f"{E['open_sel']}{r}"] = f"=IF(AND({E['created_epoch']}{r}<={CUT},{E['resolved_epoch']}{r}>{CUT}),1,0)"
        ws[f"{E['owned_sel']}{r}"] = (f"=IF(COUNTIFS({A.rng('Assignments', 'exception_id', last['A'])},{E['exception_id']}{r},"
                                      f"{A.rng('Assignments', 'assigned_epoch', last['A'])},\"<=\"&{CUT})>0,1,0)")
    ws = wb.create_sheet("Assignments")
    header(ws, 1, A.names)
    put(ws, [[a["exception_id"], a["owner_id"], epoch(parse_ts(a["assigned_at"]))] for a in asg])
    return {"I": I, "L": L, "D": D, "P": P, "R": R, "C": C, "Q": Q, "X": X, "E": E, "A": A, "last": last}


# metric id, label, unit, kind ('count' | 'money' | 'rate' | 'minutes' | 'qty'), scopes ('cur' per currency and ALL, 'all' ALL only), definition
METRICS = [
    ("eligible_invoices", "Invoices in the book", "invoices", "count", "cur", "validated, not merged, not held for review, not rejected, imported by the cutoff"),
    ("matched_invoices", "Matched or further", "invoices", "count", "cur", "the three-way match passes (or the invoice is already approved)"),
    ("matched_amount_before_tax", "Matched: amount before tax", "money", "money", "cur", "sum of the subtotals of the matched invoices, own currency"),
    ("matched_tax", "Matched: tax", "money", "money", "cur", "tax once per invoice on the unrounded subtotal, half up"),
    ("matched_total", "Matched: total", "money", "money", "cur", "subtotal + tax + freight - invoice discount"),
    ("obligations_count", "Obligations (approved)", "invoices", "count", "cur", "a policy-valid approval exists at the cutoff (approver is not the creator; role and limit by the policy in force at the decision)"),
    ("obligations_total", "Obligations: total", "money", "money", "cur", "own currency"),
    ("credit_total", "Credit notes applied", "money", "money", "cur", "credit notes dated by the cutoff against an obligation; once per credit id, no cash movement"),
    ("approved_payable", "Approved payable", "money", "money", "cur", "obligations total - credits"),
    ("exported_invoices", "Exported invoices", "invoices", "count", "cur", "taken by a non-repeat export request at or after a valid approval"),
    ("accepted_export_total", "Accepted export total", "money", "money", "cur", "the computed totals of the export rows"),
    ("export_difference", "Export - approved difference", "money", "money", "cur", "0.00 by construction of the generator's closing batch, never a finding"),
    ("open_exceptions", "Open exceptions (derived)", "exceptions", "count", "cur", "one per (invoice, reason) the match still raises at the cutoff"),
    ("open_QUANTITY_BLOCKED", "  of which QUANTITY_BLOCKED", "exceptions", "count", "cur", ""),
    ("open_PRICE_VARIANCE", "  of which PRICE_VARIANCE", "exceptions", "count", "cur", ""),
    ("open_NO_PO_LINE", "  of which NO_PO_LINE", "exceptions", "count", "cur", ""),
    ("open_ROUNDING_EXCEPTION", "  of which ROUNDING_EXCEPTION", "exceptions", "count", "cur", ""),
    ("open_CURRENCY_MISMATCH", "  of which CURRENCY_MISMATCH", "exceptions", "count", "cur", ""),
    ("straight_through_rate", "Straight-through rate", "ratio", "rate", "cur", "invoices in the book NOT blocked at the cutoff / invoices in the book (D-09.10)"),
    ("exception_rate", "Exception rate", "ratio", "rate", "cur", "blocked at the cutoff / invoices in the book"),
    ("ever_excepted_rate", "Ever excepted rate", "ratio", "rate", "cur", "blocked now or ever (the stored history) / invoices in the book; published only where the history adds invoices"),
    ("never_excepted_rate", "Never excepted rate", "ratio", "rate", "cur", "1 - ever excepted rate"),
    ("mean_minutes_to_approval", "Mean minutes to approval", "minutes", "minutes", "cur", "approval instant - import instant, approved invoices"),
    ("pending_invoices", "Pending (not approved)", "invoices", "count", "cur", "invoices in the book without a valid approval"),
    ("duplicates_merged", "Duplicates merged", "invoices", "count", "all", "invoices merged into an earlier canonical copy at the cutoff (all currencies)"),
    ("ordered_quantity", "Ordered quantity", "units", "qty", "cur", "PO lines ordered by the cutoff"),
    ("net_received_quantity", "Net received quantity", "units", "qty", "cur", "receipts not reversed and received by the cutoff"),
    ("accepted_invoiced_quantity", "Accepted invoiced quantity", "units", "qty", "cur", "quantities on the lines of matched invoices (cumulative consumption)"),
    ("open_to_receive_quantity", "Open to receive", "units", "qty", "cur", "sum of max(ordered - net received, 0)"),
    ("open_to_invoice_quantity", "Open to invoice", "units", "qty", "cur", "sum of max(ordered - accepted invoiced, 0)"),
    ("received_not_invoiced_quantity", "Received not invoiced", "units", "qty", "cur", "sum of max(net received - accepted invoiced, 0)"),
    ("received_not_invoiced_value", "Received not invoiced: value at PO price", "money", "money", "cur", "own currency of the PO line"),
    ("open_exceptions_history", "Open exceptions (stored history)", "exceptions", "count", "all", "created by the cutoff and not yet resolved, from the nightly run's history (differs from the derived count by 1 at the March month-end)"),
    ("unowned_open_exceptions_history", "Unowned open exceptions (stored history)", "exceptions", "count", "all", "open and no assignment at or before the cutoff"),
]
SCOPES = ("USD", "EUR", "GBP", "ALL")


def metric_formula(mid: str, scope: str, cols: dict, cell: dict) -> str | None:
    """The live formula of one Metrics cell. ``cell`` maps metric id -> the cell address of the same scope."""
    I, P, E, last = cols["I"], cols["P"], cols["E"], cols["last"]
    cur = f"{I.rng('Invoices', 'currency', last['I'])},\"{scope}\"," if scope != "ALL" else ""
    pcur = f"{P.rng('PO lines', 'currency', last['P'])},\"{scope}\"," if scope != "ALL" else ""

    def flag_sum(flag: str) -> str:
        if scope == "ALL":
            return f"SUM({I.rng('Invoices', flag, last['I'])})"
        return f"SUMIFS({I.rng('Invoices', flag, last['I'])},{cur[:-1]})"

    def money(col: str, flag: str) -> str:
        if scope == "ALL":
            return ""
        return f"ROUND(SUMIFS({I.rng('Invoices', col, last['I'])},{cur}{I.rng('Invoices', flag, last['I'])},1),2)"
    if mid == "eligible_invoices":
        return "=" + flag_sum("sel_in_book")
    if mid == "matched_invoices":
        return "=" + flag_sum("sel_good")
    if mid == "obligations_count":
        return "=" + flag_sum("sel_approved")
    if mid == "exported_invoices":
        return "=" + flag_sum("sel_exported")
    if mid in ("matched_amount_before_tax", "matched_tax", "matched_total"):
        return "=" + money({"matched_amount_before_tax": "subtotal", "matched_tax": "tax", "matched_total": "total"}[mid], "sel_good") if scope != "ALL" else None
    if mid == "obligations_total":
        return "=" + money("total", "sel_approved") if scope != "ALL" else None
    if mid == "credit_total":
        return f"=ROUND(SUMIFS({I.rng('Invoices', 'sel_credit', last['I'])},{cur[:-1]}),2)" if scope != "ALL" else None
    if mid == "approved_payable":
        return f"=ROUND({cell['obligations_total']}-{cell['credit_total']},2)" if scope != "ALL" else None
    if mid == "accepted_export_total":
        return "=" + money("total", "sel_exported") if scope != "ALL" else None
    if mid == "export_difference":
        return f"=ROUND({cell['accepted_export_total']}-{cell['obligations_total']},2)" if scope != "ALL" else None
    if mid == "open_exceptions":
        return "=" + "+".join(flag_sum(f"has_{c}") for c in ps.REASONS)
    if mid.startswith("open_") and mid[5:] in ps.REASONS:
        return "=" + flag_sum(f"has_{mid[5:]}")
    if mid == "straight_through_rate":
        return f"=IF({cell['eligible_invoices']}=0,\"\",ROUND(({cell['eligible_invoices']}-{flag_sum('sel_flagged')})/{cell['eligible_invoices']},6))"
    if mid == "exception_rate":
        return f"=IF({cell['eligible_invoices']}=0,\"\",ROUND({flag_sum('sel_flagged')}/{cell['eligible_invoices']},6))"
    if mid == "ever_excepted_rate":
        return f"=IF({cell['eligible_invoices']}=0,\"\",ROUND({flag_sum('sel_ever')}/{cell['eligible_invoices']},6))"
    if mid == "never_excepted_rate":
        return f"=IF({cell['eligible_invoices']}=0,\"\",ROUND(({cell['eligible_invoices']}-{flag_sum('sel_ever')})/{cell['eligible_invoices']},6))"
    if mid == "mean_minutes_to_approval":
        return f"=IF({cell['obligations_count']}=0,\"\",ROUND({flag_sum('sel_minutes')}/{cell['obligations_count']},4))"
    if mid == "pending_invoices":
        return f"={cell['eligible_invoices']}-{cell['obligations_count']}"
    if mid == "duplicates_merged":
        return "=" + flag_sum("sel_merged") if scope == "ALL" else None
    qty = {"ordered_quantity": "ordered_in", "net_received_quantity": "net_received", "accepted_invoiced_quantity": "consumed",
           "open_to_receive_quantity": "open_to_receive", "open_to_invoice_quantity": "open_to_invoice",
           "received_not_invoiced_quantity": "received_not_invoiced"}
    if mid in qty:
        rng = P.rng("PO lines", qty[mid], last["P"])
        return f"=ROUND(SUM({rng}),4)" if scope == "ALL" else f"=ROUND(SUMIFS({rng},{pcur[:-1]}),4)"
    if mid == "received_not_invoiced_value":
        return f"=ROUND(SUMIFS({P.rng('PO lines', 'received_not_invoiced_value', last['P'])},{pcur[:-1]}),2)" if scope != "ALL" else None
    if mid == "open_exceptions_history":
        return f"=SUM({E.rng('Exceptions', 'open_sel', last['E'])})" if scope == "ALL" else None
    if mid == "unowned_open_exceptions_history":
        return f"=COUNTIFS({E.rng('Exceptions', 'open_sel', last['E'])},1,{E.rng('Exceptions', 'owned_sel', last['E'])},0)" if scope == "ALL" else None
    raise KeyError(mid)


def build_domain_reference(wb, ref):
    ws = wb.create_sheet("Domain reference")
    title(ws, "PASTED VALUES from the Python domain (compute_metrics and window_metrics), not formulas",
          "One column per cutoff and scope: the domain's figure over web/public/data, exact decimals. Metrics reads these only for its check column. Blank = the domain publishes no such figure for that scope (money is never summed across currencies).")
    labels = [f"{c}|{s}" for c, _ in CUTS for s in SCOPES]
    header(ws, 4, ["metric"] + labels)
    rows = {}
    for r, (mid, *_rest) in enumerate(METRICS, start=5):
        rows[mid] = r
        ws.cell(row=r, column=1, value=mid)
        for c, (cname, _) in enumerate(CUTS):
            for s, scope in enumerate(SCOPES):
                v = ref[(cname, scope)].get(mid)
                if v is not None:
                    ws.cell(row=r, column=2 + c * 4 + s, value=v)
    widths(ws, {"A": 38})
    return rows, len(labels)


def build_metrics(wb, cols, dref_rows, n_ref_cols):
    ws = wb.create_sheet("Metrics", 1)
    title(ws, "Headline metrics for the selected cutoff, per currency (live formulas)",
          "Cutoff: Cutoff!B3. Money is per currency and never summed across currencies; the ALL column holds counts and units only. 'check' is TRUE when every figure of the row equals the Python domain's value for the same cutoff and scope within the tolerance.")
    ws["A3"], ws["B3"] = "selected cutoff", "=Cutoff!B3"
    ws["D3"], ws["E3"] = "all checks TRUE", None
    header(ws, 5, ["metric id", "metric", "unit", "USD", "EUR", "GBP", "ALL (counts and units)", "definition", "tolerance", "check"])
    rows = {mid: 6 + i for i, (mid, *_r) in enumerate(METRICS)}
    last_row = 5 + len(METRICS)
    letters = {"USD": "D", "EUR": "E", "GBP": "F", "ALL": "G"}
    ref_a, ref_z = "B", get_column_letter(1 + n_ref_cols)
    for mid, label, unit, kind, scope_kind, definition in METRICS:
        r = rows[mid]
        ws.cell(row=r, column=1, value=mid)
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=8, value=definition).alignment = Alignment(wrap_text=True, vertical="top")
        tol = {"count": 0, "qty": 0.00005, "money": 0.005, "rate": 0.0000005, "minutes": 0.00005}[kind]
        ws.cell(row=r, column=9, value=tol).font = BLUE
        checks = []
        for scope in SCOPES:
            cell = {m: f"{letters[scope]}{rows[m]}" for m in rows}
            f = metric_formula(mid, scope, cols, cell)
            if scope != "ALL" and scope_kind == "all":
                f = None
            if f is None:
                continue
            ws[f"{letters[scope]}{r}"] = f
            ws[f"{letters[scope]}{r}"].number_format = {"count": "0", "qty": "0.0000", "money": "#,##0.00", "rate": "0.000000", "minutes": "0.0000"}[kind]
            col = f"MATCH(Cutoff!$B$3&\"|{scope}\",'Domain reference'!$B$4:${ref_z}$4,0)"
            ref = f"INDEX('Domain reference'!$B${dref_rows[mid]}:${ref_z}${dref_rows[mid]},{col})"
            cellref = f"{letters[scope]}{r}"
            checks.append(f"IF(ISNUMBER({ref}),IF(ISNUMBER({cellref}),ABS({cellref}-{ref})<=$I{r}+0.0000000001,FALSE),IF(ISNUMBER({cellref}),FALSE,TRUE))")
        ws.cell(row=r, column=10, value="=AND(" + ",".join(checks) + ")")
    ws["E3"] = f"=AND(J6:J{last_row})"
    ws["E3"].font = BOLD
    widths(ws, {"A": 32, "B": 40, "C": 12, "D": 16, "E": 16, "F": 16, "G": 20, "H": 80, "I": 10, "J": 8})
    ws.freeze_panes = "D6"
    return rows


def build_golden(wb):
    ws = wb.create_sheet("Golden case")
    title(ws, "The known answer, with formulas (P05-GOLD-01)",
          "PO 10 x 20.00; receipt 6; invoice of 6 -> 120.00 before tax, 12.00 tax, 132.00; a second invoice of 4 is QUANTITY_BLOCKED; the first imported again leaves one obligation. Blue = typed inputs, black = formulas, last column = expected.json.")
    inp = [("PO line quantity", 10), ("PO unit price", 20.00), ("net received quantity", 6), ("tax rate", 0.10), ("freight", 0), ("invoice discount", 0),
           ("invoice 1 quantity", 6), ("invoice 2 quantity", 4), ("approver is not the creator (1 = yes)", 1), ("policy limit", 5000)]
    header(ws, 3, ["input", "value"])
    for i, (k, v) in enumerate(inp, start=4):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v).font = BLUE
    # inputs: B4 qty, B5 price, B6 net, B7 rate, B8 freight, B9 discount, B10 q1, B11 q2, B12 approver ok, B13 limit
    header(ws, 16, ["figure", "formula", "expected.json", "check"])
    docs = EXPECTED["documents"]["invoices"]
    mq = EXPECTED["metrics"]["quantities"]
    mres = {m["invoice_line_id"]: m for m in EXPECTED["documents"]["match_results"]}
    rows = [
        ("invoice 1 subtotal", "=ROUND(B10*B5,2)", float(docs["inv-001"]["subtotal"])),
        ("invoice 1 tax", "=ROUND(B10*B5*B7,2)", float(docs["inv-001"]["tax"])),
        ("invoice 1 total", "=ROUND(B17+B18+B8-B9,2)", float(docs["inv-001"]["total"])),
        ("available to invoice 1 (net received - consumed 0)", "=B6-0", float(mres["il-001"]["available_quantity"])),
        ("invoice 1 quantity variance (invoiced - available)", "=B10-B20", float(mres["il-001"]["quantity_variance"])),
        ("invoice 1 outcome (variance <= 0 means matched)", '=IF(B21<=0,"matched","QUANTITY_BLOCKED")', "matched"),
        ("quantity consumed by invoice 1", "=IF(B22=\"matched\",B10,0)", 6.0),
        ("invoice 2 subtotal", "=ROUND(B11*B5,2)", float(docs["inv-002"]["subtotal"])),
        ("invoice 2 total", "=ROUND(B24+ROUND(B11*B5*B7,2)+B8-B9,2)", float(docs["inv-002"]["total"])),
        ("available to invoice 2 (net received - consumed)", "=B6-B23", float(mres["il-002"]["available_quantity"])),
        ("invoice 2 quantity variance", "=B11-B26", float(mres["il-002"]["quantity_variance"])),
        ("invoice 2 outcome", '=IF(B27<=0,"matched","QUANTITY_BLOCKED")', "QUANTITY_BLOCKED"),
        ("duplicate: invoice 3 total equals invoice 1 total (same identity) -> merged (1 = merged)", "=IF(ROUND(B10*B5,2)+ROUND(B10*B5*B7,2)+B8-B9=B19,1,0)", 1.0),
        ("obligations (invoice 1 approved: approver differs, total within the limit; invoice 3 merged; invoice 2 blocked)",
         '=IF(AND(B22="matched",B12=1,B19<=B13),1,0)', float(EXPECTED["metrics"]["obligations"]["count"])),
        ("obligations total", '=IF(B30=1,B19,0)', float(EXPECTED["metrics"]["obligations"]["total"])),
        ("ordered quantity", "=B4", float(mq["ordered_quantity"])),
        ("open to receive (ordered - received)", "=B4-B6", float(mq["open_to_receive_quantity"])),
        ("received not invoiced (net received - consumed)", "=B6-B23", float(mq["received_not_invoiced_quantity"])),
    ]
    for i, (label, f, want) in enumerate(rows, start=17):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=f)
        c = ws.cell(row=i, column=3, value=want)
        c.font = BLUE
        ws.cell(row=i, column=4, value=f'=IF(ISNUMBER(B{i}),ABS(B{i}-C{i})<=0.000001,B{i}=C{i})')
    ws[f"A{17 + len(rows) + 1}"] = "all golden checks TRUE"
    ws[f"B{17 + len(rows) + 1}"] = f"=AND(D17:D{16 + len(rows)})"
    ws[f"B{17 + len(rows) + 1}"].font = BOLD
    widths(ws, {"A": 86, "B": 18, "C": 16, "D": 10})
    return 17 + len(rows) + 1


ASSUMPTIONS = [
    ("D-01", "Identity is vendor, normalised invoice number, invoice date and the COMPUTED total; the same source bytes are checked first. A merged invoice consumes nothing and keeps its lineage."),
    ("D-02", "All money is decimal; tax is computed once per invoice on the unrounded subtotal, subtotal and tax rounded half up to 2 dp; total = subtotal + tax + freight - invoice discount. Rounding tolerance 0.01 per line and 0.02 per invoice."),
    ("D-03", "Three-way match: a line matches when the invoiced quantity does not exceed net received minus the quantity already consumed by earlier accepted invoices and the price is within the tolerance (0.00 per unit). Only a matched invoice consumes quantity."),
    ("D-04", "Approval needs an actor other than the creator and a role allowed by the policy version effective at decision time for the invoice total; the first valid approval counts. Export is an idempotent batch (a repeated key or batch id returns the prior batch); a row needs a valid approval at or before its request. Credit notes apply once and reduce the payable; no cash moves."),
    ("D-05", "Tolerances travel with the data; an obligation is an approved, exported or adjusted invoice; rates have the invoices in the book as denominator."),
    ("D-07", "Synthetic world, seed 20261026, 12 months, as-of 2026-10-01T02:00:00Z; the generator approves only uninterrupted matches, rejects only never-matched invoices and closes each period with a batch of everything approved."),
    ("D-09", "Committed outcomes (an invoice with a policy-valid approval) consume received quantity first and are never displaced; the straight-through rate keeps its meaning (not blocked at the cutoff) and the never-excepted rate is added beside it; money is per currency."),
    ("W-01", "Money is never summed across currencies: Metrics has one column per currency and the ALL column carries counts and units only."),
    ("W-02", "The three-way match outcome of each invoice at each cutoff (matched, exception with its reasons, merged, held for review, rejected) is TYPED from the Python domain, not recomputed by formulas: the cumulative consumption is an ordered chain that a spreadsheet cannot reproduce row by row for three cutoffs. data/sql/analysis recomputes the same book from facts in SQL and reconciles it."),
    ("W-03", "Floating point: line extended amounts are double products; subtotal and tax add 1E-11 before ROUND so that exact half-cent ties round up as the domain's ROUND_HALF_UP does. tests/test_artifacts.py compares the total of every one of the 2,800 invoices with the domain after recalculation."),
    ("W-04", "Recalculation: openpyxl writes formulas only. The delivered file was recalculated by LibreOffice headless and carries its cached values; Excel was NOT run (docs/workbook.md). fullCalcOnLoad is set so Excel recalculates on open. Only functions common to both engines are used (SUMIFS, COUNTIFS, INDEX, MATCH, IFERROR, ROUND, SEARCH, CHOOSE)."),
    ("W-05", "The stored exception history is the nightly run's (02:00 UTC); the derived count can differ from it for an invoice imported between the last run and the cutoff: 1 at the March month-end, 0 at the other two."),
    ("W-06", "A policy limit is compared with the invoice total in the invoice's own currency with no conversion (the contract has no dated rate)."),
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Decisions applied (docs/contracts.md) and the workbook's own limits")
    header(ws, 3, ["id", "statement"])
    for i, (k, v) in enumerate(ASSUMPTIONS, start=4):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 8, "B": 150})


def build_sql_results(wb):
    ws = wb.create_sheet("SQL results (pasted)")
    ws["A1"] = "PASTED VALUES, not formulas, from data/sql/analysis/run_all.py"
    ws["A1"].font = Font(bold=True, size=14)
    if not RESULTS_PATH.exists():
        ws["A2"] = "artifacts/analysis-results.json was not found: run data/sql/analysis/run_all.py --json artifacts/analysis-results.json first"
        return
    res = json.loads(RESULTS_PATH.read_text())
    meta = res["meta"][0]
    ws["A2"] = f"dataset hash {meta['dataset_hash']} | code revision {meta['code_revision']} | contract v{meta['contract_version']} | goes stale if the data or the queries change"
    ws["A3"], ws["B3"], ws["C3"] = "reconciliation rows | equal", meta["reconciliation_rows"], meta["reconciliation_equal"]
    r = 5
    for key in ("01_match_outcomes_by_vendor_and_reason__reason_open_vs_ever", "01_match_outcomes_by_vendor_and_reason__vendor_group_open_vs_ever",
                "01_match_outcomes_by_vendor_and_reason__by_currency", "06_export_batches_and_reconciliation__reconciliation_by_currency",
                "07_credits_and_payable_by_currency__payable_by_currency", "04_approvals_under_policy_versions__approvals_by_policy_and_currency"):
        rows = res[key]
        ws.cell(row=r, column=1, value=key).font = BOLD
        header(ws, r + 1, list(rows[0]))
        for i, row in enumerate(rows, start=r + 2):
            for j, (k, v) in enumerate(row.items(), start=1):
                try:
                    v = float(v) if v not in ("",) and k not in ("currency", "reason", "vendor_group", "policy_id") and not any(ch.isalpha() for ch in v.replace("e", "")) else v
                except (ValueError, AttributeError):
                    pass
                ws.cell(row=i, column=j, value=v)
        r += len(rows) + 4
    widths(ws, {"A": 56})


def recalc(path: Path) -> None:
    """Open with LibreOffice headless and save a recalculated copy over ``path`` (cached values included)."""
    if shutil.which("soffice") is None:
        raise SystemExit("soffice not installed; the file has formulas but no cached values")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in" / path.name
        src.parent.mkdir()
        shutil.copy(path, src)
        out = Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=1800)
        shutil.copy(out / path.name, path)
    set_full_calc_on_load(path)


def set_full_calc_on_load(path: Path) -> None:
    """LibreOffice drops fullCalcOnLoad; put it back (workbook.xml only) so Excel recalculates when it opens the file."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode("utf-8")
                assert "<calcPr " in text and "fullCalcOnLoad" not in text
                data = text.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ', 1).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


SHEET_ORDER = ["Read me", "Cutoff", "Metrics", "Domain reference", "Golden case", "Assumptions", "Invoices", "Lines", "Decisions", "PO lines", "Receipts", "Credits",
               "Export requests", "Export lines", "Exceptions", "Assignments", "Policies", "SQL results (pasted)"]


def build(path: Path) -> None:
    ctx = ps.load()
    ref = domain_reference(ctx)
    wb = Workbook()
    build_readme(wb)
    build_cutoff(wb)
    pol_rng = build_policies(wb, ctx)
    cols = build_data(wb, ctx, pol_rng)
    dref_rows, n_ref = build_domain_reference(wb, ref)
    build_metrics(wb, cols, dref_rows, n_ref)
    build_golden(wb)
    build_assumptions(wb)
    build_sql_results(wb)
    wb._sheets = [wb[n] for n in SHEET_ORDER]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main(argv: list[str]) -> int:
    build(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    if "--recalc" in argv:
        recalc(OUT)
        print("recalculated with LibreOffice headless; cached values written")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
