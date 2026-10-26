"""Build artifacts/onboard-studio-analyst-workbook.xlsx with openpyxl (P02-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: tests/fixtures/golden/inputs.json and expected.json (golden fixture), artifacts/analysis-results.json (written by
data/sql/analysis/run_all.py), web/public/data/manifest.json. openpyxl cannot calculate, so the file it writes holds formulas
only. With --recalc, LibreOffice (soffice) opens it headless and saves a recalculated copy over the same path, which then also
carries cached values. Overlap and Figures never hard-code a computed quantity: every duration is a formula over Inputs.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "onboard-studio-analyst-workbook.xlsx"
GOLD = json.loads((ROOT / "tests/fixtures/golden/inputs.json").read_text())
EXPECTED = json.loads((ROOT / "tests/fixtures/golden/expected.json").read_text())
RESULTS = json.loads((ROOT / "artifacts/analysis-results.json").read_text())
MANIFEST = json.loads((ROOT / "web/public/data/manifest.json").read_text())

HDR = PatternFill("solid", fgColor="1F3A5F")
SEC = PatternFill("solid", fgColor="DCE6F1")
OKFILL = PatternFill("solid", fgColor="E2F0D9")
BLUE = Font(color="0000FF")            # typed inputs
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
TS = "yyyy-mm-dd hh:mm"
H4 = "0.0000"

ENTITIES = ["organizations", "workflow_versions", "step_definitions", "dependencies", "sla_rules", "calendar_intervals",
            "cases", "step_instances", "assignments", "document_versions", "approvals", "waiting_intervals", "case_events"]
TS_FIELDS = {"created_at", "published_at", "starts_at", "ends_at", "opened_at", "completed_at", "ready_at", "started_at",
             "approved_at", "assigned_at", "ended_at", "submitted_at", "reviewed_at", "decided_at", "occurred_at"}
POS: dict[str, dict] = {}      # entity -> {"first": row, "last": row, "col": {field: letter}}


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def parse(v):
    """Fixture value -> cell value: ISO instants become naive-UTC datetimes, lists become comma text."""
    if isinstance(v, str) and len(v) == 20 and v.endswith("Z") and v[4] == "-" and "T" in v:
        return datetime.strptime(v, "%Y-%m-%dT%H:%M:%SZ")
    if isinstance(v, list):
        return ", ".join(v)
    if isinstance(v, bool):
        return str(v).lower()          # LibreOffice rewrites real booleans as =TRUE(); keep the fixture's literal as text
    return v


def build_inputs(wb):
    ws = wb.active
    ws.title = "Inputs"
    title(ws, "Inputs: the golden fixture rows (tests/fixtures/golden/inputs.json, fixture onboard-studio-golden-v1)",
          "SYNTHETIC. Every cell below is a typed input copied from the fixture (blue). Instants are UTC, shown as yyyy-mm-dd hh:mm. Contract v1, docs/contracts.md.")
    r = 4
    for ent in ENTITIES:
        rows = GOLD[ent]
        cols = list(rows[0])
        ws.cell(row=r, column=1, value=ent).font = BOLD
        header(ws, r + 1, cols)
        POS[ent] = {"first": r + 2, "last": r + 1 + len(rows), "col": {c: L(i + 1) for i, c in enumerate(cols)}}
        for i, row in enumerate(rows):
            for j, c in enumerate(cols):
                cell = ws.cell(row=r + 2 + i, column=j + 1, value=parse(row[c]))
                cell.font = BLUE
                if c in TS_FIELDS and cell.value is not None:
                    cell.number_format = TS
        r += len(rows) + 3
    widths(ws, {"A": 22, "B": 22, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18})


def ref(ent: str, field: str, idx: int = 0) -> str:
    p = POS[ent]
    return f"Inputs!${p['col'][field]}${p['first'] + idx}"


def build_calendar(wb):
    ws = wb.create_sheet("Calendar")
    title(ws, "Calendar: working intervals of cal-001 (half-open [starts_at, ends_at); contract D-01)",
          "Live links to Inputs (calendar_intervals). `kind` is typed: the fixture rows carry none and D-06 defaults it to 'working'. No holidays in the fixture.")
    header(ws, 4, ["calendar_interval_id", "calendar_id", "starts_at (UTC)", "ends_at (UTC)", "kind", "hours (formula)"])
    n = len(GOLD["calendar_intervals"])
    for i in range(n):
        r = 5 + i
        ws.cell(r, 1, f"={ref('calendar_intervals', 'calendar_interval_id', i)}")
        ws.cell(r, 2, f"={ref('calendar_intervals', 'calendar_id', i)}")
        ws.cell(r, 3, f"={ref('calendar_intervals', 'starts_at', i)}").number_format = TS
        ws.cell(r, 4, f"={ref('calendar_intervals', 'ends_at', i)}").number_format = TS
        k = ws.cell(r, 5, "working")
        k.font = BLUE
        ws.cell(r, 6, f"=ROUND((D{r}-C{r})*24,6)").number_format = H4
    ws.cell(5 + n, 5, "working hours in the calendar").font = BOLD
    ws.cell(5 + n, 6, f'=ROUND(SUMIF(E5:E{4 + n},"working",F5:F{4 + n}),4)').number_format = H4
    widths(ws, {"A": 22, "B": 12, "C": 18, "D": 18, "E": 30, "F": 16})
    return n


def build_overlap(wb, n):
    ws = wb.create_sheet("Overlap")
    title(ws, "Overlap: the interval-overlap arithmetic as live formulas (docs/contracts.md D-01; golden P02-GOLD-01)",
          "overlap(range, working interval) = MAX(0, MIN(range_end, interval_end) - MAX(range_start, interval_start)) in hours. Intervals that are not 'working' contribute 0.")
    # Block 1: the case span
    ws["A4"] = "Case span"
    ws["A4"].font = BOLD
    ws["A5"], ws["B5"] = "case_id", f"={ref('cases', 'case_id')}"
    ws["A6"], ws["B6"] = "opened_at", f"={ref('cases', 'opened_at')}"
    ws["A7"], ws["B7"] = "completed_at", f"={ref('cases', 'completed_at')}"
    for c in ("B6", "B7"):
        ws[c].number_format = TS
    # Block 2: the waiting interval, clipped to the case span (D-06)
    ws["D4"] = "Waiting interval clipped to the case span (D-06)"
    ws["D4"].font = BOLD
    ws["D5"], ws["E5"] = "waiting_interval_id", f"={ref('waiting_intervals', 'waiting_interval_id')}"
    ws["D6"], ws["E6"] = "reason_code", f"={ref('waiting_intervals', 'reason_code')}"
    ws["D7"], ws["E7"] = "clipped start", f"=MAX({ref('waiting_intervals', 'starts_at')},B6)"
    ws["D8"], ws["E8"] = "clipped end", f"=MIN({ref('waiting_intervals', 'ends_at')},B7)"
    for c in ("E7", "E8"):
        ws[c].number_format = TS
    header(ws, 10, ["calendar_interval_id", "kind", "starts_at (UTC)", "ends_at (UTC)",
                    "case overlap (h)", "waiting overlap (h)"])
    first = 11
    for i in range(n):
        r = first + i
        cr = 5 + i                                   # the Calendar row
        ws.cell(r, 1, f"=Calendar!A{cr}")
        ws.cell(r, 2, f"=Calendar!E{cr}")
        ws.cell(r, 3, f"=Calendar!C{cr}").number_format = TS
        ws.cell(r, 4, f"=Calendar!D{cr}").number_format = TS
        ws.cell(r, 5, f'=IF(B{r}="working",ROUND(MAX(0,MIN($B$7,D{r})-MAX($B$6,C{r}))*24,6),0)').number_format = H4
        ws.cell(r, 6, f'=IF(B{r}="working",ROUND(MAX(0,MIN($E$8,D{r})-MAX($E$7,C{r}))*24,6),0)').number_format = H4
    last = first + n - 1
    t = last + 1
    ws.cell(t, 1, "TOTAL business hours").font = BOLD
    ws.cell(t, 5, f"=ROUND(SUM(E{first}:E{last}),4)").number_format = H4
    ws.cell(t, 6, f"=ROUND(SUM(F{first}:F{last}),4)").number_format = H4
    ws.cell(t + 1, 1, "working intervals touched").font = BOLD
    ws.cell(t + 1, 5, f'=COUNTIF(E{first}:E{last},">0")')
    ws.cell(t + 1, 6, f'=COUNTIF(F{first}:F{last},">0")')
    for c in (ws.cell(t, 5), ws.cell(t, 6)):
        c.font, c.fill = BOLD, OKFILL
    ws.cell(t + 3, 1, "Reading: E = business cycle time (expected 2.0000: Fri 16:00-17:00 + Mon 09:00-10:00). "
                      "F = business overlap of the waiting interval (expected 1.0000: Fri 16:30-17:00 + Mon 09:00-09:30).")
    widths(ws, {"A": 26, "B": 22, "C": 18, "D": 22, "E": 20, "F": 20})
    return {"total_row": t, "case_cell": f"Overlap!$E${t}", "wait_cell": f"Overlap!$F${t}",
            "wait_start": "Overlap!$E$7", "wait_end": "Overlap!$E$8"}


def build_figures(wb, ov):
    ws = wb.create_sheet("Figures")
    title(ws, "Figures: the golden case's durations from the Overlap formulas (contract D-01, D-04, D-05)",
          "Change B3 to see the effect of the pause policy: 1 = the waiting reason is pause-eligible, 0 = it is not. Columns C and D evaluate both states at once.")
    ws["A3"], ws["B3"] = "Pause-eligible switch (1 / 0)", 1
    ws["B3"].font = BLUE
    ws["A4"] = "Policy check (formula)"
    ws["B4"] = (f'=IF(ISNUMBER(SEARCH({ref("waiting_intervals", "reason_code")},'
                f'{ref("sla_rules", "pause_policy")})),1,0)')
    ws["C4"] = "<- 1 when the SLA rule's pause_policy in Inputs lists the waiting reason"
    header(ws, 6, ["Quantity (hours unless stated)", "Selected state (switch B3)", "State A: pause-eligible",
                   "State B: not eligible", "Expected A (expected.json)", "Expected B (expected.json)",
                   "A matches", "B matches"])
    ws["B7"], ws["C7"], ws["D7"] = "=$B$3", 1, 0
    ws["A7"] = "switch used in the column"
    ws["C7"].font = ws["D7"].font = BLUE
    opened, done = ref("cases", "opened_at"), ref("cases", "completed_at")
    rows = [
        ("Calendar cycle time", lambda sw: f"=ROUND(({done}-{opened})*24,4)", "calendar_cycle_time_hours"),
        ("Business cycle time", lambda sw: f"={ov['case_cell']}", "business_cycle_time_hours"),
        ("Paused time (calendar)", lambda sw: f"=IF({sw}=1,ROUND(MAX(0,{ov['wait_end']}-{ov['wait_start']})*24,4),0)", "paused_time_hours"),
        ("Paused business overlap", lambda sw: f"=IF({sw}=1,{ov['wait_cell']},0)", "paused_business_overlap_hours"),
        ("Active elapsed time", None, "active_elapsed_hours"),
    ]
    r0 = 8
    for i, (label, fn, key) in enumerate(rows):
        r = r0 + i
        ws.cell(r, 1, label)
        for col, sw in (("B", "B$7"), ("C", "C$7"), ("D", "D$7")):
            if fn is None:
                f = f"=ROUND(MAX(0,{col}{r0 + 1}-{col}{r0 + 3}),4)"
            else:
                f = fn(sw)
            ws[f"{col}{r}"] = f
            ws[f"{col}{r}"].number_format = H4
        ea, eb = EXPECTED["states"]["pause_eligible"]["metrics"][key], EXPECTED["states"]["pause_not_eligible"]["metrics"][key]
        ws.cell(r, 5, float(ea)).font = BLUE
        ws.cell(r, 6, float(eb)).font = BLUE
        ws.cell(r, 5).number_format = ws.cell(r, 6).number_format = H4
        ws.cell(r, 7, f"=ABS(C{r}-E{r})<0.00005")
        ws.cell(r, 8, f"=ABS(D{r}-F{r})<0.00005")
    r = r0 + len(rows)                              # 13: SLA
    ws.cell(r, 1, "SLA business minutes (sla_rules)")
    for col in "BCD":
        ws[f"{col}{r}"] = f"={ref('sla_rules', 'business_minutes')}"
    ws.cell(r + 1, 1, "Active minutes (active hours x 60)")
    for col in "BCD":
        ws[f"{col}{r + 1}"] = f"=ROUND({col}{r0 + 4}*60,2)"
    ws.cell(r + 2, 1, "SLA met (active minutes <= SLA minutes)")
    for col in "BCD":
        ws[f"{col}{r + 2}"] = f'=IF({col}{r + 1}<={col}{r},"met","missed")'
    ws.cell(r + 3, 1, "SLA attainment (met / completed eligible, 1 case)")
    for col in "BCD":
        ws[f"{col}{r + 3}"] = f'=IF({col}{r + 2}="met",1,0)/1'
        ws[f"{col}{r + 3}"].number_format = "0.000000"
    for col, st in ((5, "pause_eligible"), (6, "pause_not_eligible")):
        ws.cell(r + 3, col, float(EXPECTED["states"][st]["metrics"]["sla_attainment"])).font = BLUE
        ws.cell(r + 3, col).number_format = "0.000000"
    ws.cell(r + 3, 7, f"=ABS(C{r + 3}-E{r + 3})<0.0000005")
    ws.cell(r + 3, 8, f"=ABS(D{r + 3}-F{r + 3})<0.0000005")
    ws.cell(r + 5, 1, "All checks (A and B) pass").font = BOLD
    ws.cell(r + 5, 2, f"=AND(G{r0}:H{r0 + 4},G{r + 3}:H{r + 3})")
    ws.cell(r + 7, 1, "Why 66 / 2 / 65 / 1: calendar = Fri 16:00 -> Mon 10:00 on the wall clock; business keeps only the working "
                      "intervals (Fri 16:00-17:00, Mon 09:00-10:00); the waiting interval Fri 16:30 -> Mon 09:30 is 65 h of wall clock "
                      "and 1 h of working time, so active = 2 - 1 = 1 when its reason is pause-eligible and 2 when not.")
    widths(ws, {"A": 46, "B": 24, "C": 24, "D": 22, "E": 24, "F": 24, "G": 12, "H": 12})
    return {"figures_first": r0}


ASSUMPTIONS = [
    ("D-01", "Time, calendars and intervals. UTC instants; calendar = Monday-Friday 09:00-17:00 UTC with no holidays in the fixture; "
             "durations in hours at 4 decimals; calendar cycle = completed - opened; business cycle = sum of overlap with working intervals; "
             "paused = waiting intervals whose reason is in the SLA rule's pause_policy; active = business - business overlap of eligible pauses; "
             "customer waiting is never counted as employee processing."),
    ("D-02", "Workflow versions and cases. A published version is immutable and a case keeps the version it was created with. "
             "The workbook uses the single version wfv-001 of the golden fixture."),
    ("D-03", "Steps, documents, approvals. Documents are append-only; approval needs an independent reviewer of the step's role; completion needs every "
             "required step approved and every required document accepted. Not modelled in the workbook (the fixture satisfies them)."),
    ("D-04", "Rates and backlog. SLA attainment = completed eligible cases within SLA / completed eligible cases; open cases are shown beside it. "
             "The workbook has one completed case, so attainment is 0 or 1."),
    ("D-05", "Fixture-driven clarifications. SLA test: active hours x 60 <= business minutes; step waiting = started_at - ready_at in business hours; "
             "population durations are means over accepted completed cases."),
    ("D-06", "Schema clarifications. calendar_intervals.kind in {working, holiday} (default working); waiting intervals are clipped to the case span "
             "and merged before deduction; active hours floor at 0. The Overlap sheet clips the waiting interval but has one interval, so no merge is needed."),
    ("W-01", "Workbook limit: one case, one waiting interval, no holidays. Holiday subtraction and merging of overlapping waits are done by the SQL and Python "
             "implementations (docs/sql-analysis.md), not by these formulas."),
    ("W-02", "Workbook limit: instants are stored as Excel date serials (UTC, no zone). Each overlap row is ROUND(...,6) hours to remove binary floating-point noise "
             "(1 s = 0.000278 h, so 6 decimals cannot hide a real second); totals are ROUND(...,4)."),
    ("W-03", "Recalculation: openpyxl writes formulas only. The delivered file was recalculated by LibreOffice headless and carries its cached values; Excel was NOT run "
             "(see docs/workbook.md). fullCalcOnLoad is set so Excel recalculates on open."),
    ("W-04", "Snapshot summary is a paste of the output of data/sql/analysis/run_all.py (artifacts/analysis-results.json) and goes stale if the data changes."),
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions: contract decisions D-01..D-06 and the workbook's own limits")
    header(ws, 3, ["ID", "Statement"])
    for i, (k, v) in enumerate(ASSUMPTIONS):
        ws.cell(4 + i, 1, k).font = BOLD
        c = ws.cell(4 + i, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 8, "B": 150})


def put_table(ws, r, name, rows, cols=None):
    ws.cell(r, 1, name).font = BOLD
    for c in range(1, 14):
        ws.cell(r, c).fill = SEC
    cols = cols or list(rows[0])
    header(ws, r + 1, cols)
    for i, row in enumerate(rows):
        for j, c in enumerate(cols):
            v = row.get(c)
            if isinstance(v, bool):
                v = str(v).lower()       # LibreOffice would rewrite a real boolean as =TRUE()
            elif isinstance(v, str):
                try:
                    v = float(v) if v.replace(".", "", 1).replace("-", "", 1).isdigit() else v
                except ValueError:
                    pass
            ws.cell(r + 2 + i, j + 1, v)
    return r + len(rows) + 4


def build_snapshot(wb):
    ws = wb.create_sheet("Snapshot summary")
    ws["A1"] = "PASTED VALUES (not formulas) from data/sql/analysis/run_all.py; each section is titled with its query file"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (f"SYNTHETIC. Dataset {MANIFEST['dataset_version']}, generator dataset_hash {MANIFEST['dataset_hash']}, contract {MANIFEST['contract_version']}; "
                "database seeded through the domain's validate_batch (733 of 740 cases accepted; 269 rejected rows kept).")
    r = 4
    for key in ("01_cycle_time_by_opening_month", "02_version_comparison_comparable_window", "03_sla_by_segment_region",
                "04_step_waiting_by_step_assignee", "05_rework_first_pass_by_segment", "06_paused_time_by_reason_region",
                "07_backlog_month_ends", "08_flags_and_rejections_by_reason", "09_golden_case_business_hours"):
        r = put_table(ws, r, f"{key}.sql", RESULTS[key])
    r = put_table(ws, r, "golden case-001 per-case hours: analysis SQL 09 vs data/sql/reference_metrics.sql vs Python domain vs expected.json",
                  RESULTS["golden_three_sources"])
    r = put_table(ws, r, "reference_metrics.sql over the whole snapshot (cross-check of query 01 TOTAL)", RESULTS["reference_metrics_snapshot"])
    ok = sum(1 for x in RESULTS["reconciliation"] if x["equal"] is True)
    r = put_table(ws, r, "reconciliation (run_all.py): checks equal / checks run",
                  [{"checks_equal": ok, "checks_run": len(RESULTS["reconciliation"])}])
    widths(ws, {"A": 26, "B": 22, "C": 22, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16, "J": 16, "K": 16})


def recalc(path: Path) -> None:
    """Open with LibreOffice headless and save a recalculated copy over ``path`` (cached values included)."""
    if shutil.which("soffice") is None:
        raise SystemExit("soffice not installed; the file has formulas but no cached values")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in" / path.name
        src.parent.mkdir()
        shutil.copy(path, src)
        out = Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=240)
        shutil.copy(out / path.name, path)
    set_full_calc_on_load(path)


def set_full_calc_on_load(path: Path) -> None:
    """LibreOffice drops fullCalcOnLoad; put it back (workbook.xml only) so Excel recalculates when it opens the file.
    The cached values LibreOffice computed stay in the sheets."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode("utf-8")
                assert "<calcPr " in text and "fullCalcOnLoad" not in text
                data = text.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ', 1).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def main(argv: list[str]) -> int:
    wb = Workbook()
    build_inputs(wb)
    n = build_calendar(wb)
    ov = build_overlap(wb, n)
    build_figures(wb, ov)
    build_assumptions(wb)
    build_snapshot(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(wb.sheetnames)} sheets: {', '.join(wb.sheetnames)})")
    if "--recalc" in argv:
        recalc(OUT)
        print("recalculated with LibreOffice headless; cached values written")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
