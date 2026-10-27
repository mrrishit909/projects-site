"""Build artifacts/pipeline-observatory-analyst-workbook.xlsx with openpyxl (P08-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: the published snapshot ``web/public/data`` through the independent Python side (``data/sql/analysis/python_side.py``): the six sources, the 1,311 batches with their per-batch
counts (COMPLETE world), the 1,302 checkpoints, the 1,343 runs with their gate and task counts, the 25 incidents, the backfill partitions, the 49 rate-limited fetches and the sample's
own classification; the recorded latest accepted event per source (manifest, RECORDED scope); the golden fixture's ten rows (tests/fixtures/golden/inputs.json); the Python domain over the
published files and the manifest's recorded statistics (pasted into ``Domain reference`` for the check columns only) and artifacts/analysis-results.json (data/sql/analysis/run_all.py).
openpyxl cannot calculate, so the file it writes holds formulas only. With --recalc, LibreOffice (soffice) opens it headless and saves a recalculated copy over the same path, which then
also carries cached values.

What the formulas do, live: each source's received / accepted / duplicate / quarantined / deleted from the batches, the rejection and duplicate rates, the scope switch ``Controls!B5``
(complete world or sample) that selects which rates the Selected columns show, the ingestion age from the checkpoints, the freshness from the recorded latest event, the over-target and
silent flags (``Controls!B3`` = how many declared cadences make a source silent), completeness against the declared expectation; the gate decision, stage and failure kind of every one of 1,343
runs from its check counts; the run funnel, task success against attempt success, the verified recovery minutes of the 25 incidents, backfill coverage, retry-after compliance of the 49
retries; the publish verdict by the printed rule of the site (blocked latest gate or a source past its target = do not publish; a silent source = publish with a stated exception); and the
golden case: ten rows classified by formula (8 accepted, 1 duplicate, 1 quarantined), freshness 52,200 s against an ingestion age of 42,900 s, 8 of 9, 1.000000 against 0.800000, the residual
0.00, 475.0000 minutes, 2 of 3, 0.666667.
What the spreadsheet does not do: classify the 40,191 raw rows (the per-batch counts are the generator's domain run over all of them, typed) or find the latest accepted event of a source
(typed from the generator's record, labelled RECORDED); those inputs are compared with the Python domain and the manifest after recalculation (limits W-02, W-03). Excel is never run.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "domain"), str(ROOT), str(ROOT / "data" / "sql" / "analysis")]

import python_side as ps  # noqa: E402
from pipeline_observatory import run_pipeline  # noqa: E402
from prepare_snapshot import snapshot_inputs  # noqa: E402

NAME = "pipeline-observatory-analyst-workbook.xlsx"
OUT = ROOT / "artifacts" / NAME
RESULTS_PATH = ROOT / "artifacts" / "analysis-results.json"
MANIFEST = json.loads((ROOT / "web/public/data/manifest.json").read_text())
GOLDEN = json.loads((ROOT / "tests/fixtures/golden/inputs.json").read_text())
EXPECTED = json.loads((ROOT / "tests/fixtures/golden/expected.json").read_text())
CODE_REVISION = "ed7eb5a"
DATASET_HASH = "0a90687778634365c402c5431661b56625de20785e28d9b2388d36c0cad19aab"
HDR = PatternFill("solid", fgColor="23303A")
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
EXCEL_EPOCH = datetime(1899, 12, 30, tzinfo=timezone.utc)
SHEET_ORDER = ["Read me", "Controls", "Checks", "Publish decision", "Source health", "Sample scope", "Golden case", "Assumptions", "Sources", "Batches", "Checkpoints", "Runs",
               "Incidents", "Backfills", "Retries", "Fetch log", "Domain reference", "SQL results (pasted)"]
KINDS = ("freshness", "relationships", "residual", "ingest", "transform", "gate_infra")


def serial(s: str | None):
    return None if s is None else (ps.ts(s) - EXCEL_EPOCH).total_seconds() / 86400


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def put(ws, rows: list[list], first_row: int, blue: bool = True):
    for r, row in enumerate(rows, start=first_row):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if blue and v is not None and not (isinstance(v, str) and v.startswith("=")):
                cell.font = BLUE


def fl(x) -> float | None:
    return None if x is None else float(x)


class Data:
    """Typed inputs from the snapshot through the independent Python side, plus the domain's and the manifest's figures for the check columns."""

    def __init__(self) -> None:
        self.c = ps.Ctx()
        inputs, self.manifest = snapshot_inputs()
        self.dom = run_pipeline(inputs).metrics               # the domain over the PUBLISHED files: complete for runs, tasks, incidents, backfills, checkpoints; the sample for raw rows
        self.sources = sorted(self.c.contracts)
        self.rec = self.manifest["stats"]
        self.sample_cls = self.sample_classes()

    def sample_classes(self) -> dict[str, dict[str, int]]:
        out = {s: {"received": 0, "accepted": 0, "duplicate": 0, "quarantined": 0, "deleted": 0} for s in self.sources}
        for r in self.c.raw:
            if r["batch_id"] in self.c.committed:
                v = self.c.sample[r["raw_record_id"]]
                o = out[self.c.batches[r["batch_id"]]["source_id"]]
                o["received"] += 1
                o[v] += 1
        return out


# ---- data sheets -------------------------------------------------------------------------------------------------------------------------------------------------------------
def build_sources(wb, data: Data):
    ws = wb.create_sheet("Sources")
    title(ws, "Source contracts (6) and the RECORDED latest accepted event of each", "latest_accepted_event is RECORDED by the generator (manifest stats.freshness_at_as_of): the published raw rows are a sample and cannot show it. expected_records = the declared expectation over 56 partitions (complete world).")
    header(ws, 3, ["source_id", "expected_frequency", "cadence_seconds", "freshness_target_seconds", "schema_version", "expected_records", "latest_accepted_event (recorded, serial)", "sample_rows_in_the_files"])
    rows = []
    for s in data.sources:
        ct = data.c.contracts[s]
        exp = sum(e["expected_records"] for e in data.c.expect if e["source_id"] == s)
        rows.append([s, ct["expected_frequency"], ps.cadence_seconds(ct["expected_frequency"]), ct["freshness_target_seconds"], ct["schema_version"], exp,
                     serial(data.rec["freshness_at_as_of"][s]["latest_accepted_event_at"]), sum(1 for r in data.c.raw if data.c.batches[r["batch_id"]]["source_id"] == s)])
    put(ws, rows, 4)
    for i in range(len(rows)):
        ws.cell(row=4 + i, column=7).number_format = "yyyy-mm-dd hh:mm:ss"
    widths(ws, {"A": 16, "B": 20, "G": 34})
    return 4 + len(rows) - 1


def build_batches(wb, data: Data):
    ws = wb.create_sheet("Batches")
    title(ws, "Ingestion batches (1,311): the per-batch counts the generator's domain run computed over ALL raw rows (COMPLETE world)", "received counts committed batches only; the schema-quarantined batch (status quarantined) and the aborted batches count for nothing.")
    header(ws, 3, ["batch_id", "source_id", "status", "kind", "started_at (serial)", "written_rows", "received", "accepted_unique", "duplicate", "quarantined", "deleted", "sampled"])
    bs = sorted(data.c.batches)
    rows = []
    for b in bs:
        r, s = data.c.batches[b], data.c.summaries[b]
        rows.append([b, r["source_id"], r["status"], s["kind"], serial(r["started_at"]), s["written_rows"], s["received"], s["accepted_unique"], s["duplicate"], s["quarantined"], s["deleted"], bool(s["sampled"])])
    put(ws, rows, 4)
    widths(ws, {"A": 10, "B": 16, "C": 12, "D": 16, "E": 20})
    return 4 + len(rows) - 1


def build_checkpoints(wb, data: Data):
    ws = wb.create_sheet("Checkpoints")
    title(ws, "Checkpoints (1,302): one per committed batch, committed with its rows in one transaction", "The last successful ingestion of a source is its latest committed_at.")
    header(ws, 3, ["batch_id", "source_id", "committed_at (serial)", "version_number"])
    cps = sorted(data.c.checkpoints, key=lambda c: (c["source_id"], c["version_number"]))
    put(ws, [[c["batch_id"], c["source_id"], serial(c["committed_at"]), c["version_number"]] for c in cps], 4)
    widths(ws, {"A": 10, "B": 16, "C": 22})
    return 4 + len(cps) - 1


def build_runs(wb, data: Data):
    ws = wb.create_sheet("Runs")
    title(ws, "Pipeline runs (1,343): typed gate and task counts, the stage and failure kind by formula (COMPLETE world)",
          "gate_decision: not_reached when the run has no check result; allowed when a critical result exists and none failed; else blocked (D-P8-08). failure_kind: the failed critical check, or the task that failed last.")
    header(ws, 3, ["run_id", "started_at (serial)", "status", "duration_seconds", "checks", "critical_checks", "critical_failed", "failed_task", "first_failed_critical", "published", "publication_id",
                   "tasks_terminal", "tasks_successful", "attempts", "attempts_successful", "gate_decision", "stage", "failure_kind"])
    c = data.c
    tk: dict[tuple[str, str], list[dict]] = {}
    for t in c.tasks:
        tk.setdefault((t["run_id"], t["task_name"]), []).append(t)
    rows = []
    for r in c.runs:
        rid = r["run_id"]
        info = c.run[rid]
        g = c.gate.get(rid, {})
        mine = {n: v for (run, n), v in tk.items() if run == rid}
        fin = {n: max(v, key=lambda t: t["attempt"])["status"] for n, v in mine.items()}
        atts = [t for v in mine.values() for t in v]
        rows.append([rid, serial(r["started_at"]), r["status"], info["duration"], len(g), sum(q["severity"] == "critical" for q in g.values()), sum(q["severity"] == "critical" and q["failure_count"] > 0 for q in g.values()),
                     ",".join(sorted(n for n, s in fin.items() if s == "failed")) or None, (info["failed_critical"] or [None])[0], info["published"], c.pub_by_run[rid]["publication_id"] if info["published"] else None,
                     len(mine), sum(s == "success" for s in fin.values()), len(atts), sum(t["status"] == "success" for t in atts)])
    put(ws, rows, 4)
    last = 4 + len(rows) - 1
    for i in range(len(rows)):
        r = 4 + i
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row=r, column=16, value=f'=IF(E{r}=0,"not_reached",IF(AND(F{r}>0,G{r}=0),"allowed","blocked"))')
        ws.cell(row=r, column=17, value=f'=IF(C{r}="success","published",IF(E{r}>0,"blocked_at_gate","failed_before_gate"))')
        ws.cell(row=r, column=18, value=f'=IF(C{r}="success","",IF(E{r}>0,IF(I{r}="freshness_within_target","freshness",IF(I{r}="relationships_intact","relationships",IF(I{r}="reconciliation_residual_zero","residual",I{r}))),IF(H{r}="quality_gate","gate_infra",H{r})))')
    widths(ws, {"A": 10, "B": 20, "H": 14, "I": 28, "K": 12, "P": 14, "Q": 20, "R": 14})
    return last


def build_incidents(wb, data: Data):
    ws = wb.create_sheet("Incidents")
    title(ws, "Incidents (25): detection and the verified recovery (the recovery run's publication), the minutes by formula", "Recovery time = verified recovery - detection, never a process restart (D-P8-06).")
    header(ws, 3, ["incident_id", "run_id", "detected_at (serial)", "verified_at (serial)", "recovery_seconds", "recovery_minutes", "failure_kind"])
    c = data.c
    rows = [[i["incident_id"], i["run_id"], serial(i["opened_at"]), serial(c.pub_by_run[i["recovery_run_id"]]["published_at"])] for i in c.incidents]
    put(ws, rows, 4)
    last = 4 + len(rows) - 1
    for i in range(len(rows)):
        r = 4 + i
        ws.cell(row=r, column=3).number_format = ws.cell(row=r, column=4).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row=r, column=5, value=f"=ROUND((D{r}-C{r})*86400,0)")
        ws.cell(row=r, column=6, value=f"=ROUND(E{r}/60,4)")
        ws.cell(row=r, column=7, value=f"=INDEX(Runs!$R$4:$R${4 + len(c.runs) - 1},MATCH(B{r},Runs!$A$4:$A${4 + len(c.runs) - 1},0))")
    widths(ws, {"A": 12, "B": 10, "C": 20, "D": 20, "G": 16})
    return last


def build_backfills(wb, data: Data):
    ws = wb.create_sheet("Backfills")
    title(ws, "Backfill partitions (19): the 15 partitions of the 13 accepted requests and the 4 of the 2 refused for overlap", "Coverage counts the accepted requests only; the refused partitions are beside it, never inside it (D-09 R8-11).")
    header(ws, 3, ["backfill_id", "source_id", "outcome", "partition", "verified"])
    c = data.c
    rows = []
    for outcome, reqs in (("accepted", c.backfills), ("refused", c.refused)):
        for r in reqs:
            for p in sorted(r["partitions"], key=lambda p: p["partition"]):
                rows.append([r["backfill_id"], r["source_id"], outcome, p["partition"], bool(p["verified"]) if outcome == "accepted" else False])
    rows.sort(key=lambda x: (x[0], x[3]))
    put(ws, rows, 4)
    widths(ws, {"A": 12, "B": 16, "C": 10, "D": 12})
    return 4 + len(rows) - 1


def build_retries(wb, data: Data):
    ws = wb.create_sheet("Retries")
    title(ws, "Rate-limited fetches (49) and the next fetch of the same cursor in the same session", "waited and slack by formula; a retry honours the hint when waited >= retry_after (P08-TEST-03).")
    header(ws, 3, ["fetch_id", "source_id", "retry_after_seconds", "rate_limited_at (serial)", "retried_at (serial)", "waited_seconds", "slack_seconds", "honoured"])
    rows = []
    fl_ = data.c.fetch
    key = lambda f: (f["fetched_at"], f["fetch_id"])
    for f in fl_:
        if f["status"] != "rate_limited":
            continue
        n = min((x for x in fl_ if x["source_id"] == f["source_id"] and x["session_id"] == f["session_id"] and x["cursor_value"] == f["cursor_value"] and key(x) > key(f)), key=key)
        rows.append([f["fetch_id"], f["source_id"], f["retry_after_seconds"], serial(f["fetched_at"]), serial(n["fetched_at"])])
    put(ws, rows, 4)
    for i in range(len(rows)):
        r = 4 + i
        ws.cell(row=r, column=4).number_format = ws.cell(row=r, column=5).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row=r, column=6, value=f"=ROUND((E{r}-D{r})*86400,0)")
        ws.cell(row=r, column=7, value=f"=F{r}-C{r}")
        ws.cell(row=r, column=8, value=f"=F{r}>=C{r}")
    widths(ws, {"D": 22, "E": 22})
    return 4 + len(rows) - 1


def build_fetch_log(wb, data: Data):
    ws = wb.create_sheet("Fetch log")
    title(ws, "Fetch log (1,361 rows): the status of every fetch", "ok, rate_limited, refetch (a cursor served before in another session, or the re-poll of the last committed cursor), cursor_loop (a cursor repeated inside one session).")
    header(ws, 3, ["fetch_id", "source_id", "session_id", "status"])
    put(ws, [[f["fetch_id"], f["source_id"], f["session_id"], f["status"]] for f in data.c.fetch], 4)
    widths(ws, {"B": 16, "C": 26, "D": 14})
    return 4 + len(data.c.fetch) - 1


# ---- the domain's and the manifest's figures, pasted for the check columns ---------------------------------------------------------------------------------------------
def domain_reference(data: Data) -> dict:
    d, st = data.dom, data.rec
    by_src = {}
    for s in data.sources:
        i, f = st["ingestion"]["sources"][s], st["freshness_at_as_of"][s]
        by_src[s] = {"received": i["received"], "accepted_unique": i["accepted_unique"], "duplicate": i["duplicate"], "quarantined": i["quarantined"], "rejection_rate": float(i["rejection_rate"]),
                     "duplicate_rate": float(i["duplicate_rate"]), "ingestion_age_seconds": d["freshness"][s]["ingestion_age_seconds"], "freshness_seconds": f["freshness_seconds"],
                     "completeness": float(st["completeness_by_source"][s]["value"])}
    qg = d["quality_gate"]
    reached = [v for v in qg.values() if v["results"]]
    return {"sources": by_src,
            "funnel": {"runs": len(qg), "reached_gate": len(reached), "allowed": sum(v["decision"] == "allowed" for v in reached), "blocked_at_gate": sum(v["decision"] == "blocked" for v in reached),
                       "failed_before_gate": len(qg) - len(reached)},
            "runs": {k: (float(v) if isinstance(v, str) else v) for k, v in d["runs"]["overall"].items()},
            "incidents": {"total": d["incidents"]["total"], "mean_recovery_minutes": float(d["incidents"]["mean_recovery_minutes"])},
            "backfills": {"requested": d["backfills"]["overall"]["requested_partitions"], "verified": d["backfills"]["overall"]["verified_partitions"], "coverage": float(d["backfills"]["overall"]["coverage"]),
                          "refused": len(d["backfills"]["refused"])},
            "pointer": d["publications"]["current_publication_id"], "by_type": st["incidents"]["by_type"], "retry": st["retry"]}


def build_domain_reference(wb, ref: dict, data: Data):
    ws = wb.create_sheet("Domain reference")
    title(ws, "PASTED VALUES: the Python domain over the published files and the manifest's recorded statistics (check columns only)",
          "Run, task, incident, backfill and ingestion-age figures are the domain's over the published files (complete for those entities). Per-source ingestion counts, freshness and completeness are the generator's RECORD for the complete world (manifest stats): the files hold only a sample of the raw rows.")
    header(ws, 3, ["source_id", "received", "accepted_unique", "duplicate", "quarantined", "rejection_rate", "duplicate_rate", "ingestion_age_seconds", "freshness_seconds", "completeness", "source of the counts"])
    rows = []
    for s in data.sources:
        v = ref["sources"][s]
        rows.append([s, v["received"], v["accepted_unique"], v["duplicate"], v["quarantined"], v["rejection_rate"], v["duplicate_rate"], v["ingestion_age_seconds"], v["freshness_seconds"], v["completeness"],
                     "manifest stats (recorded, complete world); ingestion age: domain"])
    put(ws, rows, 4)
    r0 = 4 + len(rows) + 1
    ws.cell(row=r0, column=1, value="Run funnel, tasks and attempts (domain over the published files)").font = BOLD
    header(ws, r0 + 1, ["runs", "reached_gate", "allowed", "blocked_at_gate", "failed_before_gate", "task_success_rate", "attempt_success_rate", "tasks_terminal", "attempts"])
    f, ru = ref["funnel"], ref["runs"]
    put(ws, [[f["runs"], f["reached_gate"], f["allowed"], f["blocked_at_gate"], f["failed_before_gate"], ru["task_success_rate"], ru["attempt_success_rate"], ru["tasks_terminal"], ru["attempts"]]], r0 + 2)
    r1 = r0 + 4
    ws.cell(row=r1, column=1, value="Incidents, backfills, retries").font = BOLD
    header(ws, r1 + 1, ["incidents", "mean_recovery_minutes", "backfill_requested", "backfill_verified", "backfill_coverage", "refused_requests", "rate_limited", "retries_honoured", "pointer"])
    put(ws, [[ref["incidents"]["total"], ref["incidents"]["mean_recovery_minutes"], ref["backfills"]["requested"], ref["backfills"]["verified"], ref["backfills"]["coverage"], ref["backfills"]["refused"],
              ref["retry"]["rate_limited_responses"], ref["retry"]["retries_honouring_retry_after"], ref["pointer"]]], r1 + 2)
    widths(ws, {"A": 16, "K": 56})
    return {"src_first": 4, "src_last": 4 + len(rows) - 1, "funnel_row": r0 + 2, "inc_row": r1 + 2}


# ---- formula sheets ---------------------------------------------------------------------------------------------------------------------------------------------------------
def build_source_health(wb, data: Data, bat_last: int, cp_last: int, dref: dict):
    ws = wb.create_sheet("Source health")
    title(ws, "Source health at the as-of: counts and rates (complete world), ingestion age, freshness (recorded), flags",
          "Controls!B3 = the number of declared cadences that make a source silent; Controls!B5 = the scope of the Selected rates. Check = the row agrees with the domain / the manifest's record (Domain reference).")
    header(ws, 3, ["source_id", "cadence_seconds", "target_seconds", "committed_batches", "received", "accepted_unique", "duplicate", "quarantined", "deleted", "rejection_rate", "duplicate_rate",
                   "selected_rejection_rate", "selected_duplicate_rate", "ingestion_age_seconds", "freshness_seconds", "over_target", "silent", "expected_records", "completeness", "check"])
    b = f"Batches!$%s$4:$%s${bat_last}"
    n = len(data.sources)
    last = 4 + n - 1
    for i, s in enumerate(data.sources):
        r = 4 + i
        ws.cell(row=r, column=1, value=s).font = BLUE
        ws.cell(row=r, column=2, value=f"=INDEX(Sources!$C$4:$C$9,MATCH(A{r},Sources!$A$4:$A$9,0))")
        ws.cell(row=r, column=3, value=f"=INDEX(Sources!$D$4:$D$9,MATCH(A{r},Sources!$A$4:$A$9,0))")
        ws.cell(row=r, column=4, value=f'=COUNTIFS({b % ("B", "B")},A{r},{b % ("C", "C")},"committed")')
        for col, src in zip("EFGHI", "GHIJK"):
            ws[f"{col}{r}"] = f'=SUMIFS({b % (src, src)},{b % ("B", "B")},$A{r},{b % ("C", "C")},"committed")'
        ws[f"J{r}"] = f"=ROUND(H{r}/E{r},6)"
        ws[f"K{r}"] = f"=ROUND(G{r}/E{r},6)"
        sr = 4 + i
        ws[f"L{r}"] = f'=IF(Controls!$B$5="sample",\'Sample scope\'!G{sr},J{r})'
        ws[f"M{r}"] = f'=IF(Controls!$B$5="sample",\'Sample scope\'!H{sr},K{r})'
        ws[f"N{r}"] = f"=ROUND((Controls!$B$4-_xlfn.MAXIFS(Checkpoints!$C$4:$C${cp_last},Checkpoints!$B$4:$B${cp_last},A{r}))*86400,0)"
        ws[f"O{r}"] = f"=ROUND((Controls!$B$4-INDEX(Sources!$G$4:$G$9,MATCH(A{r},Sources!$A$4:$A$9,0)))*86400,0)"
        ws[f"P{r}"] = f"=O{r}>C{r}"
        ws[f"Q{r}"] = f"=O{r}>Controls!$B$3*B{r}"
        ws[f"R{r}"] = f"=INDEX(Sources!$F$4:$F$9,MATCH(A{r},Sources!$A$4:$A$9,0))"
        ws[f"S{r}"] = f"=ROUND(F{r}/R{r},6)"
        d = 4 + i
        ws[f"T{r}"] = (f"=AND(E{r}='Domain reference'!B{d},F{r}='Domain reference'!C{d},G{r}='Domain reference'!D{d},H{r}='Domain reference'!E{d},ROUND(J{r}-'Domain reference'!F{d},6)=0,"
                       f"ROUND(K{r}-'Domain reference'!G{d},6)=0,N{r}='Domain reference'!H{d},O{r}='Domain reference'!I{d},ROUND(S{r}-'Domain reference'!J{d},6)=0)")
    t = last + 1
    ws[f"A{t}"] = "all sources"
    ws[f"A{t}"].font = BOLD
    for col in "DEFGHI":
        ws[f"{col}{t}"] = f"=SUM({col}4:{col}{last})"
    ws[f"J{t}"] = f"=ROUND(H{t}/E{t},6)"
    ws[f"K{t}"] = f"=ROUND(G{t}/E{t},6)"
    ws[f"P{t}"] = f"=COUNTIF(P4:P{last},TRUE)"
    ws[f"Q{t}"] = f"=COUNTIF(Q4:Q{last},TRUE)"
    ws[f"T{t}"] = f"=AND(T4:T{last})"
    widths(ws, {"A": 16, "T": 8})
    return last, t


def build_sample_scope(wb, data: Data):
    ws = wb.create_sheet("Sample scope")
    title(ws, "The same rows under two scopes: the 2,000-row sample classified from the loaded rows only, and the complete world's classification of the SAME sampled batches",
          "The sample overstates accepted and understates duplicates because a key whose first copy sits in an unsampled batch looks new. Typed from the independent Python side (python_side.classify_sample) and the batch summaries.")
    header(ws, 3, ["source_id", "sample_received", "sample_accepted", "sample_quarantined", "sample_duplicate", "sample_deleted", "sample_rejection_rate", "sample_duplicate_rate",
                   "world_accepted (same batches)", "world_duplicate (same batches)", "accepted_delta", "duplicate_delta"])
    rows = []
    for s in data.sources:
        sm = data.sample_cls[s]
        wb_ = [x for b, x in data.c.summaries.items() if x["source_id"] == s and x["sampled"] and b in data.c.committed]
        rows.append([s, sm["received"], sm["accepted"], sm["quarantined"], sm["duplicate"], sm["deleted"], None, None, sum(x["accepted_unique"] for x in wb_), sum(x["duplicate"] for x in wb_)])
    put(ws, rows, 4)
    last = 4 + len(rows) - 1
    for r in range(4, last + 1):
        ws[f"G{r}"] = f"=ROUND(D{r}/B{r},6)"
        ws[f"H{r}"] = f"=ROUND(E{r}/B{r},6)"
        ws[f"K{r}"] = f"=C{r}-I{r}"
        ws[f"L{r}"] = f"=E{r}-J{r}"
    t = last + 1
    ws[f"A{t}"] = "all sources"
    ws[f"A{t}"].font = BOLD
    for col in "BCDEFIJKL":
        ws[f"{col}{t}"] = f"=SUM({col}4:{col}{last})"
    ws[f"G{t}"] = f"=ROUND(D{t}/B{t},6)"
    ws[f"H{t}"] = f"=ROUND(E{t}/B{t},6)"
    ws[f"A{t + 2}"] = "Selected scope (Controls!B5) headline duplicate rate"
    ws[f"B{t + 2}"] = f'=IF(Controls!$B$5="sample",H{t},\'Source health\'!K10)'
    widths(ws, {"A": 16})
    return t


def build_publish_decision(wb, data: Data, runs_last: int, inc_last: int, bf_last: int, rt_last: int, fl_last: int, sh_last: int, dref: dict):
    ws = wb.create_sheet("Publish decision")
    title(ws, "Is this dataset safe, complete and fresh enough to publish? The printed rule of the site, as formulas",
          "Rule: a blocked latest gate or a source past its declared freshness target means DO NOT PUBLISH; a silent source (freshness above Controls!B3 declared cadences) means PUBLISH WITH A STATED EXCEPTION. Scope: complete world except freshness (recorded).")
    R = f"Runs!$%s$4:$%s${runs_last}"
    rr = lambda col: R % (col, col)
    rows = [
        ("Latest run", f"=INDEX({rr('A')},ROWS({rr('A')}))"),
        ("Latest run: status", f"=INDEX({rr('C')},ROWS({rr('C')}))"),
        ("Latest run: gate decision", f"=INDEX({rr('P')},ROWS({rr('P')}))"),
        ("Publication pointer (publication of the latest published run)", f"=LOOKUP(2,1/({rr('J')}=TRUE),{rr('K')})"),
        ("Sources past their declared freshness target", f"='Source health'!P{sh_last + 1}"),
        ("Silent sources", f"='Source health'!Q{sh_last + 1}"),
        ("First silent source", f'=IFERROR(INDEX(\'Source health\'!A4:A{sh_last},MATCH(TRUE,\'Source health\'!Q4:Q{sh_last},0)),"none")'),
        ("Verdict", '=IF(B5<>"allowed","DO NOT PUBLISH: the latest gate did not allow the run",IF(B7>0,"DO NOT PUBLISH: a source is past its declared freshness target",IF(B8>0,"PUBLISH WITH A STATED EXCEPTION: "&B8&IF(B8=1," source is"," sources are")&" silent while the jobs stay green, the first is "&B9,"PUBLISH")))'),
    ]
    for i, (k, f) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = BOLD if k == "Verdict" else Font()
        ws.cell(row=i, column=2, value=f)
    # funnel
    ws["A13"] = "Run funnel (D-09 R8-07)"
    ws["A13"].font = BOLD
    funnel = [("runs", f"=COUNTA({rr('A')})"), ("reached the gate", f'=COUNTIF({rr("P")},"allowed")+COUNTIF({rr("P")},"blocked")'), ("allowed", f'=COUNTIF({rr("P")},"allowed")'),
              ("blocked at the gate", f'=COUNTIF({rr("Q")},"blocked_at_gate")'), ("failed before the gate", f'=COUNTIF({rr("Q")},"failed_before_gate")'), ("published", f"=COUNTIF({rr('J')},TRUE)")]
    for i, (k, f) in enumerate(funnel, start=14):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=f)
    ws["A21"] = "Failed runs by kind"
    ws["A21"].font = BOLD
    for i, k in enumerate(KINDS, start=22):
        ws.cell(row=i, column=1, value=k).font = BLUE
        ws.cell(row=i, column=2, value=f"=COUNTIF({rr('R')},A{i})")
        ws.cell(row=i, column=3, value=f'=ROUND(AVERAGEIFS(Incidents!$E$4:$E${inc_last},Incidents!$G$4:$G${inc_last},A{i})/60,4)')
    ws["C21"] = "mean recovery minutes"
    ws["A29"] = "Tasks against attempts"
    ws["A29"].font = BOLD
    ws["A30"], ws["B30"] = "task success rate", f"=ROUND(SUM({rr('M')})/SUM({rr('L')}),6)"
    ws["A31"], ws["B31"] = "attempt success rate", f"=ROUND(SUM({rr('O')})/SUM({rr('N')}),6)"
    ws["A33"] = "Incidents and recovery"
    ws["A33"].font = BOLD
    ws["A34"], ws["B34"] = "incidents", f"=COUNTA(Incidents!$A$4:$A${inc_last})"
    ws["A35"], ws["B35"] = "mean recovery minutes (verified recovery minus detection)", f"=ROUND(AVERAGE(Incidents!$E$4:$E${inc_last})/60,4)"
    ws["A37"] = "Backfill coverage (accepted requests only)"
    ws["A37"].font = BOLD
    ws["A38"], ws["B38"] = "requested partitions", f'=COUNTIF(Backfills!$C$4:$C${bf_last},"accepted")'
    ws["A39"], ws["B39"] = "verified partitions", f'=COUNTIFS(Backfills!$C$4:$C${bf_last},"accepted",Backfills!$E$4:$E${bf_last},TRUE)'
    ws["A40"], ws["B40"] = "coverage", "=ROUND(B39/B38,6)"
    ws["A41"], ws["B41"] = "refused partitions (beside, never inside)", f'=COUNTIF(Backfills!$C$4:$C${bf_last},"refused")'
    ws["A43"] = "Retry-after"
    ws["A43"].font = BOLD
    ws["A44"], ws["B44"] = "rate-limited fetches", f'=COUNTIF(\'Fetch log\'!$D$4:$D${fl_last},"rate_limited")'
    ws["A45"], ws["B45"] = "retries honouring the hint", f"=COUNTIF(Retries!$H$4:$H${rt_last},TRUE)"
    ws["A46"], ws["B46"] = "refetch rows | loop rows", f'=COUNTIF(\'Fetch log\'!$D$4:$D${fl_last},"refetch")&" | "&COUNTIF(\'Fetch log\'!$D$4:$D${fl_last},"cursor_loop")'
    widths(ws, {"A": 62, "B": 90, "C": 22})


def build_golden(wb, dref: dict):
    ws = wb.create_sheet("Golden case")
    title(ws, "Golden known answer (P08-GOLD-01, tests/fixtures/golden): ten rows by formula", "Typed inputs are blue; the class of each row, every count, rate, instant difference and the residual are formulas. Check = equals the hand answer of CALCULATION.md.")
    header(ws, 3, ["source_key", "quantity", "unit_price", "amount", "event_ts (serial)", "class (formula)"])
    rows = [r for r in GOLDEN["raw_records"] if r["batch_id"] == "b-001"]
    rows.sort(key=lambda r: r["raw_record_id"])
    first, last = 4, 4 + len(rows) - 1
    for i, r in enumerate(rows):
        p = r["payload_json"]
        for j, v in enumerate([r["source_key"], p["quantity"], float(p["unit_price"]), float(p["amount"]), serial(p["event_ts"])], start=1):
            ws.cell(row=first + i, column=j, value=v).font = BLUE
        ws.cell(row=first + i, column=5).number_format = "yyyy-mm-dd hh:mm"
        row = first + i
        ws.cell(row=row, column=6, value=f'=IF(COUNTIFS(A${first - 1}:A{row - 1},A{row},F${first - 1}:F{row - 1},"accepted")>0,"duplicate",IF(B{row}<0,"quarantined","accepted"))')
    g = last + 2
    ws.cell(row=g, column=1, value="figure").font = BOLD
    ws.cell(row=g, column=2, value="formula result").font = BOLD
    ws.cell(row=g, column=3, value="hand answer (typed)").font = BOLD
    ws.cell(row=g, column=4, value="check").font = BOLD
    k = f"F{first}:F{last}"
    figs = [
        ("received", f"=COUNTA(A{first}:A{last})", 10),
        ("accepted unique", f'=COUNTIF({k},"accepted")', 8),
        ("duplicate", f'=COUNTIF({k},"duplicate")', 1),
        ("quarantined", f'=COUNTIF({k},"quarantined")', 1),
        ("rejection rate", f"=ROUND(B{g + 4}/B{g + 1},6)", 0.1),
        ("duplicate rate", f"=ROUND(B{g + 3}/B{g + 1},6)", 0.1),
        ("freshness seconds (as-of minus the latest accepted event)", f'=ROUND((Controls!$B$4-_xlfn.MAXIFS(E{first}:E{last},{k},"accepted"))*86400,0)', 52200),
        ("ingestion age seconds (as-of minus the last committed batch, b-002 at 12:05)", f"=ROUND((Controls!$B$4-{serial('2026-10-04T12:05:00Z')})*86400,0)", 42900),
        ("completeness (accepted / 9 expected)", f"=ROUND(B{g + 2}/9,6)", 0.888889),
        ("run-001 task success (4 tasks, 4 successful)", "=ROUND(4/4,6)", 1.0),
        ("run-001 attempt success (5 attempts: transform twice; 4 successful)", "=ROUND(4/(1+2+1+1),6)", 0.8),
        ("source total (every received row)", f"=ROUND(SUM(D{first}:D{last}),2)", 289.94),
        ("published total (accepted rows)", f'=ROUND(SUMIF({k},"accepted",D{first}:D{last}),2)', 289.97),
        ("quarantined amount (exclusion)", f'=ROUND(SUMIF({k},"quarantined",D{first}:D{last}),2)', -30.0),
        ("duplicate amount (exclusion)", f'=ROUND(SUMIF({k},"duplicate",D{first}:D{last}),2)', 29.97),
        ("residual = source - published - exclusions", f"=ROUND(B{g + 12}-B{g + 13}-(B{g + 14}+B{g + 15}),2)", 0.0),
        ("recovery minutes (verified 20:15 minus detected 12:20)", f"=ROUND(ROUND(({serial('2026-10-04T20:15:00Z')}-{serial('2026-10-04T12:20:00Z')})*86400,0)/60,4)", 475.0),
        ("backfill coverage (2 of 3 partitions)", "=ROUND(2/3,6)", 0.666667),
        ("processing rate (8 accepted / 12 s)", f"=ROUND(B{g + 2}/12,6)", 0.666667),
    ]
    for i, (name, f, hand) in enumerate(figs, start=g + 1):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=f)
        ws.cell(row=i, column=3, value=hand).font = BLUE
        ws.cell(row=i, column=4, value=f"=ROUND(B{i}-C{i},6)=0")
    widths(ws, {"A": 74, "B": 18, "C": 22, "D": 10, "E": 20, "F": 14})
    return g + 1, g + len(figs)


def build_checks(wb, sh_last: int, g_first: int, g_last: int, dref: dict, sh_total: int):
    ws = wb.create_sheet("Checks")
    title(ws, "Checks: TRUE when every comparison with the Python domain, the manifest's record and the golden hand answer holds")
    ws["A2"] = "All checks pass"
    ws["A2"].font = BOLD
    f, inc = dref["funnel_row"], dref["inc_row"]
    checks = [
        ("Every source row of Source health agrees with the domain / the manifest's record (counts, rates, ingestion age, freshness, completeness)", f"=COUNTIF('Source health'!T4:T{sh_last},TRUE)=ROWS('Source health'!T4:T{sh_last})"),
        ("Run funnel: runs | reached the gate | allowed | blocked at the gate | failed before it equal the domain's gate results",
         f"=AND('Publish decision'!B14='Domain reference'!A{f},'Publish decision'!B15='Domain reference'!B{f},'Publish decision'!B16='Domain reference'!C{f},'Publish decision'!B17='Domain reference'!D{f},'Publish decision'!B18='Domain reference'!E{f})"),
        ("Task success and attempt success equal the domain", f"=AND(ROUND('Publish decision'!B30-'Domain reference'!F{f},6)=0,ROUND('Publish decision'!B31-'Domain reference'!G{f},6)=0)"),
        ("Failed runs by kind equal the manifest's by_type (freshness, relationships, residual, ingest, transform, gate_infra)",
         "=AND(" + ",".join(f"'Publish decision'!B{22 + i}={dref['_by_type'][k]}" for i, k in enumerate(KINDS)) + ")"),
        ("Incidents and mean recovery minutes equal the domain", f"=AND('Publish decision'!B34='Domain reference'!A{inc},ROUND('Publish decision'!B35-'Domain reference'!B{inc},4)=0)"),
        ("Backfill coverage (requested, verified, coverage) equals the domain", f"=AND('Publish decision'!B38='Domain reference'!C{inc},'Publish decision'!B39='Domain reference'!D{inc},ROUND('Publish decision'!B40-'Domain reference'!E{inc},6)=0,'Publish decision'!B41=4)"),
        ("Retry-after: every rate-limited fetch was retried no sooner than its hint", "=AND('Publish decision'!B44='Domain reference'!G" + str(inc) + ",'Publish decision'!B45='Domain reference'!H" + str(inc) + ",'Publish decision'!B44='Publish decision'!B45)"),
        ("Publication pointer equals the domain's current publication", f"='Publish decision'!B6='Domain reference'!I{inc}"),
        ("Publish verdict at the default controls is the stated exception for the one silent source", "=IF(AND(Controls!B3=2,Controls!B4=" + str(serial("2026-10-05T00:00:00Z")) + "),LEFT('Publish decision'!B10,31)=\"PUBLISH WITH A STATED EXCEPTION\",TRUE)"),
        ("Golden case: every formula figure equals the hand answer", f"=COUNTIF('Golden case'!D{g_first}:D{g_last},TRUE)=ROWS('Golden case'!D{g_first}:D{g_last})"),
        ("Sample scope: the sample's accepted count exceeds the world's for the same batches and the duplicate count falls short (the scope warning holds)", f"=AND('Sample scope'!K{4 + len(dref['sources'])}>0,'Sample scope'!L{4 + len(dref['sources'])}<0)"),
    ]
    ws["B2"] = f"=AND(B5:B{4 + len(checks)})"
    ws["A4"], ws["B4"] = "check", "result"
    ws["A4"].font = ws["B4"].font = BOLD
    for i, (name, f_) in enumerate(checks, start=5):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=f_)
    widths(ws, {"A": 150, "B": 12})
    return len(checks)


def build_controls(wb):
    ws = wb.create_sheet("Controls")
    title(ws, "Controls", "Change B3 (cadences) or B5 (scope): Source health and Publish decision follow. Blue cells are inputs.")
    ws["A3"], ws["B3"] = "A source is silent above this many declared cadences (display rule, D-08)", 2
    ws["B3"].font = BLUE
    dv = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    ws["A4"], ws["B4"] = "As-of (UTC)", serial("2026-10-05T00:00:00Z")
    ws["B4"].font = BLUE
    ws["B4"].number_format = "yyyy-mm-dd hh:mm"
    ws["A5"], ws["B5"] = "Scope of the Selected rates (complete world | sample)", "complete world"
    ws["B5"].font = BLUE
    dv2 = DataValidation(type="list", formula1='"complete world,sample"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add("B5")
    ws["A7"], ws["B7"] = "All checks pass (Checks!B2)", "=Checks!B2"
    ws["A7"].font = BOLD
    ws["A8"], ws["B8"] = "Publish verdict (Publish decision!B10)", "='Publish decision'!B10"
    widths(ws, {"A": 74, "B": 80})


def build_readme(wb):
    ws = wb.create_sheet("Read me")
    title(ws, "Pipeline Observatory analyst workbook (P08-ART-05)")
    rows = [
        ("What it is", f"A formula-driven recomputation of the Pipeline Observatory headline measures over the published synthetic snapshot po-20261026-20261026 (dataset hash {DATASET_HASH}, code revision {CODE_REVISION}, as-of 2026-10-05T00:00:00Z): 6 sources, 1,311 batches, 1,343 hourly runs, 25 incidents, 15 backfill requests, 1,361 fetches."),
        ("The decision", "Is this dataset safe, complete and fresh enough to publish? Publish decision applies the printed rule of the site: a blocked latest gate or a source past its freshness target means do not publish; a silent source means publish with a stated exception."),
        ("Synthetic data and scope", "Everything is synthetic (a seeded generator). Every figure carries a scope: COMPLETE world (published in full), 2,000-row SAMPLE (raw rows), or RECORDED (the generator's statement). Observed = what this one draw shows; constructed = the designed relationships and injected defects; no real-world benefit is claimed."),
        ("How to use it", "Open the .xlsx in any spreadsheet application. Controls!B3 changes how many cadences make a source silent; Controls!B5 switches the scope of the Selected rates; Checks!B2 is TRUE when every figure agrees with the Python domain, the manifest's record and the golden hand answer. Nothing needs installing to read the cached values."),
        ("Recalculation status", "openpyxl wrote formulas only; LibreOffice (headless) recalculated the delivered file. Microsoft Excel was NOT run: nothing here was opened or recalculated by Excel (docs/workbook.md)."),
        ("Sheets", "Controls, Checks, Publish decision, Source health, Sample scope, Golden case, Assumptions (read-this-first limits), data sheets (Sources, Batches, Checkpoints, Runs, Incidents, Backfills, Retries, Fetch log), Domain reference (pasted values), SQL results (pasted)."),
        ("Colour", "Blue = typed input or pasted value; black = formula."),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 24, "B": 140})


ASSUMPTIONS = [
    ("A-01", "Gate (D-P8-08): a run reached the gate when it has quality results; allowed iff at least one critical result exists and none failed, else blocked; a run that never reached the gate has no decision (not_reached). 'Blocked' in the domain's pointer_history counts those 15 runs too; the funnel keeps them apart (D-09 R8-07)."),
    ("A-02", "Failure kind: a failed run that reached the gate is named for its failed critical check (freshness_within_target = freshness, relationships_intact = relationships, reconciliation_residual_zero = residual); otherwise for the task that failed last (ingest, transform; quality_gate = gate_infra)."),
    ("A-03", "Rates divide by received (committed batches only; the schema-quarantined batch and the aborted batches count for nothing). A zero denominator is null, never 0. Deleted rows are beside accepted, duplicate and quarantined, not inside a rate's numerator."),
    ("A-04", "Ingestion age = as-of minus the latest checkpoint of a committed batch; freshness = as-of minus the latest ACCEPTED event, which is RECORDED by the generator (the files hold a 2,000-row sample of the raw rows). Silent = freshness above Controls!B3 declared cadences (a display rule, not a contract rule)."),
    ("A-05", "Task success = successful terminal tasks / terminal tasks (the status of a task's last attempt); attempt success = successful attempts / attempts. An attempt is not a task."),
    ("A-06", "Recovery minutes = (verified recovery - detection) / 60, the verified recovery being the publication of the first later run that succeeded, was allowed and reconciled (D-P8-06); the recovery publication instants are typed from the published files."),
    ("A-07", "Backfill coverage = verified / requested partitions of the ACCEPTED requests; the 4 partitions of the 2 requests refused for overlap are shown beside it (D-09 R8-11)."),
    ("W-01", "Observed, constructed and hypothetical results are kept apart: every figure is observed in one draw of a generated world; the designed relationships (DR-01..DR-05) and the twelve injected defects are constructed; no benefit to a real organisation is claimed."),
    ("W-02", "TYPED, not replayed by formulas: the per-batch counts (the generator's domain run over 40,191 raw rows), the gate and task counts of each run, the recorded latest accepted event of each source, the recovery publication instants and the sample's own classification (python_side.py). Each is compared with the Python domain, the manifest or the SQL after recalculation."),
    ("W-03", "Recalculation: openpyxl writes formulas only; the delivered file was recalculated by LibreOffice headless and carries its cached values; Excel was NOT run. fullCalcOnLoad is set so Excel recalculates on open. Functions used: SUMIFS, COUNTIFS, COUNTIF, MAXIFS, AVERAGEIFS, INDEX, MATCH, LOOKUP, ROUND, IF, AND, IFERROR (MAXIFS needs Excel 2019 or later)."),
    ("W-04", "The switches change what the Selected columns and the verdict show; the check cell does not depend on the scope switch, and its verdict check applies at the default controls."),
    ("W-05", "The workbook has no macros, no external links and no credentials."),
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions and limits")
    header(ws, 3, ["id", "statement"])
    for i, (k, v) in enumerate(ASSUMPTIONS, start=4):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 8, "B": 150})


def build_sql_results(wb):
    ws = wb.create_sheet("SQL results (pasted)")
    title(ws, "PASTED VALUES from data/sql/analysis/run_all.py (artifacts/analysis-results.json)", "The headline result sets; reconciliation: every row equal across the SQL, the reference SQL, the Python domain, the manifests and the independent Python side.")
    res = json.loads(RESULTS_PATH.read_text())
    meta = res["meta"][0]
    r = 4
    ws.cell(row=r, column=1, value=f"reconciliation rows {meta['reconciliation_rows']}, equal {meta['reconciliation_equal']}; result sets {meta['result_sets']}, rows {meta['result_rows']}; dataset {meta['dataset_version']} / {meta['dataset_hash']}; code {meta['code_revision']}")
    r += 2
    for key in ("01_ingestion_by_source__overall", "02_run_funnel_by_failure_kind__funnel", "03_task_vs_attempt_success__overall", "04_freshness_vs_ingestion_age__silent_source",
                "05_publication_pointer_and_blocked_runs__summary", "06_incidents_and_verified_recovery__summary", "07_backfill_coverage_and_refused_overlaps__coverage_summary",
                "08_fetch_log_retry_after__retry_summary", "10_sample_scope_reconciliation__rates_by_scope", "09_golden_known_answer__runs_overall"):
        rows = res[key]
        ws.cell(row=r, column=1, value=key).font = BOLD
        r += 1
        cols = list(rows[0])
        header(ws, r, cols)
        r += 1
        for row in rows:
            for j, c in enumerate(cols, start=1):
                v = row[c]
                if isinstance(v, str) and re.fullmatch(r"-?\d+(\.\d+)?", v):
                    v = float(v) if "." in v else int(v)
                ws.cell(row=r, column=j, value=v).font = BLUE
            r += 1
        r += 1
    widths(ws, {"A": 34})


def recalc(path: Path) -> None:
    """Open with LibreOffice headless and save a recalculated copy over ``path`` (cached values included)."""
    if shutil.which("soffice") is None:
        raise SystemExit("soffice not installed; the file has formulas but no cached values")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in" / path.name
        src.parent.mkdir()
        shutil.copy(path, src)
        out = Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=1800)
        shutil.copy(out / path.name, path)
    set_full_calc_on_load(path)


def set_full_calc_on_load(path: Path) -> None:
    """LibreOffice drops fullCalcOnLoad; put it back (workbook.xml only) so Excel recalculates when it opens the file."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode("utf-8")
                assert "<calcPr " in text and "fullCalcOnLoad" not in text
                data = text.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ', 1).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build(path: Path) -> None:
    data = Data()
    ref = domain_reference(data)
    wb = Workbook()
    build_readme(wb)
    build_controls(wb)
    build_sources(wb, data)
    bat_last = build_batches(wb, data)
    cp_last = build_checkpoints(wb, data)
    runs_last = build_runs(wb, data)
    inc_last = build_incidents(wb, data)
    bf_last = build_backfills(wb, data)
    rt_last = build_retries(wb, data)
    fl_last = build_fetch_log(wb, data)
    dref = build_domain_reference(wb, ref, data)
    dref["_by_type"] = ref["by_type"]
    dref["sources"] = ref["sources"]
    sh_last, sh_total = build_source_health(wb, data, bat_last, cp_last, dref)
    build_sample_scope(wb, data)
    build_publish_decision(wb, data, runs_last, inc_last, bf_last, rt_last, fl_last, sh_last, dref)
    g_first, g_last = build_golden(wb, dref)
    build_checks(wb, sh_last, g_first, g_last, dref, sh_total)
    build_assumptions(wb)
    build_sql_results(wb)
    wb._sheets = [wb[n] for n in SHEET_ORDER]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main(argv: list[str]) -> int:
    build(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    if "--recalc" in argv:
        recalc(OUT)
        print("recalculated with LibreOffice headless; cached values written")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
