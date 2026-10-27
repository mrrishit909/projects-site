"""Build artifacts/retentionlab-analyst-workbook.xlsx with openpyxl (P10-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: the published snapshot ``web/public/data`` through the independent Python side (the 1,513-row test table, the 97 scored accounts of the 2026-09-01 queue, the three
scenarios, the per-cutoff counts), tests/fixtures/golden (the four-account known answer is typed from CALCULATION.md) and artifacts/analysis-results.json (data/sql/analysis/run_all.py:
the figures every check column compares with). openpyxl cannot calculate, so the file it writes holds formulas only. With --recalc, LibreOffice (soffice) opens it headless and saves a
recalculated copy over the same path, which then also carries cached values.

What the formulas do, live: the tie-broken rank of every test row for each model (probability descending, then account id, then the earlier cutoff); average precision in the
scikit-learn step form (the mean, over churning rows, of the precision at that row's tied threshold); the Brier score; the ten calibration bins with counts and gaps before and after the
Platt mapping; precision, recall and lift at the capacities of the scenarios; the margins against their declared values, labelled met BY CONSTRUCTION (D-G02); the hypothetical expected
intervention value of the 97-account queue (ROUND once, at the end) and a what-if of the same; the golden four-account answer; the censored share and mature churn by cutoff; the
population stability index of usage_events_30d by region from the test table, floored bins counted; and the purge arithmetic. Controls!B3 switches the model whose block is shown.
What the spreadsheet does not do: refit a model, rebuild a feature from events, or recompute the full-world figures (1,241 scored accounts, the aggregates of 2,400 accounts): those
are TYPED from the published files and labelled so. Excel is never run.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "data" / "sql" / "analysis")]

import python_side as ps  # noqa: E402

OUT = ROOT / "artifacts" / "retentionlab-analyst-workbook.xlsx"
RESULTS_PATH = ROOT / "artifacts" / "analysis-results.json"
MANIFEST = json.loads((ROOT / "web/public/data/manifest.json").read_text())
CODE_REVISION = "ce96e65"
DATASET_HASH = "9ebaaadab8a2f70e5f25820616391ee0aec55fbf0b121dee25460de469821b1c"
HDR = PatternFill("solid", fgColor="23303A")
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
EXCEL_EPOCH = datetime(1899, 12, 30, tzinfo=timezone.utc)
SHEET_ORDER = ["Read me", "Controls", "Checks", "Models", "Calibration", "Capacity", "Economics", "Censored tail", "Drift", "Purge", "Golden case", "Assumptions", "Scenarios",
               "Test table", "SQL results (pasted)"]
MODELS = ["random_baseline", "business_rule", "logistic_regression", "prevalence_baseline", "logistic_uncalibrated"]
PCOL = {"random_baseline": "L", "business_rule": "M", "logistic_regression": "N", "prevalence_baseline": "O", "logistic_uncalibrated": "P"}
RANKCOL = {"random_baseline": "Q", "business_rule": "R", "logistic_regression": "S", "prevalence_baseline": "T"}
APCOL = {"random_baseline": "U", "business_rule": "V", "logistic_regression": "W", "prevalence_baseline": "X", "logistic_uncalibrated": "Y"}
SECOL = {"random_baseline": "Z", "business_rule": "AA", "logistic_regression": "AB", "prevalence_baseline": "AC", "logistic_uncalibrated": "AD"}
FIRST, T_SHEET = 5, "'Test table'"
CHECKS: list[tuple[str, str, object, float, str]] = []      # (description, workbook formula, reference value, tolerance, source)


def serial(s: str | None):
    return None if s is None else (ps.ts(s) - EXCEL_EPOCH).total_seconds() / 86400


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def put(ws, rows: list[list], first_row: int, blue: bool = True, fmt: dict[int, str] | None = None):
    for r, row in enumerate(rows, start=first_row):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if blue and v is not None and not (isinstance(v, str) and v.startswith("=")):
                cell.font = BLUE
            if fmt and c in fmt and v is not None:
                cell.number_format = fmt[c]


def fl(x):
    return None if x is None or x == "" else float(x)


def check(desc: str, formula: str, ref, tol: float, source: str) -> None:
    CHECKS.append((desc, formula, ref, tol, source))


class Data:
    def __init__(self) -> None:
        self.res = json.loads(RESULTS_PATH.read_text())
        self.T = ps.Test()
        self.C = ps.Cohort()
        self.N = len(self.T.rows)
        self.LAST = FIRST + self.N - 1
        latest = max(self.C.cutoff[p["cutoff_id"]]["pa"] for p in self.C.pred)
        self.queue = sorted(((self.C.cutoff[p["cutoff_id"]]["account_id"], p["probability"]) for p in self.C.pred
                             if p["model_name"] == "logistic_regression" and self.C.cutoff[p["cutoff_id"]]["pa"] == latest), key=lambda x: x[0])
        self.scen = sorted(self.C.scen, key=lambda s: s["capacity"])
        self.agg = json.loads((ps.DATA / "cutoff_aggregates.json").read_text())
        self.cal = json.loads((ps.DATA / "calibration.json").read_text())

    def rows(self, key: str) -> list[dict]:
        return self.res[key]

    def rng(self, col: str) -> str:
        return f"{T_SHEET}!${col}${FIRST}:${col}${self.LAST}"


# ---- the test table with the live per-row columns -------------------------------------------------------------------------------------------------------------------------------------
def build_test_table(wb, d: Data):
    ws = wb.create_sheet("Test table")
    title(ws, "Test table: the 1,513 mature eligible rows of the test range (full world, cutoffs 2026-02-01 .. 2026-05-01)",
          "Blue = published values (test_cohort.json). Black = formulas: tie-broken rank (rows within the 25 highest probabilities of a model only, enough for K <= 20), average-precision term, squared error, calibration bin.")
    header(ws, 4, ["row_no", "account_id", "cutoff_index", "prediction_at", "segment", "region", "usage_events_30d", "days_since_last_usage", "high_severity_support_90d", "billing_total_90d",
                   "churned", "p_random_baseline", "p_business_rule", "p_logistic (Platt)", "p_prevalence_baseline", "p_logistic (before Platt)", "rank random", "rank rule", "rank logistic",
                   "rank prevalence", "AP term random", "AP term rule", "AP term logistic", "AP term prevalence", "AP term logistic (before Platt)", "sq.err random", "sq.err rule",
                   "sq.err logistic", "sq.err prevalence", "sq.err logistic (before Platt)", "bin logistic (Platt)", "bin logistic (before Platt)"])
    L = d.LAST
    r_ = lambda col: f"${col}${FIRST}:${col}${L}"
    for i, r in enumerate(d.T.rows):
        n = FIRST + i
        vals = [i + 1, r["account_id"], r["cutoff_index"], serial(r["pa"].strftime("%Y-%m-%dT%H:%M:%SZ")), r["segment"], r["region"], r["usage_events_30d"], r["days_since_last_usage"],
                r["high_severity_support_90d"], fl(r["billing_total_90d"]), r["churned"]] + [float(r[c]) for c in ("p_random_baseline", "p_business_rule", "p_logistic_calibrated", "p_prevalence_baseline", "p_logistic_uncalibrated")]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=n, column=c, value=v)
            cell.font = BLUE
        ws.cell(row=n, column=4).number_format = "yyyy-mm-dd"
        for m, rc in RANKCOL.items():
            p = PCOL[m]
            thr = f"Models!$G${5 + MODELS.index(m)}"
            ws[f"{rc}{n}"] = (f'=IF({p}{n}>={thr},1+COUNTIF({r_(p)},">"&{p}{n})+COUNTIFS({r_(p)},{p}{n},{r_("B")},"<"&B{n})'
                              f'+COUNTIFS({r_(p)},{p}{n},{r_("B")},B{n},{r_("C")},"<"&C{n}),"")')
        for m, ac in APCOL.items():
            p = PCOL[m]
            ws[f"{ac}{n}"] = f'=IF($K{n}=1,COUNTIFS({r_(p)},">="&{p}{n},{r_("K")},1)/COUNTIF({r_(p)},">="&{p}{n}),0)'
        for m, sc in SECOL.items():
            p = PCOL[m]
            ws[f"{sc}{n}"] = f"=({p}{n}-$K{n})^2"
        ws[f"AE{n}"] = f"=MIN(INT(ROUND(N{n}*1000000,0)/100000),9)"
        ws[f"AF{n}"] = f"=MIN(INT(ROUND(P{n}*1000000,0)/100000),9)"
    for col in "LMNOP":
        for n in range(FIRST, L + 1):
            ws[f"{col}{n}"].number_format = "0.000000"
    widths(ws, {"A": 8, "B": 11, "C": 10, "D": 12, "E": 11, "F": 8})


def build_models(wb, d: Data):
    ws = wb.create_sheet("Models")
    title(ws, "Model comparison on the test population (D-03, D-07 DR-01)", "Average precision = mean over churning rows of the precision at that row's tied threshold (scikit-learn step form). Brier = mean((p - y)^2). The margins are met BY CONSTRUCTION (D-G02).")
    header(ws, 4, ["model", "test-table column", "rows", "churning rows", "prevalence", "average precision", "top-25 probability threshold (rank only inside it)", "Brier score", "mean probability",
                   "reference AP (SQL 01)", "reference Brier (SQL 01)"])
    ref = {r["model"]: r for r in d.rows("01_model_comparison__models")}
    labels = {"random_baseline": "random_baseline", "business_rule": "business_rule", "logistic_regression": "logistic_regression (Platt)", "prevalence_baseline": "prevalence_baseline",
              "logistic_uncalibrated": "logistic_regression (before Platt)"}
    K, Np = d.rng("K"), d.N
    for i, m in enumerate(MODELS):
        n = 5 + i
        p = d.rng(PCOL[m])
        ws[f"A{n}"], ws[f"B{n}"] = m, PCOL[m]
        ws[f"C{n}"] = f"=COUNT({p})"
        ws[f"D{n}"] = f"=SUM({K})"
        ws[f"E{n}"] = f"=D{n}/C{n}"
        ws[f"F{n}"] = f"=SUM({d.rng(APCOL[m])})/D{n}"
        ws[f"G{n}"] = f"=LARGE({p},25)"
        ws[f"H{n}"] = f"=SUM({d.rng(SECOL[m])})/C{n}"
        ws[f"I{n}"] = f"=AVERAGE({p})"
        ws[f"J{n}"], ws[f"K{n}"] = fl(ref[labels[m]]["average_precision"]), fl(ref[labels[m]]["brier_score"])
        ws[f"J{n}"].font = ws[f"K{n}"].font = BLUE
        check(f"Average precision of {m} = the SQL/Python figure", f"Models!F{n}", fl(ref[labels[m]]["average_precision"]), 5.0000001e-7, "analysis-results.json 01_model_comparison__models")
        check(f"Brier score of {m} = the SQL/Python figure", f"Models!H{n}", fl(ref[labels[m]]["brier_score"]), 5.0000001e-7, "analysis-results.json 01_model_comparison__models")
    ws["A11"] = "Margins of the Platt-calibrated logistic regression"
    ws["A11"].font = BOLD
    header(ws, 12, ["margin", "declared before the first run (D-G01)", "realised (6-place figures)", "status", "reference (SQL 08)"])
    mg = {r["margin"]: r for r in d.rows("08_rows_and_margins__margins")}
    rowsm = [("ap_over_rule", 0.03, "=ROUND(F7,6)-ROUND(F6,6)"), ("ap_over_random", 0.05, "=ROUND(F7,6)-ROUND(F5,6)"), ("brier_under_rule", 0.002, "=ROUND(H6,6)-ROUND(H7,6)"),
             ("brier_under_random", 0.05, "=ROUND(H5,6)-ROUND(H7,6)"), ("ap_over_prevalence_constant (reference)", None, "=ROUND(F7,6)-ROUND(F8,6)"),
             ("brier_under_prevalence_constant (reference)", None, "=ROUND(H8,6)-ROUND(H7,6)")]
    for i, (name, declared, f) in enumerate(rowsm):
        n = 13 + i
        ws[f"A{n}"] = name
        ws[f"B{n}"] = declared
        if declared is not None:
            ws[f"B{n}"].font = BLUE
        ws[f"C{n}"] = f
        ws[f"D{n}"] = f'=IF(ISNUMBER(B{n}),IF(C{n}>=B{n},"met, by construction (D-G02)","not met"),"reference, not a target")'
        ws[f"E{n}"] = fl(mg[name if name in mg else name]["realised_margin"])
        ws[f"E{n}"].font = BLUE
        check(f"Margin {name} = the SQL figure", f"Models!C{n}", fl(mg[name]["realised_margin"]), 5.0000001e-7, "analysis-results.json 08_rows_and_margins__margins")
    ws["A20"] = "A met margin is a property of the synthetic world: the generator's engagement signal was strengthened until the margins were met (manifest D-G02). The Brier margin over the random baseline is uninformative (a uniform random score cannot beat about 0.33); the honest reference is the prevalence constant."
    ws["A20"].alignment = Alignment(wrap_text=True)
    check("Models: rows = 1,513 for every model", f"SUMPRODUCT(--(Models!C5:C9={d.N}))", 5, 0, "manifest row_counts.test_cohort")
    check("Models: churning rows = 189 for every model", "SUMPRODUCT(--(Models!D5:D9=189))", 5, 0, "manifest DR-01 realised")
    widths(ws, {"A": 44, "B": 22, "C": 18, "D": 28, "E": 16, "F": 18, "G": 26, "H": 14})


def build_scenarios(wb, d: Data):
    ws = wb.create_sheet("Scenarios")
    title(ws, "Intervention scenarios (assumptions, D-03): immutable; the effect is ASSUMED, not estimated", "Capacities 5, 10, 20 were fixed on validation before any outcome (D-G08). All economics are hypothetical.")
    header(ws, 4, ["scenario_id", "capacity K", "contact_cost (USD)", "assumed_effect (share of risk removed)", "margin_value (USD per retained account)", "label"])
    put(ws, [[s["scenario_id"], s["capacity"], fl(s["contact_cost"]), fl(s["assumed_effect"]), fl(s["margin_value"]), "hypothetical"] for s in d.scen], 5)
    widths(ws, {"A": 14, "B": 12, "C": 18, "D": 34, "E": 36, "F": 14})


def build_capacity(wb, d: Data):
    ws = wb.create_sheet("Capacity")
    title(ws, "Precision, recall and risk lift at the three capacities, test population (outcomes known)", "Rank: probability descending, then account id, then the earlier cutoff. Risk lift is a ranking statistic, NOT treatment uplift. The 2026-09-01 queue has no known outcomes (Economics).")
    header(ws, 4, ["model", "scenario", "K", "churners in top K", "precision@K", "recall@K", "risk lift@K", "reference precision (SQL 03)", "reference recall", "reference lift", "tie across the cut (reference)"])
    ref = d.rows("03_capacity_economics__test_population_at_k")
    names = {"random_baseline": "random_baseline", "business_rule": "business_rule", "logistic_regression": "logistic_regression (Platt)", "prevalence_baseline": "prevalence_baseline"}
    n = 5
    first = n
    for m in ("random_baseline", "business_rule", "logistic_regression", "prevalence_baseline"):
        for j, sc in enumerate(d.scen):
            row = next(r for r in ref if r["model"] == names[m] and r["scenario_id"] == sc["scenario_id"])
            ws[f"A{n}"], ws[f"B{n}"] = m, f"=Scenarios!A{5 + j}"
            ws[f"C{n}"] = f"=Scenarios!B{5 + j}"
            ws[f"D{n}"] = f'=COUNTIFS({d.rng(RANKCOL[m])},"<="&C{n},{d.rng("K")},1)'
            ws[f"E{n}"] = f"=D{n}/C{n}"
            ws[f"F{n}"] = f"=D{n}/Models!$D${5 + MODELS.index(m)}"
            ws[f"G{n}"] = f"=E{n}/Models!$E${5 + MODELS.index(m)}"
            ws[f"H{n}"], ws[f"I{n}"], ws[f"J{n}"], ws[f"K{n}"] = fl(row["precision_at_k"]), fl(row["recall_at_k"]), fl(row["risk_lift_at_k"]), row["tie_split_at_k"] == "true" or row["tie_split_at_k"] is True
            for c in "HIJK":
                ws[f"{c}{n}"].font = BLUE
            n += 1
    last = n - 1
    check("Capacity: precision@K of every model and capacity = the SQL figures (largest absolute difference)", f"SUMPRODUCT(MAX(ABS(Capacity!E{first}:E{last}-Capacity!H{first}:H{last})))", 0, 5.0000001e-7, "analysis-results.json 03_capacity_economics__test_population_at_k")
    check("Capacity: recall@K of every model and capacity = the SQL figures", f"SUMPRODUCT(MAX(ABS(Capacity!F{first}:F{last}-Capacity!I{first}:I{last})))", 0, 5.0000001e-7, "same")
    check("Capacity: risk lift@K of every model and capacity = the SQL figures", f"SUMPRODUCT(MAX(ABS(Capacity!G{first}:G{last}-Capacity!J{first}:J{last})))", 0, 5.0000001e-7, "same")
    ws["A19"] = "Selected model (Controls!B3)"
    ws["A19"].font = BOLD
    header(ws, 20, ["scenario", "K", "churners in top K", "precision@K", "recall@K", "risk lift@K"])
    for j in range(3):
        r = 21 + j
        ws[f"A{r}"] = f"=Scenarios!A{5 + j}"
        ws[f"B{r}"] = f"=Scenarios!B{5 + j}"
        sel = lambda col: f"=INDEX({col}${first}:{col}${last},(Controls!$B$4-1)*3+{j + 1})"
        ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"] = sel("D"), sel("E"), sel("F"), sel("G")
    check("Capacity: the selected-model block equals the grid row of the selected model (precision@K, K = 5, 10, 20)", "ABS(Capacity!D21-INDEX(Capacity!H5:H16,(Controls!B4-1)*3+1))+ABS(Capacity!D22-INDEX(Capacity!H5:H16,(Controls!B4-1)*3+2))+ABS(Capacity!D23-INDEX(Capacity!H5:H16,(Controls!B4-1)*3+3))", 0, 1.5000001e-6,
          "analysis-results.json 03_capacity_economics__test_population_at_k")
    widths(ws, {"A": 30, "B": 14, "C": 8, "D": 18, "E": 14, "F": 14, "G": 14, "H": 22, "I": 18, "J": 16, "K": 22})


def build_calibration(wb, d: Data):
    ws = wb.create_sheet("Calibration")
    title(ws, "Calibration bins with their counts, and the effect of the Platt mapping (D-03, D-09 R10-08)", "bin = floor(10 p), the last closed; gap = (mean predicted - observed) x 100 pp. An empty bin has no mean, rate or gap (blank, never 0).")
    bins = d.rows("02_calibration_platt__bins")
    for k, (model, col, binc) in enumerate((("logistic_regression (Platt)", "N", "AE"), ("logistic_regression (before Platt)", "P", "AF"))):
        top = 4 + k * 14
        ws[f"A{top - 1}"] = model
        ws[f"A{top - 1}"].font = BOLD
        header(ws, top, ["bin", "lower", "upper", "rows", "churning rows", "mean predicted", "observed rate", "gap (pp)", "reference rows (SQL 02)", "reference gap (pp)", "row diff", "gap diff"])
        ref = [b for b in bins if b["model"] == model]
        for i in range(10):
            n = top + 1 + i
            ws[f"A{n}"], ws[f"B{n}"], ws[f"C{n}"] = i, i / 10, (i + 1) / 10
            ws[f"D{n}"] = f"=COUNTIF({d.rng(binc)},A{n})"
            ws[f"E{n}"] = f"=COUNTIFS({d.rng(binc)},A{n},{d.rng('K')},1)"
            ws[f"F{n}"] = f'=IF(D{n}=0,"",AVERAGEIFS({d.rng(col)},{d.rng(binc)},A{n}))'
            ws[f"G{n}"] = f'=IF(D{n}=0,"",E{n}/D{n})'
            ws[f"H{n}"] = f'=IF(D{n}=0,"",(F{n}-G{n})*100)'
            ws[f"I{n}"], ws[f"J{n}"] = int(ref[i]["rows_in_bin"]), fl(ref[i]["gap_pp"])
            ws[f"I{n}"].font = ws[f"J{n}"].font = BLUE
            ws[f"K{n}"] = f"=ABS(D{n}-I{n})"
            ws[f"L{n}"] = f'=IF(D{n}=0,0,ABS(H{n}-J{n}))'
            for c in "BC":
                ws[f"{c}{n}"].font = BLUE
        check(f"Calibration bins of {model}: row counts = the SQL counts (sum of differences)", f"SUM(Calibration!K{top + 1}:K{top + 10})", 0, 0, "analysis-results.json 02_calibration_platt__bins")
        check(f"Calibration bins of {model}: gaps of the occupied bins = the SQL gaps (each within the 2-place rounding)", f"MAX(Calibration!L{top + 1}:L{top + 10})", 0, 0.0050001, "same")
        check(f"Calibration bins of {model}: all 1,513 rows fall in a bin", f"SUM(Calibration!D{top + 1}:D{top + 10})", d.N, 0, "test table")
    check("Calibration: six of the ten bins of the Platt-calibrated model hold no rows", f"COUNTIF(Calibration!D5:D14,0)", 6, 0, "analysis-results.json 02_calibration_platt__bins")
    pe = d.rows("02_calibration_platt__platt_effect")[0]
    ws["A34"] = "Platt effect on the test rows"
    ws["A34"].font = BOLD
    header(ws, 35, ["figure", "workbook", "reference (SQL 02 / calibration.json)"])
    items = [("Brier before Platt", "=Models!H9", fl(pe["brier_before"]), 5.0000001e-7), ("Brier after Platt", "=Models!H7", fl(pe["brier_after"]), 5.0000001e-7),
             ("Brier change", "=ROUND(Models!H7,6)-ROUND(Models!H9,6)", fl(pe["brier_change"]), 5.0000001e-7),
             ("Overall gap before (pp)", f"=(Models!I9-Models!E7)*100", fl(pe["overall_gap_pp_before"]), 0.0050001), ("Overall gap after (pp)", f"=(Models!I7-Models!E7)*100", fl(pe["overall_gap_pp_after"]), 0.0050001),
             ("Platt made the Brier worse on test (TRUE/FALSE)", "=ROUND(Models!H7,6)>ROUND(Models!H9,6)", True, 0)]
    for i, (name, f, refv, tol) in enumerate(items):
        n = 36 + i
        ws[f"A{n}"], ws[f"B{n}"], ws[f"C{n}"] = name, f, refv
        ws[f"C{n}"].font = BLUE
        check(f"Platt effect: {name}", f"Calibration!B{n}", refv, tol, "analysis-results.json 02_calibration_platt__platt_effect")
    ws["A43"] = "The mapping was fitted on 289 retained validation rows, all from one cutoff (2025-10-01), because the label-window purge removes the later validation cutoffs; it made the model slightly worse on the four test cutoffs."
    ws["A43"].alignment = Alignment(wrap_text=True)
    # gaps by segment and region
    ws["A46"] = "Calibration gap by segment and region (Platt-calibrated model)"
    ws["A46"].font = BOLD
    header(ws, 47, ["dimension", "group", "rows", "churning rows", "mean predicted", "observed rate", "gap (pp)", "reference gap (pp)", "diff"])
    refs = d.rows("10_segment_region_calibration__by_segment") + d.rows("10_segment_region_calibration__by_region")
    for i, r in enumerate(refs):
        n = 48 + i
        col = "E" if r["dimension"] == "segment" else "F"
        ws[f"A{n}"], ws[f"B{n}"] = r["dimension"], r["grp"]
        ws[f"C{n}"] = f"=COUNTIF({d.rng(col)},B{n})"
        ws[f"D{n}"] = f"=COUNTIFS({d.rng(col)},B{n},{d.rng('K')},1)"
        ws[f"E{n}"] = f"=AVERAGEIFS({d.rng('N')},{d.rng(col)},B{n})"
        ws[f"F{n}"] = f"=D{n}/C{n}"
        ws[f"G{n}"] = f"=(E{n}-F{n})*100"
        ws[f"H{n}"] = fl(r["gap_pp"])
        ws[f"H{n}"].font = BLUE
        ws[f"I{n}"] = f"=ABS(G{n}-H{n})"
    last = 48 + len(refs) - 1
    check("Calibration gaps by segment and region = the SQL gaps (largest difference, within the 2-place rounding)", f"MAX(Calibration!I48:I{last})", 0, 0.0050001, "analysis-results.json 10_segment_region_calibration")
    widths(ws, {"A": 46, "B": 16, "C": 12, "D": 12, "E": 14, "F": 16, "G": 14, "H": 16, "I": 18, "J": 16})


def build_economics(wb, d: Data):
    ws = wb.create_sheet("Economics")
    title(ws, "Hypothetical expected intervention value (D-03, D-09 R10-04): assumption-based, by population",
          "Expected value = sum over the top-K accounts of (probability x assumed effect x margin - contact cost), rounded ONCE at the end. A ranking by risk is not a treatment effect; no randomised experiment supports the assumed effect.")
    header(ws, 4, ["account_id", "probability (logistic, Platt)", "rank (p desc, account id asc)"])
    q = d.queue
    for i, (acc, p) in enumerate(q):
        n = 5 + i
        ws[f"A{n}"], ws[f"B{n}"] = acc, float(p)
        ws[f"A{n}"].font = ws[f"B{n}"].font = BLUE
        ws[f"C{n}"] = f'=1+COUNTIF($B$5:$B${4 + len(q)},">"&B{n})+COUNTIFS($B$5:$B${4 + len(q)},B{n},$A$5:$A${4 + len(q)},"<"&A{n})'
    lq = 4 + len(q)
    header(ws, 4, ["scenario", "capacity K", "contact budget (K x cost)", "COHORT expected value (97 scored accounts)", "published cohort value (manifest)", "WORLD expected value (1,241 scored accounts, published)", "world scored accounts", "label", "population note"], col=5)
    econ = MANIFEST["published_cohort"]["economics"]
    world = MANIFEST["economics_full_world"]["scenarios"]
    for j, s in enumerate(d.scen):
        n = 5 + j
        ws[f"E{n}"] = f"=Scenarios!A{5 + j}"
        ws[f"F{n}"] = f"=Scenarios!B{5 + j}"
        ws[f"G{n}"] = f"=ROUND(F{n}*Scenarios!C{5 + j},2)"
        ws[f"H{n}"] = f"=ROUND(SUMPRODUCT(($C$5:$C${lq}<=F{n})*($B$5:$B${lq}*Scenarios!D{5 + j}*Scenarios!E{5 + j}-Scenarios!C{5 + j})),2)"
        ws[f"I{n}"] = fl(econ[s["scenario_id"]]["expected_value"])
        ws[f"J{n}"] = fl(world[s["scenario_id"]]["expected_value"])
        ws[f"K{n}"] = world[s["scenario_id"]]["scored_accounts"]
        ws[f"L{n}"] = "hypothetical"
        ws[f"M{n}"] = "cohort: computed here; world: TYPED from cutoff_aggregates.json (not recomputable from the cohort)"
        for c in "IJK":
            ws[f"{c}{n}"].font = BLUE
        check(f"Economics {s['scenario_id']}: the cohort expected value (97 scored accounts) = the published cohort figure", f"Economics!H{n}", fl(econ[s["scenario_id"]]["expected_value"]), 0.0000001, "manifest published_cohort.economics; analysis-results.json 03_capacity_economics__cohort_economics")
        check(f"Economics {s['scenario_id']}: the contact budget = K x contact cost", f"Economics!G{n}", fl(econ[s["scenario_id"]]["contact_budget"]), 0.0000001, "manifest published_cohort.economics")
    check("Economics: the queue holds 97 scored accounts", f"COUNT(Economics!B5:B{lq})", len(q), 0, "analysis-results.json 03_capacity_economics__cohort_economics")
    ws["E10"] = "What-if (Controls!B11:B14)"
    ws["E10"].font = BOLD
    ws["E11"], ws["F11"] = "capacity K", "=Controls!B11"
    ws["E12"], ws["F12"] = "contact budget", "=ROUND(Controls!B11*Controls!B12,2)"
    ws["E13"], ws["F13"] = "HYPOTHETICAL expected value, 97 cohort accounts", f"=ROUND(SUMPRODUCT(($C$5:$C${lq}<=Controls!B11)*($B$5:$B${lq}*Controls!B13*Controls!B14-Controls!B12)),2)"
    ws["E14"] = "A what-if is an assumption, not an estimate. The same K and economics on the whole world cannot be recomputed here."
    ws["E16"] = "Outcomes: every row of the 2026-09-01 cutoff is censored at the as-of (label window closes 2026-12-14), so precision and recall of these queues are unknown, not zero."
    check("Economics: the what-if with the default Controls (K 5, cost 8, effect 0.3, margin 200) reproduces scenario sc-001", "IF(AND(Controls!B11=5,Controls!B12=8,Controls!B13=0.3,Controls!B14=200),Economics!F13,Economics!H5)", fl(econ["sc-001"]["expected_value"]), 0.0000001, "manifest published_cohort.economics")
    widths(ws, {"A": 12, "B": 20, "C": 24, "E": 44, "F": 14, "G": 20, "H": 28, "I": 24, "J": 30, "K": 14, "L": 14, "M": 70})


def build_censored(wb, d: Data):
    ws = wb.create_sheet("Censored tail")
    title(ws, "The censored tail by cutoff (D-01, D-07 DR-05): rows whose label window has not closed are excluded, never counted as retained",
          "World columns are TYPED from cutoff_aggregates.json (2,400 accounts); the cohort columns are TYPED from the loaded 200-account cohort (SQL 04). Shares and rates are formulas.")
    header(ws, 4, ["cutoff", "range", "WORLD eligible", "world mature", "world censored", "world mature churned", "world censored share", "world mature churn rate", "published rate if censored rows were counted as non-churn",
                   "implied cancellations seen so far", "COHORT eligible", "cohort censored", "cohort cancellations seen so far", "cohort censored share", "distance of the implied count from a whole number"])
    naive = MANIFEST["designed_relationships"][4]["realised"]["churn_rate_if_censored_were_non_churn"]
    coh = {r["prediction_at"]: r for r in d.rows("04_censored_tail__cohort_by_cutoff")}
    world = d.rows("04_censored_tail__world_by_cutoff")
    for i, w in enumerate(world):
        n = 5 + i
        key = w["prediction_at"]
        iso = key.replace("+00:00", "Z")
        c = coh[key]
        ws[f"A{n}"], ws[f"B{n}"] = serial(iso), w["split_range"]
        ws[f"A{n}"].number_format = "yyyy-mm-dd"
        ws[f"C{n}"], ws[f"D{n}"], ws[f"E{n}"], ws[f"F{n}"] = int(w["eligible"]), int(w["mature"]), int(w["censored"]), int(w["churned"])
        ws[f"G{n}"] = f"=E{n}/C{n}"
        ws[f"H{n}"] = f'=IF(D{n}=0,"",F{n}/D{n})'
        ws[f"I{n}"] = fl(naive.get(iso))
        ws[f"J{n}"] = f'=IF(ISNUMBER(I{n}),ROUND(I{n}*C{n},0),"")'
        ws[f"K{n}"], ws[f"L{n}"], ws[f"M{n}"] = int(c["eligible"]), int(c["censored"]), int(c["cancellations_seen_so_far_in_censored_rows"])
        ws[f"N{n}"] = f"=L{n}/K{n}"
        ws[f"O{n}"] = f"=IF(ISNUMBER(I{n}),ABS(J{n}-I{n}*C{n}),0)"
        for col in "ABCDEFIKLM":
            if ws[f"{col}{n}"].value is not None:
                ws[f"{col}{n}"].font = BLUE
    last = 4 + len(world)
    t = last + 1
    ws[f"A{t}"] = "Total"
    ws[f"A{t}"].font = BOLD
    for col in "CDEF":
        ws[f"{col}{t}"] = f"=SUM({col}5:{col}{last})"
    ws[f"G{t}"] = f"=E{t}/C{t}"
    ws[f"H{t}"] = f"=F{t}/D{t}"
    for col in "KLM":
        ws[f"{col}{t}"] = f"=SUM({col}5:{col}{last})"
    ws[f"N{t}"] = f"=L{t}/K{t}"
    ws[f"A{t + 2}"] = "If censored rows were counted as retained, the last three cutoffs would show 9.4 %, 6.9 % and 3.9 % churn (world) instead of an unknown: the rate collapses towards zero the nearer the cutoff is to the as-of."
    ws[f"A{t + 2}"].alignment = Alignment(wrap_text=True)
    ws0 = d.rows("04_censored_tail__world_summary")[0]
    check("Censored tail: eligible rows (world) = 9,039", f"'Censored tail'!C{t}", int(ws0["eligible_rows"]), 0, "analysis-results.json 04_censored_tail__world_summary; manifest population.feature_rows")
    check("Censored tail: censored rows (world) = 3,594", f"'Censored tail'!E{t}", int(ws0["censored_rows"]), 0, "same; manifest DR-05")
    check("Censored tail: censored share (world) = 0.397610", f"'Censored tail'!G{t}", fl(ws0["censored_share"]), 5.0000001e-7, "same")
    check("Censored tail: mature churn rate (world) = 0.104500", f"'Censored tail'!H{t}", fl(ws0["mature_churn_rate"]), 5.0000001e-7, "same")
    check("Censored tail: implied cancellations seen so far are whole numbers (109, 83, 48)", f"MAX('Censored tail'!O5:O{last})", 0, 0.001, "manifest DR-05")
    check("Censored tail: the censored cutoffs are exactly the last three (rows with censored > 0)", f"COUNTIF('Censored tail'!E5:E{last},\">0\")", 3, 0, "manifest DR-05 immature_cutoffs")
    check("Censored tail: cohort censored rows = 286 (94 + 95 + 97)", f"'Censored tail'!L{t}", 286, 0, "analysis-results.json 04_censored_tail__cohort_by_cutoff")
    widths(ws, {"A": 12, "B": 12, "C": 12, "D": 12, "E": 12, "F": 14, "G": 14, "H": 14, "I": 28, "J": 18, "K": 12, "L": 12, "M": 18, "N": 14})


def build_drift(wb, d: Data):
    ws = wb.create_sheet("Drift")
    title(ws, "Population drift of usage_events_30d by region (D-03 P10-MET-11, D-P10-10)", "PSI = sum (actual - expected) x LN(actual / expected) over bins [0,1), [1,3), [3,inf); a zero proportion is replaced by 0.0001 (floored, counted). The 0.2 threshold is a diagnostic, not proof. Reference = the 2026-02-01 cohort.")
    header(ws, 4, ["group", "cutoff index", "cur 0", "cur 1-2", "cur 3+", "ref 0", "ref 1-2", "ref 3+", "current rows", "reference rows", "floored bins", "PSI (workbook)", "exceeds 0.2", "PSI published (full world)", "diff"])
    pub = {(r["prediction_at"], r["grp"]): r for r in d.rows("07_drift_by_region__psi_test_range")}
    agg = {(x["prediction_at"], g): x[g] for x in d.agg["drift"]["by_cutoff"] for g in ("apac", "other_regions", "all_accounts")}
    reg, cut, usage = d.rng("F"), d.rng("C"), d.rng("G")

    def cnt(group: str, idx_expr: str, lo: str, hi: str | None) -> str:
        rcrit = {"apac": f',{reg},"apac"', "other_regions": f',{reg},"<>apac"', "all_accounts": ""}[group]
        crit = f'{usage},">={lo}"' + (f',{usage},"<{hi}"' if hi else "")
        return f"COUNTIFS({cut},{idx_expr}{rcrit},{crit})"
    n = 5
    first = n
    for group in ("apac", "other_regions", "all_accounts"):
        for idx in range(4):
            pa = d.T.cutoffs[idx].strftime("%Y-%m-%dT%H:%M:%SZ")
            ws[f"A{n}"], ws[f"B{n}"] = group, idx
            for k, (lo, hi) in enumerate((("0", "1"), ("1", "3"), ("3", None))):
                ws[f"{'CDE'[k]}{n}"] = "=" + cnt(group, f"$B{n}", lo, hi)
                ws[f"{'FGH'[k]}{n}"] = "=" + cnt(group, "0", lo, hi)
            ws[f"I{n}"], ws[f"J{n}"] = f"=SUM(C{n}:E{n})", f"=SUM(F{n}:H{n})"
            ws[f"K{n}"] = f"=SUMPRODUCT(--((C{n}:E{n}=0)+(F{n}:H{n}=0)>0))"
            term = lambda c, r: f"((IF({c}{n}=0,0.0001,{c}{n}/$I{n})-IF({r}{n}=0,0.0001,{r}{n}/$J{n}))*LN(IF({c}{n}=0,0.0001,{c}{n}/$I{n})/IF({r}{n}=0,0.0001,{r}{n}/$J{n})))"
            ws[f"L{n}"] = "=" + "+".join(term(c, r) for c, r in (("C", "F"), ("D", "G"), ("E", "H")))
            ws[f"M{n}"] = f"=L{n}>0.2"
            ws[f"N{n}"] = fl(agg[(pa, group)]["psi"])
            ws[f"N{n}"].font = BLUE
            ws[f"O{n}"] = f"=ABS(L{n}-N{n})"
            n += 1
    last = n - 1
    check("Drift: PSI of apac, other regions and all accounts at the four test cutoffs, recomputed in the workbook from the test table = the published series (largest difference)", f"MAX(Drift!O{first}:O{last})", 0, 5.0000001e-7, "cutoff_aggregates.json drift.by_cutoff; analysis-results.json 07_drift_by_region")
    ws["A19"] = "Exceedances in the test-range recomputation (apac)"
    ws["B19"] = f'=COUNTIFS(A{first}:A{last},"apac",M{first}:M{last},TRUE)'
    check("Drift: apac exceeds 0.2 only at 2026-05-01 among the test cutoffs", "Drift!B19", 1, 0, "analysis-results.json 07_drift_by_region__psi_test_range")
    check("Drift: the whole population and the other regions never exceed 0.2 at the test cutoffs", f'COUNTIFS(A{first}:A{last},"<>apac",M{first}:M{last},TRUE)', 0, 0, "same")
    ws["A21"] = "Published full-world apac series (TYPED from cutoff_aggregates.json; 21 cutoffs)"
    ws["A21"].font = BOLD
    header(ws, 22, ["cutoff", "apac PSI", "floored bins", "exceeds 0.2 (formula)", "all accounts PSI", "other regions PSI", "false-positive reading: exceeds with a floored bin"])
    ser = d.rows("07_drift_by_region__published_series_apac")
    for i, r in enumerate(ser):
        m = 23 + i
        ws[f"A{m}"] = serial(r["prediction_at"].replace("+00:00", "Z"))
        ws[f"A{m}"].number_format = "yyyy-mm-dd"
        ws[f"B{m}"], ws[f"C{m}"], ws[f"E{m}"], ws[f"F{m}"] = fl(r["psi"]), int(r["floored_bins"]), fl(r["all_accounts_psi"]), fl(r["other_regions_psi"])
        for c in "ABCEF":
            ws[f"{c}{m}"].font = BLUE
        ws[f"D{m}"] = f"=B{m}>0.2"
        ws[f"G{m}"] = f"=AND(D{m},C{m}>0)"
    e = 22 + len(ser)
    ws[f"A{e + 2}"] = "Exceedances before the change (comparison cutoffs before 2026-03-31, the reference month excluded)"
    ws[f"B{e + 2}"] = f"=COUNTIFS(A23:A{e},\"<\"&DATE(2026,3,31),D23:D{e},TRUE)"
    ws[f"A{e + 3}"] = "... of which with a floored bin (false positives of the 0.0001 floor)"
    ws[f"B{e + 3}"] = f"=COUNTIFS(A23:A{e},\"<\"&DATE(2026,3,31),G23:G{e},TRUE)"
    ws[f"A{e + 4}"] = "Exceedances from 2026-05-01 (the shift)"
    ws[f"B{e + 4}"] = f"=COUNTIFS(A23:A{e},\">=\"&DATE(2026,5,1),D23:D{e},TRUE)"
    ws[f"A{e + 5}"] = "Largest whole-population PSI in the series"
    ws[f"B{e + 5}"] = f"=MAX(E23:E{e})"
    sm = d.rows("07_drift_by_region__published_summary")[0]
    check("Drift: five apac exceedances before the change (13 months before the reference, 14 comparison cutoffs)", f"Drift!B{e + 2}", int(sm["exceedances_before_the_change"]), 0, "analysis-results.json 07_drift_by_region__published_summary")
    check("Drift: four of the five carry a floored bin", f"Drift!B{e + 3}", int(sm["exceedances_before_the_change_with_a_floored_bin"]), 0, "same")
    check("Drift: five exceedances from 2026-05-01", f"Drift!B{e + 4}", int(sm["exceedances_from_the_change"]), 0, "same")
    check("Drift: the largest whole-population PSI is 0.043730", f"Drift!B{e + 5}", fl(sm["all_accounts_max_psi"]), 5.0000001e-7, "same")
    widths(ws, {"A": 56, "B": 14, "C": 10, "D": 12, "E": 12, "F": 12, "G": 40, "L": 14, "N": 18})


def build_purge(wb, d: Data):
    ws = wb.create_sheet("Purge")
    title(ws, "Rows dropped by the chronological split (D-02): label-window purge, then group purge", "Counts are TYPED (cohort: the loaded split assignments, recomputed in SQL 06; world: split_summary.json). Retained and dropped are formulas; the 104-day embargo = horizon 90 + grace 14.")
    header(ws, 4, ["population", "range", "cutoff rows", "retained", "LABEL_WINDOW", "GROUP_TEST", "GROUP_VALIDATION", "dropped (formula)", "adds up (formula)"])
    rows = [("cohort_200", r) for r in d.rows("06_purge_counts__cohort_stored")] + [("world_2400", r) for r in d.rows("06_purge_counts__world_published")]
    for i, (pop, r) in enumerate(rows):
        n = 5 + i
        ws[f"A{n}"], ws[f"B{n}"] = pop, r["split_range"]
        for c, k in zip("CDEFG", ("cutoff_rows", "retained", "label_window", "group_test", "group_validation")):
            ws[f"{c}{n}"] = int(r[k])
            ws[f"{c}{n}"].font = BLUE
        ws[f"H{n}"] = f"=C{n}-D{n}"
        ws[f"I{n}"] = f"=D{n}+E{n}+F{n}+G{n}=C{n}"
    last = 4 + len(rows)
    check("Purge: retained + dropped = rows in every range of both populations", f"COUNTIF(Purge!I5:I{last},TRUE)", len(rows), 0, "analysis-results.json 06_purge_counts")
    check("Purge: whole world, train range loses 879 + 124 + 52 = 1,055 rows", "SUM(Purge!E9:G9)", 1055, 0, "split_summary.json population.train")
    check("Purge: whole world, validation range loses 1,291 + 13 = 1,304 rows", "SUM(Purge!E10:G10)", 1304, 0, "split_summary.json population.validation")
    check("Purge: no row of the test range or after it is dropped (both populations)", "SUM(Purge!E7:G8)+SUM(Purge!E11:G12)", 0, 0, "analysis-results.json 06_purge_counts")
    ws["A20"] = "The test range keeps all its rows; account roles are assigned by a hash of the id (D-G03), which shapes the test population and is disclosed."
    widths(ws, {"A": 14, "B": 12, "C": 12, "D": 12, "E": 14, "F": 12, "G": 18, "H": 16, "I": 16})


def build_golden(wb):
    ws = wb.create_sheet("Golden case")
    title(ws, "The golden known answer (P10-GOLD-01): four accounts, risks .8 .6 .4 .1, outcomes 1 0 1 0, capacity 2", "Typed from tests/fixtures/golden/CALCULATION.md; the formulas rebuild every figure; the hand answers are typed in column G.")
    header(ws, 4, ["account", "risk", "outcome", "rank", "squared error", "AP term"])
    data = [(1, 0.8, 1), (2, 0.6, 0), (3, 0.4, 1), (4, 0.1, 0)]
    for i, (a, p, y) in enumerate(data):
        n = 5 + i
        ws[f"A{n}"], ws[f"B{n}"], ws[f"C{n}"] = a, p, y
        for c in "ABC":
            ws[f"{c}{n}"].font = BLUE
        ws[f"D{n}"] = f'=1+COUNTIF($B$5:$B$8,">"&B{n})+COUNTIFS($B$5:$B$8,B{n},$A$5:$A$8,"<"&A{n})'
        ws[f"E{n}"] = f"=(B{n}-C{n})^2"
        ws[f"F{n}"] = f'=IF(C{n}=1,COUNTIFS($B$5:$B$8,">="&B{n},$C$5:$C$8,1)/COUNTIF($B$5:$B$8,">="&B{n}),0)'
    ws["A11"] = "scenario sc-001 (assumptions)"
    for c, v in zip("BCDE", ("capacity K", "contact cost (USD)", "assumed effect", "margin (USD)")):
        ws[f"{c}11"] = v
        ws[f"{c}11"].font = BOLD
    for c, v in zip("BCDE", (2, 5.0, 0.1, 131.75)):
        ws[f"{c}12"] = v
        ws[f"{c}12"].font = BLUE
    header(ws, 14, ["figure", "workbook", "hand answer (CALCULATION.md)", "equal"])
    figs = [("churn rate", "=SUM(C5:C8)/COUNT(C5:C8)", 0.5), ("precision at K=2", "=SUMPRODUCT((D5:D8<=B12)*C5:C8)/B12", 0.5), ("recall at K=2", "=SUMPRODUCT((D5:D8<=B12)*C5:C8)/SUM(C5:C8)", 0.5),
            ("risk lift at K=2", "=(SUMPRODUCT((D5:D8<=B12)*C5:C8)/B12)/(SUM(C5:C8)/COUNT(C5:C8))", 1.0), ("Brier score", "=AVERAGE(E5:E8)", 0.1925), ("average precision (5/6)", "=SUM(F5:F8)/SUM(C5:C8)", 5 / 6),
            ("contact budget", "=ROUND(B12*C12,2)", 10.0), ("expected intervention value (hypothetical)", "=ROUND(SUMPRODUCT((D5:D8<=B12)*(B5:B8*D12*E12-C12)),2)", 8.45)]
    for i, (name, f, hand) in enumerate(figs):
        n = 15 + i
        ws[f"A{n}"], ws[f"B{n}"], ws[f"C{n}"] = name, f, hand
        ws[f"C{n}"].font = BLUE
        ws[f"D{n}"] = f"=ABS(B{n}-C{n})<=0.0000005"
        check(f"Golden case: {name} = the hand answer", f"'Golden case'!B{n}", hand, 5.0000001e-7, "tests/fixtures/golden/CALCULATION.md")
    ws["A24"] = "8.445 rounds to 8.45 (ROUND_HALF_UP, once at the end); rounding each account first would give 8.44 or 8.45 by accident, which is why the fixture is built this way."
    widths(ws, {"A": 44, "B": 14, "C": 28, "D": 14, "E": 16, "F": 12})


def build_controls(wb):
    ws = wb.create_sheet("Controls")
    title(ws, "Controls", "Change B3 (the model whose block Capacity shows), B5 (a scenario) and the what-if inputs B11:B14. Blue cells are inputs.")
    ws["A3"], ws["B3"] = "Selected model", "logistic_regression"
    ws["B3"].font = BLUE
    dv = DataValidation(type="list", formula1='"random_baseline,business_rule,logistic_regression,prevalence_baseline"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    ws["A4"], ws["B4"] = "Model index (1..4)", "=MATCH(B3,Models!A5:A8,0)"
    ws["A5"], ws["B5"] = "Selected scenario", "sc-001"
    ws["B5"].font = BLUE
    dv2 = DataValidation(type="list", formula1='"sc-001,sc-002,sc-003"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add("B5")
    ws["A6"], ws["B6"] = "Scenario capacity K", "=INDEX(Scenarios!B5:B7,MATCH(B5,Scenarios!A5:A7,0))"
    ws["A7"], ws["B7"] = "Scenario contact cost (USD)", "=INDEX(Scenarios!C5:C7,MATCH(B5,Scenarios!A5:A7,0))"
    ws["A8"], ws["B8"] = "Scenario assumed effect", "=INDEX(Scenarios!D5:D7,MATCH(B5,Scenarios!A5:A7,0))"
    ws["A9"], ws["B9"] = "Scenario margin (USD per retained account)", "=INDEX(Scenarios!E5:E7,MATCH(B5,Scenarios!A5:A7,0))"
    ws["A10"] = "What-if (hypothetical, on the 97 cohort accounts)"
    ws["A10"].font = BOLD
    for r, (a, v) in enumerate((("What-if capacity K", 5), ("What-if contact cost (USD)", 8.0), ("What-if assumed effect (share of risk removed, 0 to 1)", 0.3), ("What-if margin (USD per retained account)", 200.0)), start=11):
        ws[f"A{r}"], ws[f"B{r}"] = a, v
        ws[f"B{r}"].font = BLUE
    dv3 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=False)
    ws.add_data_validation(dv3)
    dv3.add("B13")
    ws["A16"], ws["B16"] = "As-of (UTC)", serial("2026-10-05T00:00:00Z")
    ws["B16"].font = BLUE
    ws["B16"].number_format = "yyyy-mm-dd hh:mm"
    ws["A18"], ws["B18"] = "All checks pass (Checks!B2)", "=Checks!B2"
    ws["A18"].font = BOLD
    widths(ws, {"A": 56, "B": 24})


def build_readme(wb):
    ws = wb.create_sheet("Read me")
    title(ws, "RetentionLab analyst workbook (P10-ART-05)")
    rows = [
        ("What it is", f"A formula-driven recomputation of the RetentionLab evaluation over the published synthetic snapshot rl-20261026-20261026 (dataset hash {DATASET_HASH}, code revision {CODE_REVISION}, as-of 2026-10-05T00:00:00Z): the 1,513-row held-out test table of the full 2,400-account world, the 200-account cohort and its 97-account 2026-09-01 queue."),
        ("The decision", "Which accounts deserve limited retention attention under uncertain intervention value? The workbook shows how well four models rank and score the held-out rows (the margin of the logistic model over its baselines is met by construction), how well calibrated the probabilities are, what a capacity of 5, 10 or 20 gets, what the censored tail hides, where the population drifts, and what a HYPOTHETICAL contact programme would be worth under stated assumptions."),
        ("Synthetic data", "Everything is synthetic (a seeded generator). Simulated = what this one draw shows; hypothetical = an assumption-based value (the expected intervention value); no retention improvement is claimed and the ranking by risk is not a treatment effect."),
        ("How to use it", "Open the .xlsx in any spreadsheet application. Controls!B3 selects the model, B5 a scenario, B11:B14 a what-if; Checks!B2 is TRUE when every figure agrees with the SQL/Python results and the golden hand answer. Nothing needs installing to read the cached values."),
        ("Recalculation status", "openpyxl wrote formulas only; LibreOffice (headless) recalculated the delivered file. Microsoft Excel was NOT run: nothing here was opened or recalculated by Excel (docs/workbook.md)."),
        ("Sheets", "Controls, Checks, Models, Calibration, Capacity, Economics, Censored tail, Drift, Purge, Golden case, Assumptions, Scenarios, Test table (the 1,513 rows with the live per-row columns), SQL results (pasted)."),
        ("Colour", "Blue = typed input or pasted value; black = formula."),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 24, "B": 140})


ASSUMPTIONS = [
    ("A-01", "Average precision is the scikit-learn step form: tied probabilities are one threshold, so the figure does not depend on how account ids break a tie. The workbook computes it as the mean, over churning rows, of tp(p >= row) / n(p >= row)."),
    ("A-02", "Ranking for precision, recall and lift at K: probability descending, then account id ascending, then the earlier cutoff (the union rule of D-08). The rank is evaluated only for rows inside the 25 highest probabilities of a model, which is exact for K <= 25; a tie that crosses the cut is flagged in the Capacity reference column."),
    ("A-03", "The test table is the whole test range of the full world: 1,513 rows from 442 accounts (189 churning rows from 95 accounts). Rows are correlated across overlapping cutoffs; no interval is given, because none would be honest."),
    ("A-04", "Calibration bins: index = floor(10 p) with p = 1 in the last bin; computed with integer arithmetic on p x 10^6. An empty bin shows blank, not 0. Gaps in percentage points."),
    ("A-05", "Economics are HYPOTHETICAL and assumption-based: value = sum over the top K of (probability x assumed effect x margin - contact cost), rounded once at the end (ROUND, half away from zero). The cohort value is computed from the 97 scored accounts of the 2026-09-01 cutoff; the world value (1,241 scored accounts) is typed from cutoff_aggregates.json and cannot be recomputed from the cohort. No treatment effect is estimated."),
    ("A-06", "The censored tail: a row is mature only when its label window (horizon 90 days + grace 14 days) has closed at the as-of 2026-10-05; the last three cutoffs are fully censored. Censored rows are excluded from every evaluation and are never counted as retained."),
    ("A-07", "PSI of usage_events_30d against the 2026-02-01 cohort on the bins [0,1), [1,3), [3,inf); a zero proportion is replaced by 0.0001 without renormalising (floored, counted); LN is the natural logarithm in floating point here (the domain uses 60-digit decimals); the threshold 0.2 is a diagnostic."),
    ("W-01", "Simulated, constructed and hypothetical results are kept apart: every figure is a property of one draw of a generated world; the margins of the logistic model over the baselines were met by construction (D-G02); the economics are assumptions."),
    ("W-02", "TYPED from the published files, not replayed by formulas: the test-table values and probabilities, the 97 queue probabilities, the scenarios, the per-cutoff counts of the whole world, the published drift series, the purge counts, and the reference columns of every check (pasted from artifacts/analysis-results.json)."),
    ("W-03", "Recalculation: openpyxl writes formulas only; the delivered file was recalculated by LibreOffice headless and carries its cached values; Excel was NOT run. fullCalcOnLoad is set so Excel recalculates on open. Functions used: COUNTIF, COUNTIFS, AVERAGEIFS, SUMPRODUCT, SUM, INDEX, MATCH, LARGE, LN, IF, AND, ABS, MAX, ROUND, INT, MIN (no _xlfn functions)."),
    ("W-04", "The model selector changes which model Capacity's block shows; every model is always computed and checked, so the check cell does not depend on the selector."),
    ("W-05", "The workbook has no macros, no external links and no credentials."),
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions and limits (read this first)")
    header(ws, 3, ["id", "assumption or limit"])
    for i, (k, v) in enumerate(ASSUMPTIONS, start=4):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 8, "B": 150})


def build_sql_results(wb, d: Data):
    ws = wb.create_sheet("SQL results (pasted)")
    title(ws, "Results of data/sql/analysis (pasted values from artifacts/analysis-results.json; the reconciliation rows count is in meta)")
    row = 3
    meta = d.res["meta"][0]
    ws.cell(row=row, column=1, value=f"reconciliation rows: {meta['reconciliation_rows']} (equal: {meta['reconciliation_equal']}); result sets: {meta['result_sets']}").font = BOLD
    row += 2
    for key in ("01_model_comparison__models", "02_calibration_platt__platt_effect", "03_capacity_economics__both_populations", "04_censored_tail__world_summary", "06_purge_counts__world_totals",
                "07_drift_by_region__published_summary", "08_rows_and_margins__rows_and_accounts", "09_golden_known_answer__known_answer"):
        rows = d.rows(key)
        ws.cell(row=row, column=1, value=key).font = BOLD
        row += 1
        cols = list(rows[0])
        header(ws, row, cols)
        for r in rows:
            row += 1
            for c, k in enumerate(cols, start=1):
                v = r[k]
                try:
                    v = float(v) if v not in ("", "true", "false") and not str(v).startswith(("20",)) and all(ch in "0123456789.-" for ch in v) else v
                except ValueError:
                    pass
                cell = ws.cell(row=row, column=c, value=v)
                cell.font = BLUE
        row += 2
    widths(ws, {"A": 34})


def build_checks(wb):
    ws = wb.create_sheet("Checks")
    title(ws, "Checks: TRUE when every comparison with the SQL/Python results and the golden hand answer holds")
    ws["A2"] = "All checks pass"
    ws["A2"].font = BOLD
    first, last = 5, 4 + len(CHECKS)
    ws["B2"] = f"=AND(F{first}:F{last})"
    header(ws, 4, ["#", "check", "workbook", "reference", "tolerance", "ok", "reference source"])
    for i, (desc, f, ref, tol, src) in enumerate(CHECKS):
        n = first + i
        ws[f"A{n}"], ws[f"B{n}"] = i + 1, desc
        ws[f"C{n}"] = "=" + f
        ws[f"D{n}"] = ref
        ws[f"D{n}"].font = BLUE
        ws[f"E{n}"] = tol
        ws[f"E{n}"].font = BLUE
        ws[f"F{n}"] = f"=C{n}=D{n}" if isinstance(ref, bool) else f"=ABS(C{n}-D{n})<=E{n}"
        ws[f"G{n}"] = src
    widths(ws, {"A": 6, "B": 110, "C": 14, "D": 14, "E": 12, "F": 8, "G": 70})


def recalc(path: Path) -> None:
    """Open with LibreOffice headless and save a recalculated copy over ``path`` (cached values included)."""
    if shutil.which("soffice") is None:
        raise SystemExit("soffice not installed; the file has formulas but no cached values")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in" / path.name
        src.parent.mkdir()
        shutil.copy(path, src)
        out = Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=1800)
        shutil.copy(out / path.name, path)
    set_full_calc_on_load(path)


def set_full_calc_on_load(path: Path) -> None:
    """LibreOffice drops fullCalcOnLoad; put it back (workbook.xml only) so Excel recalculates when it opens the file."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode("utf-8")
                assert "<calcPr " in text and "fullCalcOnLoad" not in text
                data = text.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ', 1).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build(path: Path) -> None:
    CHECKS.clear()
    d = Data()
    wb = Workbook()
    wb.remove(wb.active)
    build_readme(wb)
    build_controls(wb)
    build_scenarios(wb, d)
    build_test_table(wb, d)
    build_models(wb, d)
    build_calibration(wb, d)
    build_capacity(wb, d)
    build_economics(wb, d)
    build_censored(wb, d)
    build_drift(wb, d)
    build_purge(wb, d)
    build_golden(wb)
    build_assumptions(wb)
    build_sql_results(wb, d)
    build_checks(wb)
    wb._sheets = [wb[n] for n in SHEET_ORDER]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main(argv: list[str]) -> int:
    build(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(CHECKS)} checks)")
    if "--recalc" in argv:
        recalc(OUT)
        print("recalculated with LibreOffice headless; cached values written")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
