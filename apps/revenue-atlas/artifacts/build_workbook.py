"""Build artifacts/revenue-atlas-analyst-workbook.xlsx with openpyxl (P01-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: tests/fixtures/golden/inputs.json and expected.json (golden fixture), artifacts/analysis-results.json
(written by data/sql/analysis/run_all.py), web/public/data/manifest.json. openpyxl cannot calculate, so the file it writes
holds formulas only. With --recalc, LibreOffice (soffice) opens it headless and saves a recalculated copy over the same
path, which then also carries cached values. The workbook never hard-codes a computed total in Allocation or Bridge.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "revenue-atlas-analyst-workbook.xlsx"
GOLD = json.loads((ROOT / "tests/fixtures/golden/inputs.json").read_text())
EXPECTED = json.loads((ROOT / "tests/fixtures/golden/expected.json").read_text())
RESULTS = json.loads((ROOT / "artifacts/analysis-results.json").read_text())
MANIFEST = json.loads((ROOT / "web/public/data/manifest.json").read_text())

HDR = PatternFill("solid", fgColor="1F3A5F")
SEC = PatternFill("solid", fgColor="DCE6F1")
BLUE = Font(color="0000FF")            # typed inputs
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
MONEY = '#,##0.00'


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note
        ws["A2"].alignment = Alignment(wrap_text=False)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# --------------------------------------------------------------------------------------------- Inputs
ENTITIES = ["customers", "products", "orders", "order_lines", "discounts", "shipments", "shipment_lines",
            "returns", "cost_versions", "exchange_rates"]
POS: dict[str, dict] = {}      # entity -> {"first": row, "last": row, "col": {field: letter}}


def num(v):
    if isinstance(v, bool):
        return str(v).lower()           # LibreOffice rewrites real booleans as =TRUE(); keep the fixture's literal as text
    return float(Decimal(v)) if isinstance(v, str) and v.replace(".", "", 1).lstrip("-").isdigit() else v


def build_inputs(wb):
    ws = wb.active
    ws.title = "Inputs"
    title(ws, "Inputs: the golden fixture rows (tests/fixtures/golden/inputs.json, fixture revenue-atlas-golden-v1)",
          "SYNTHETIC. Every cell below is a typed input copied from the fixture (blue). Contract v1 / v1.1, docs/contracts.md.")
    r = 4
    for ent in ENTITIES:
        rows = GOLD[ent]
        cols = list(rows[0])
        ws.cell(row=r, column=1, value=ent).font = BOLD
        r += 1
        header(ws, r, cols)
        POS[ent] = {"first": r + 1, "last": r + len(rows), "col": {c: L(i + 1) for i, c in enumerate(cols)}}
        for row in rows:
            r += 1
            for i, c in enumerate(cols):
                cell = ws.cell(row=r, column=i + 1, value=num(row[c]) if c not in ("sku",) else row[c])
                cell.font = BLUE
        r += 2
    widths(ws, {L(i): 20 for i in range(1, 8)})


def ref(ent, field, absolute=True):
    p = POS[ent]
    c = p["col"][field]
    return f"Inputs!${c}${p['first']}:${c}${p['last']}"


def cell(ent, field, i):
    return f"Inputs!{POS[ent]['col'][field]}{POS[ent]['first'] + i}"


# ------------------------------------------------------------------------------------------ Allocation
ALLOC: dict[str, tuple[int, int]] = {}
TAIL = ["total_cents", "weight", "weight_sum", "floor_cents", "remainder", "leftover_cents", "rank", "extra_cent",
        "allocated_amount"]


def tail(ws, r, a, b):
    """Columns G..O: the largest-remainder step as live formulas (D-04). A=id, B=group; rows a..b of the block."""
    B, K, A = f"$B${a}:$B${b}", f"$K${a}:$K${b}", f"$A${a}:$A${b}"
    ws[f"I{r}"] = f"=SUMIF({B},B{r},$H${a}:$H${b})"
    ws[f"J{r}"] = f"=IF(I{r}=0,0,QUOTIENT(G{r}*H{r},I{r}))"
    ws[f"K{r}"] = f"=IF(I{r}=0,0,MOD(G{r}*H{r},I{r}))"
    ws[f"L{r}"] = f"=G{r}-SUMIF({B},B{r},$J${a}:$J${b})"
    ws[f"M{r}"] = f"=1+SUMPRODUCT(({B}=B{r})*({K}>K{r}))+SUMPRODUCT(({B}=B{r})*({K}=K{r})*({A}<A{r}))"
    ws[f"N{r}"] = f"=IF(AND(I{r}>0,M{r}<=L{r}),1,0)"
    ws[f"O{r}"] = f"=(J{r}+N{r})/100"
    ws[f"O{r}"].number_format = MONEY


def build_allocation(wb):
    ws = wb.create_sheet("Allocation")
    title(ws, "Allocation: freight and discounts by shipment line, largest remainder as live formulas (D-04)",
          "Cents are integers: floor(T x w / W), remainder mod(T x w, W); leftover cents go to the largest remainders, "
          "ties to the lowest shipment_line_id. Zero-gross fallback (weight = quantity) is in the weight formula.")
    n = len(GOLD["shipment_lines"])
    # block 1: freight
    r0 = 4
    ws.cell(row=r0, column=1, value="1. Freight: shipping_cost over the shipment's lines by gross-sales weight").font = BOLD
    header(ws, r0 + 1, ["shipment_line_id", "shipment_id", "line_id", "shipped_quantity", "unit_price", "gross_sales"] + TAIL)
    a, b = r0 + 2, r0 + 1 + n
    for i in range(n):
        r = a + i
        ws[f"A{r}"] = "=" + cell("shipment_lines", "shipment_line_id", i)
        ws[f"B{r}"] = "=" + cell("shipment_lines", "shipment_id", i)
        ws[f"C{r}"] = "=" + cell("shipment_lines", "line_id", i)
        ws[f"D{r}"] = "=" + cell("shipment_lines", "shipped_quantity", i)
        ws[f"E{r}"] = f"=INDEX({ref('order_lines', 'unit_price')},MATCH(C{r},{ref('order_lines', 'line_id')},0))"
        ws[f"F{r}"] = f"=D{r}*E{r}"
        ws[f"G{r}"] = f"=ROUND(INDEX({ref('shipments', 'shipping_cost')},MATCH(B{r},{ref('shipments', 'shipment_id')},0)),2)*100"
        ws[f"H{r}"] = f"=IF(SUMIF($B${a}:$B${b},B{r},$F${a}:$F${b})>0,ROUND(F{r}*100,0),D{r})"
        tail(ws, r, a, b)
        for col in "EF":
            ws[f"{col}{r}"].number_format = MONEY
    ALLOC["freight"] = (a, b)
    # block 2: discounts
    r0 = b + 3
    ws.cell(row=r0, column=1, value="2. Discounts: round(discount x shipped/ordered, 2) per order line, split by shipped quantity").font = BOLD
    header(ws, r0 + 1, ["shipment_line_id", "line_id", "shipment_id", "shipped_quantity", "ordered_quantity",
                        "line_discount"] + TAIL)
    a, b = r0 + 2, r0 + 1 + n
    for i in range(n):
        r = a + i
        ws[f"A{r}"] = "=" + cell("shipment_lines", "shipment_line_id", i)
        ws[f"B{r}"] = "=" + cell("shipment_lines", "line_id", i)
        ws[f"C{r}"] = "=" + cell("shipment_lines", "shipment_id", i)
        ws[f"D{r}"] = "=" + cell("shipment_lines", "shipped_quantity", i)
        ws[f"E{r}"] = f"=INDEX({ref('order_lines', 'quantity')},MATCH(B{r},{ref('order_lines', 'line_id')},0))"
        ws[f"F{r}"] = f"=SUMIF({ref('discounts', 'line_id')},B{r},{ref('discounts', 'discount_amount')})"
        ws[f"G{r}"] = (f"=ROUND(F{r}*SUMIF({ref('shipment_lines', 'line_id')},B{r},{ref('shipment_lines', 'shipped_quantity')})"
                       f"/E{r},2)*100")
        ws[f"H{r}"] = f"=D{r}"
        tail(ws, r, a, b)
        ws[f"F{r}"].number_format = MONEY
    ALLOC["discount"] = (a, b)
    # block 3: self-test of the remainder step (author-typed, not data)
    r0 = b + 3
    ws.cell(row=r0, column=1, value="3. Formula self-test (author-typed illustrative inputs, NOT fixture or snapshot data): "
            "10.00 split over three equal weights; expected 3.34 / 3.33 / 3.33 (one leftover cent to the lowest id)").font = BOLD
    header(ws, r0 + 1, ["id", "group", "amount_to_split", "n/a", "n/a", "n/a"] + TAIL)
    a, b = r0 + 2, r0 + 4
    for i in range(3):
        r = a + i
        ws[f"A{r}"] = f"t-{i + 1}"
        ws[f"B{r}"] = "g"
        ws[f"C{r}"] = 10.00
        ws[f"G{r}"] = f"=ROUND(C{r},2)*100"
        ws[f"H{r}"] = 1
        for col in "ABCH":
            ws[f"{col}{r}"].font = BLUE
        tail(ws, r, a, b)
    ALLOC["selftest"] = (a, b)
    widths(ws, {L(i): 15 for i in range(1, 16)})
    ws.freeze_panes = "A4"


# ----------------------------------------------------------------------------------------------- Bridge
BR: dict[str, int] = {}


def build_bridge(wb):
    ws = wb.create_sheet("Bridge")
    title(ws, "Bridge: gross sales to contribution margin (D-05), baseline and after the return",
          "All cells are formulas over Inputs and Allocation. Blue cells are typed switches/inputs. USD only (fixture FX 1.0).")
    ws["A3"], ws["B3"], ws["C3"] = "Returns included (1 = yes)", 0, 1
    ws["B3"].font = ws["C3"].font = BLUE
    header(ws, 4, ["Step", "Baseline (before return)", "After return", "Contract rule"])
    fa, fb = ALLOC["freight"]
    da, db = ALLOC["discount"]
    h0 = 26                          # helper table header row
    ha, hb = h0 + 1, h0 + len(GOLD["shipment_lines"])
    steps = [
        ("gross_sales", "Gross sales", f"=SUM(Allocation!$F${fa}:$F${fb})", "sum(shipped qty x unit price)"),
        ("discounts", "Allocated discounts", f"=SUM(Allocation!$O${da}:$O${db})", "line discount x shipped/ordered, largest remainder"),
        ("credits", "Merchandise credits", "={c}$3*SUM($F${ha}:$F${hb})", "posted credits, recognised on credit_date (D-01a)"),
        ("net", "Net merchandise sales", "={c}5-{c}6-{c}7", "gross - discounts - credits"),
        ("cogs_gross", "COGS before returns", "=SUM($E${ha}:$E${hb})", "shipped qty x unit cost of the line's cost version"),
        ("recov", "Recoverable return cost", "={c}$3*SUM($H${ha}:$H${hb})", "returned qty x that same unit cost"),
        ("cogs", "COGS", "={c}9-{c}10", "COGS before returns - recoverable return cost"),
        ("gp", "Gross profit", "={c}8-{c}11", "net - COGS"),
        ("ship_rev", "Shipping revenue", 0, "none in the fixture (typed input)"),
        ("ship_cost", "Shipping cost (allocated freight)", f"=SUM(Allocation!$O${fa}:$O${fb})", "sum of allocated freight"),
        ("other", "Other variable costs", 0, "none in the fixture (typed input)"),
        ("cm", "Contribution margin", "={c}12+{c}13-{c}14-{c}15", "GP + shipping revenue - shipping cost - other"),
        ("residual", "Allocation residual", f"=SUM({ref('shipments', 'shipping_cost')})-SUM(Allocation!$O${fa}:$O${fb})", "shipping_cost - sum(allocated), must be 0.00"),
        ("elig", "Eligible net revenue", "={c}8+{c}13", "net + shipping revenue"),
        ("cm_rate", "Contribution margin rate", '=IF({c}18=0,"",{c}16/{c}18)', "contribution / eligible net revenue; blank when 0"),
        ("disc_rate", "Discount rate", '=IF({c}5=0,"",{c}6/{c}5)', "discounts / gross sales; blank when 0"),
        ("ret_rate", "Return rate", f'=IF(SUM($D${ha}:$D${hb})=0,"",{{c}}$3*SUM($G${ha}:$G${hb})/SUM($D${ha}:$D${hb}))', "returned units / shipped units; blank when 0"),
    ]
    for i, (key, label, f, rule) in enumerate(steps):
        r = 5 + i
        BR[key] = r
        ws[f"A{r}"] = label
        ws[f"D{r}"] = rule
        for c in "BC":
            v = f.format(c=c, ha=ha, hb=hb) if isinstance(f, str) else f
            ws[f"{c}{r}"] = v
            if not isinstance(f, str):
                ws[f"{c}{r}"].font = BLUE
            ws[f"{c}{r}"].number_format = "0.000000" if key.endswith("rate") else MONEY
    assert BR["cm"] == 16 and BR["elig"] == 18 and BR["gross_sales"] == 5, BR
    for k in ("net", "gp", "cm"):
        ws[f"A{BR[k]}"].font = BOLD
    ws.cell(row=h0 - 1, column=1, value="Helper: cost and returns per shipment line").font = BOLD
    # helper columns: A id, B cost version, C unit cost, D shipped qty, E cogs, F credit, G returned qty, H recoverable
    header(ws, h0, ["shipment_line_id", "cost_version_id", "unit_cost", "shipped_quantity", "cogs_before_returns",
                    "credit_amount", "returned_quantity", "recoverable_cost"])
    for i in range(len(GOLD["shipment_lines"])):
        r = ha + i
        ws[f"A{r}"] = "=" + cell("shipment_lines", "shipment_line_id", i)
        ws[f"B{r}"] = "=" + cell("shipment_lines", "cost_version_id", i)
        ws[f"C{r}"] = f"=INDEX({ref('cost_versions', 'unit_cost')},MATCH(B{r},{ref('cost_versions', 'cost_version_id')},0))"
        ws[f"D{r}"] = "=" + cell("shipment_lines", "shipped_quantity", i)
        ws[f"E{r}"] = f"=C{r}*D{r}"
        ws[f"F{r}"] = f"=SUMIF({ref('returns', 'shipment_line_id')},A{r},{ref('returns', 'credit_amount')})"
        ws[f"G{r}"] = f"=SUMIF({ref('returns', 'shipment_line_id')},A{r},{ref('returns', 'returned_quantity')})"
        ws[f"H{r}"] = f"=G{r}*C{r}"
        for col in "CEFH":
            ws[f"{col}{r}"].number_format = MONEY
    widths(ws, {"A": 34, "B": 24, "C": 18, "D": 52, "E": 20, "F": 16, "G": 16, "H": 18})


# --------------------------------------------------------------------------------------------- Checks
def build_checks(wb):
    ws = wb.create_sheet("Checks")
    title(ws, "Checks: workbook results against tests/fixtures/golden/expected.json",
          "Expected values are typed from expected.json (blue); results are live formulas from Bridge. Rates compared at 6 dp, money at 2 dp.")
    header(ws, 4, ["state", "metric", "expected (expected.json)", "workbook", "match"])
    rows = [("gross_sales", "gross_sales"), ("allocated_discounts", "discounts"), ("merchandise_credits", "credits"),
            ("net_merchandise_sales", "net"), ("cogs", "cogs"), ("gross_profit", "gp"), ("shipping_cost", "ship_cost"),
            ("contribution_margin", "cm"), ("allocation_residual", "residual"),
            ("contribution_margin_rate", "cm_rate"), ("discount_rate", "disc_rate"), ("return_rate", "ret_rate")]
    r = 5
    for state, col in (("baseline", "B"), ("after_return", "C")):
        for key, bkey in rows:
            exp = Decimal(EXPECTED["states"][state]["metrics"][key])
            ws[f"A{r}"], ws[f"B{r}"] = state, key
            ws[f"C{r}"] = float(exp)
            ws[f"C{r}"].font = BLUE
            ws[f"D{r}"] = f"=Bridge!{col}{BR[bkey]}"
            dp = 6 if key.endswith("rate") else 2
            ws[f"E{r}"] = f"=ROUND(D{r},{dp})=ROUND(C{r},{dp})"
            r += 1
    ws[f"A{r + 1}"] = "all checks"
    ws[f"E{r + 1}"] = f"=AND(E5:E{r - 1})"
    ws[f"A{r + 3}"] = "Allocation self-test (block 3): 3.34 / 3.33 / 3.33, sum 10.00"
    sa, sb = ALLOC["selftest"]
    ws[f"E{r + 3}"] = (f"=AND(ROUND(Allocation!O{sa},2)=3.34,ROUND(Allocation!O{sa + 1},2)=3.33,"
                       f"ROUND(Allocation!O{sb},2)=3.33,ROUND(SUM(Allocation!O{sa}:O{sb}),2)=10)")
    widths(ws, {"A": 16, "B": 28, "C": 24, "D": 16, "E": 10})


# ------------------------------------------------------------------------------------------ Assumptions
DECISIONS = [
    ("D-01", "Grain, population, time basis",
     "Fact grain is the shipment line. Population: lines whose shipment shipped_at is in the half-open UTC window [from, to), order status shipped or "
     "partially_returned, not quarantined. Cancelled orders never enter. Returns are recognised on credit_date (D-01a, contract v1.1): lines by "
     "shipped_at, credits and recoverable cost by credit_date, a closed window stays closed."),
    ("D-02", "Money, currency, rounding",
     "Decimal money; NUMERIC(20,6) in the database. Convert with the exchange_rates row whose rate_date equals the shipment's UTC date (returns: the credit date). "
     "A missing rate quarantines the line (FX_RATE_MISSING), never 1.0 or 0. Rounding to cents only at allocation and presentation. This workbook is USD only."),
    ("D-03", "Historical cost",
     "cost_versions are half-open [valid_from, valid_to); the cost of a line is the single version containing shipped_at; zero or several quarantine the line "
     "(COST_VERSION_MISSING / COST_VERSION_AMBIGUOUS). A return reverses cost at the same version."),
    ("D-04", "Allocation of discounts and freight",
     "Discount: round(discount x shipped/ordered, 2) per order line, split by shipped quantity. Freight: shipping_cost split by gross-sales weight (quantity if "
     "gross is 0). Both use largest remainder, lowest shipment_line_id breaks ties; allocation_residual must be 0.00; keyed by shipment_line_id."),
    ("D-05", "Metric formulas",
     "Net = gross - discounts - credits. COGS = shipped qty x historical cost - recoverable return cost. Contribution = GP + shipping revenue - shipping cost - "
     "other variable costs. Rates are ratios of sums, null on a zero denominator; rates carried to six decimals."),
    ("D-06", "Workflow invariants",
     "Malformed currency rejects the order and dependents (raw rows kept); same idempotency key + different payload is 409; return quantity <= remaining and "
     "credit >= 0 (RETURN_EXCEEDS_REMAINING, CREDIT_SIGN_REVERSED); region = customers version effective at shipped_at else UNASSIGNED; quotas join at (rep, period) only."),
    ("D-07", "Scenario deltas",
     "Signed fractions discount_delta and shipping_delta apply to a copy: allocated_discount x (1 + discount_delta), freight x (1 + shipping_delta); credits and "
     "costs not rescaled; results are labelled estimates. Not implemented in this workbook."),
]
LOCAL = [
    "Everything in Inputs, Allocation and Bridge is synthetic (golden fixture revenue-atlas-golden-v1); nothing here describes a real company.",
    "Zero-gross shipments allocate freight by shipped quantity (weight formula); the explicit unallocated_freight row of D-04 is not modelled.",
    "A single FX rate of 1.0 (USD to USD) applies; non-USD conversion at shipped or credit-date rates is not modelled in the formula chain.",
    "Cached values: see docs/workbook.md. Formulas are the deliverable; Excel recalculates on open (fullCalcOnLoad is set).",
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, f"Assumptions: contract decisions D-01..D-07 (docs/contracts.md, contract v1.1)")
    header(ws, 3, ["ID", "Decision", "Rule as applied in this workbook"])
    for i, (k, t, d) in enumerate(DECISIONS):
        r = 4 + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = k, t, d
        ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    r = 4 + len(DECISIONS) + 1
    ws[f"A{r}"] = "Workbook-specific"
    ws[f"A{r}"].font = BOLD
    for j, t in enumerate(LOCAL):
        ws[f"A{r + 1 + j}"] = f"W-{j + 1}"
        ws[f"C{r + 1 + j}"] = t
        ws[f"C{r + 1 + j}"].alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 10, "B": 34, "C": 120})


# ------------------------------------------------------------------------------------ Snapshot summary
SECTIONS = [
    ("01_margin_bridge_by_month", "Margin bridge by month (USD)"),
    ("02_contribution_by_region", "Contribution by region, UNASSIGNED shown"),
    ("03_discount_rate_by_region_before_after_2026_03", "Discount rate by region, before / after 2026-03"),
    ("04_fixtures_cost_vs_price_by_month", "cat-fixtures unit cost vs price by month"),
    ("05_return_rate_by_segment_category", "Return rate by segment x category (top 12 by mature_rate)"),
    ("06_top10_customers_contribution_decline", "Top 10 customers by contribution decline, A=2025-10..2026-03 vs B=2026-04..2026-09"),
    ("07_quarantine_reasons", "Quarantine and rejection reasons"),
    ("08_allocation_residual_check", "Allocation residual check"),
    ("09_stored_vs_contract_freight", "Stored C04 freight vs contract freight (diagnostic)"),
    ("reconciliation", "Reconciliation: bridge TOTAL vs data/sql/reference_totals.sql"),
    ("python_domain_metrics", "Python domain run_pipeline metrics over the same snapshot files (all 21 contract metrics)"),
    ("memo_arithmetic", "Derived arithmetic used in docs/executive-memo.md"),
]


def conv(v):
    if isinstance(v, str):
        try:
            if v.lstrip("-").replace(".", "", 1).isdigit() and not (len(v) > 1 and v[0] == "0" and v[1] != "."):
                return float(v)
        except ValueError:
            pass
    return v


def build_summary(wb):
    ws = wb.create_sheet("Snapshot summary")
    title(ws, "Snapshot summary: PASTED VALUES (not formulas) from the SQL analysis",
          f"Source: data/sql/analysis/run_all.py on the seeded synthetic snapshot, dataset_hash {MANIFEST['dataset_hash'][:16]}..., "
          f"dataset_version {MANIFEST['dataset_version']}. Re-run the query to refresh; nothing on this sheet recalculates.")
    r = 4
    for key, label in SECTIONS:
        rows = RESULTS[key]
        ws.cell(row=r, column=1, value=f"{key}.sql: {label}" if key[0].isdigit() else f"{key}: {label}").font = BOLD
        ws.cell(row=r, column=1).fill = SEC
        r += 1
        cols = list(rows[0])
        header(ws, r, cols)
        for row in rows:
            r += 1
            for i, c in enumerate(cols):
                v = row[c]
                v = str(v).lower() if isinstance(v, bool) else v if isinstance(v, int) else conv(v)
                ws.cell(row=r, column=i + 1, value=v)
        r += 2
    widths(ws, {L(i): 18 for i in range(1, 12)})
    ws.column_dimensions["C"].width = 22


def build():
    wb = Workbook()
    build_inputs(wb)
    build_allocation(wb)
    build_bridge(wb)
    build_assumptions(wb)
    build_summary(wb)
    build_checks(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print("wrote", OUT)


def recalc():
    """Recalculate with LibreOffice headless and replace the file (adds cached values, keeps formulas)."""
    soffice = shutil.which("soffice")
    if not soffice:
        print("soffice not found: workbook left without cached values")
        return 1
    with tempfile.TemporaryDirectory() as d:
        subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", d, str(OUT)], check=True, capture_output=True, timeout=180)
        shutil.copy(Path(d) / OUT.name, OUT)
    print("recalculated with LibreOffice and saved", OUT)
    return 0


if __name__ == "__main__":
    build()
    sys.exit(recalc() if "--recalc" in sys.argv else 0)
