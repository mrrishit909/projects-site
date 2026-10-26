"""Build artifacts/serviceflow-analyst-workbook.xlsx with openpyxl (P04-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: the published snapshot ``web/public/data`` read through the Python domain (the event log is replayed once, in Python, to the
per-ticket instants of the ``Tickets`` sheet and the service intervals of ``Intervals``), tests/fixtures/golden/expected.json,
artifacts/analysis-results.json (written by data/sql/analysis/run_all.py) and web/public/data/manifest.json. openpyxl cannot
calculate, so the file it writes holds formulas only. With --recalc, LibreOffice (soffice) opens it headless and saves a recalculated
copy over the same path, which then also carries cached values. Every figure of ``Metrics``, ``Hour of week``, ``Queue by priority`` and
``Golden case`` is a live formula over the data sheets and the ``Scenario`` switch (full / training / holdout window); the only typed
numbers there are the inputs, the tolerances and the expected values of the check columns. Excel is never run.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "domain"), str(ROOT), str(ROOT / "data" / "sql" / "analysis")]

import python_side as ps  # noqa: E402
from serviceflow import compute_metrics  # noqa: E402
from serviceflow.timeutil import d4, nearest_rank  # noqa: E402

OUT = ROOT / "artifacts" / "serviceflow-analyst-workbook.xlsx"
RESULTS = json.loads((ROOT / "artifacts/analysis-results.json").read_text())
MANIFEST = json.loads((ROOT / "web/public/data/manifest.json").read_text())
EXPECTED = json.loads((ROOT / "tests/fixtures/golden/expected.json").read_text())
CODE_REVISION = "43de7f8"
DATASET_HASH = "944ef91b3dc17c04c6bb601e0faac94f956d121877de2b9864988cf7c5d5cba4"

HDR = PatternFill("solid", fgColor="0F1B33")
SEC = PatternFill("solid", fgColor="E8E1CC")
BLUE = Font(color="0000FF")            # typed inputs
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
EPOCH = datetime(1970, 1, 1, tzinfo=timezone.utc)


def epoch(dt: datetime | None) -> int | None:
    return None if dt is None else int((dt - EPOCH).total_seconds())


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---- derived rows (Python, once) ------------------------------------------------------------------------------------------
def snapshot_rows(ctx: ps.Ctx):
    tickets, intervals = [], []
    for t in sorted(ctx.pop.tickets, key=lambda x: x.ticket_id):
        h = ctx.hist[t.ticket_id]
        tickets.append([t.ticket_id, t.priority, t.channel, t.queue_id, epoch(t.created_at), epoch(h.first_human_reply_at),
                        epoch(h.first_service_at), epoch(h.final_resolution), epoch(h.abandoned_at), h.reopen_count, len(h.transfers),
                        ps.ts(t.created_at)])
        for iv in h.intervals:
            intervals.append([t.ticket_id, iv.agent_id, epoch(iv.start), epoch(iv.end)])
    intervals.sort(key=lambda r: (r[2], r[0], r[1]))
    rec = {}
    for b in ctx.breaks:
        rec[b["shift_id"]] = rec.get(b["shift_id"], 0) + int((ps.parse_ts(b["ends_at"]) - ps.parse_ts(b["starts_at"])).total_seconds() // 60)
    shifts = [[s.shift_id, s.agent_id, epoch(s.starts_at), epoch(s.ends_at), s.break_minutes, rec.get(s.shift_id, 0)]
              for s in sorted(ctx.pop.shifts, key=lambda x: x.shift_id)]
    by_ticket = {}
    for s in ctx.pop.surveys:
        by_ticket[s.ticket_id] = s
    surveys = [[s.survey_id, s.ticket_id, 1 if s.invited_at is not None else 0, s.score] for s in sorted(ctx.pop.surveys, key=lambda x: x.survey_id)]
    return tickets, intervals, shifts, surveys


def domain_reference(ctx: ps.Ctx) -> dict[str, dict[str, float]]:
    """The domain's figures for the three windows as numbers (pasted into 'Domain reference'; never read by a formula except the checks)."""
    out: dict[str, dict[str, float]] = {}
    for name, w in ps.WINDOWS.items():
        m = compute_metrics(ctx.pop, w, [w[1]], ps.ABANDONABLE)
        vals = {k: v for k, v in m.items() if not isinstance(v, (dict, bool)) and v is not None}
        vals["backlog_at_window_end"] = list(m["backlog"].values())[0]
        for key, src in (("p90_wait_minutes", "wait_minutes"), ("p90_resolution_minutes", "resolution_minutes")):
            vals[key] = d4(nearest_rank([Fraction(x) for x in m[src].values()], Fraction(9, 10)))
        out[name] = {k: float(v) for k, v in vals.items()}
    return out


# ---- sheets ---------------------------------------------------------------------------------------------------------------
def build_readme(wb):
    ws = wb.active
    ws.title = "Read me"
    title(ws, "ServiceFlow analyst workbook (P04-ART-05)", "SYNTHETIC DATA: every ticket, agent and figure is generated (seed 20261026); nothing describes a real support team.")
    rows = [
        ("What it is", "A formula-driven workbook over the published synthetic snapshot of ServiceFlow (3,817 accepted tickets, 12 weeks from 2026-07-06). It recomputes the headline metrics of the contract with live formulas so an analyst can follow a number cell by cell, and it can be switched between the whole twelve weeks, the training weeks (1-8) and the holdout weeks (9-12)."),
        ("How to read it without installing anything", "Open the .xlsx in any spreadsheet application, or read docs/workbook.md (it lists every sheet and the checked values). The delivered file already carries values calculated by LibreOffice."),
        ("How to use the switch", "Scenario!B3 takes full, training or holdout. Every figure on Metrics, Hour of week and Queue by priority follows it; the check column of Metrics compares each figure with the Python domain's value for the same window (pasted on 'Domain reference')."),
        ("Observed, simulated, hypothetical", "Everything on the data sheets and on Metrics is OBSERVED in a synthetic history. Nothing here is a simulation result: the three simulated scenarios and their replications are in docs/executive-memo.md and artifacts/bi/. No staffing benefit is claimed."),
        ("Inputs (blue) and formulas (black)", "Blue cells are typed inputs: the per-ticket instants (epoch seconds, UTC), shifts, intervals, surveys, SLA targets, the window table and the expected values of the check columns. Black cells are formulas."),
        ("Where the per-ticket instants come from", "The event log (21,893 accepted events) was replayed ONCE in Python (the domain's replay) to the first human reply, first service start, final resolution, abandonment and the service intervals; the workbook starts from those instants. The replay itself is not a spreadsheet formula (limit W-02)."),
        ("Dataset and code", f"Dataset hash {DATASET_HASH} (web/public/data/manifest.json, dataset_version {MANIFEST['dataset_version']}); code revision {CODE_REVISION}; contract v1 (docs/contracts.md D-01..D-09)."),
        ("Recalculation status", "openpyxl wrote formulas only; LibreOffice (headless) recalculated the delivered file. Microsoft Excel was NOT run: nothing here was opened or recalculated by Excel (docs/workbook.md)."),
        ("Sheets", "Scenario, Metrics, Domain reference, Hour of week, Queue by priority, Golden case, Assumptions, Tickets, Intervals, Shifts, Surveys, SQL results (pasted)."),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=a).font = BOLD
        c = ws.cell(row=i, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 34, "B": 130})


def build_scenario(wb):
    ws = wb.create_sheet("Scenario")
    title(ws, "Scenario switch, windows and SLA targets", "Change B3 to full, training or holdout. All windows are half-open [start, end) in UTC epoch seconds (D-08).")
    ws["A3"], ws["B3"] = "scenario", "full"
    ws["A3"].font = BOLD
    ws["B3"].font = Font(color="0000FF", bold=True)
    dv = DataValidation(type="list", formula1='"full,training,holdout"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    header(ws, 6, ["window", "start (epoch s)", "end (epoch s)", "start UTC", "end UTC"])
    for r, (name, (a, b)) in enumerate(ps.WINDOWS.items(), start=7):
        for c, v in enumerate([name, epoch(a), epoch(b), ps.ts(a), ps.ts(b)], start=1):
            ws.cell(row=r, column=c, value=v).font = BLUE
    ws["A11"], ws["B11"] = "window start (epoch s)", "=INDEX($B$7:$B$9,MATCH($B$3,$A$7:$A$9,0))"
    ws["A12"], ws["B12"] = "window end (epoch s)", "=INDEX($C$7:$C$9,MATCH($B$3,$A$7:$A$9,0))"
    ws["A13"], ws["B13"] = "window length (hours)", "=(B12-B11)/3600"
    ws["A14"], ws["B14"] = "backlog / open-overdue cutoff (epoch s)", "=B12"
    ws["C14"] = "the window end (the domain's default cutoff)"
    header(ws, 17, ["priority", "first response target (min)", "resolution target (min)"])
    for r, p in enumerate(ctx_policy(), start=18):
        for c, v in enumerate(p, start=1):
            ws.cell(row=r, column=c, value=v).font = BLUE
    ws["A22"], ws["B22"], ws["C22"] = "channels that can abandon", "chat", "phone"
    ws["B22"].font = ws["C22"].font = BLUE
    ws["A24"], ws["B24"] = "positive CSAT score >=", 4
    ws["B24"].font = BLUE
    widths(ws, {"A": 40, "B": 28, "C": 26, "D": 24, "E": 24})


_POLICY: list = []


def ctx_policy():
    return _POLICY


def build_tickets(wb, rows):
    ws = wb.create_sheet("Tickets")
    n = len(rows)
    labels = ["ticket_id", "priority", "channel", "final queue", "created (epoch s)", "first human reply (epoch s)", "first service (epoch s)",
              "final resolution (epoch s)", "abandoned (epoch s)", "reopens", "transfers",
              "in arrivals", "first response (min)", "wait to first service (min)", "in completions", "resolution (min)",
              "first response target (min)", "resolution target (min)", "first response within", "resolution within",
              "open at window end", "overdue at window end", "can abandon", "abandoned", "weekday (0=Mon)", "hour (UTC)", "created (UTC text)"]
    header(ws, 1, labels)
    w0, w1 = "Scenario!$B$11", "Scenario!$B$12"
    cut = "Scenario!$B$14"
    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(r[:11], start=1):
            if v is not None:
                ws.cell(row=i, column=c, value=v).font = BLUE
        ws.cell(row=i, column=27, value=r[11]).font = BLUE
        f = {
            12: f"=IF(AND(E{i}>={w0},E{i}<{w1}),1,0)",
            13: f'=IF(AND(L{i}=1,ISNUMBER(F{i})),MAX(0,F{i}-E{i})/60,"")',
            14: f'=IF(AND(L{i}=1,ISNUMBER(G{i})),MAX(0,G{i}-E{i})/60,"")',
            15: f"=IF(ISNUMBER(H{i}),IF(AND(H{i}>={w0},H{i}<{w1}),1,0),0)",
            16: f'=IF(O{i}=1,MAX(0,H{i}-E{i})/60,"")',
            17: f"=INDEX(Scenario!$B$18:$B$20,MATCH(B{i},Scenario!$A$18:$A$20,0))",
            18: f"=INDEX(Scenario!$C$18:$C$20,MATCH(B{i},Scenario!$A$18:$A$20,0))",
            19: f'=IF(ISNUMBER(M{i}),IF(M{i}<=Q{i},1,0),"")',
            20: f'=IF(ISNUMBER(P{i}),IF(P{i}<=R{i},1,0),"")',
            21: f"=IF(E{i}>{cut},0,IF(OR(AND(ISNUMBER(H{i}),H{i}<={cut}),AND(ISNUMBER(I{i}),I{i}<={cut})),0,1))",
            22: f"=IF(AND(U{i}=1,({cut}-E{i})/60>R{i}),1,0)",
            23: f"=IF(OR(C{i}=Scenario!$B$22,C{i}=Scenario!$C$22),1,0)",
            24: f"=IF(ISNUMBER(I{i}),1,0)",
            25: f"=MOD(INT(E{i}/86400)+3,7)",
            26: f"=MOD(INT(E{i}/3600),24)",
        }
        for c, v in f.items():
            ws.cell(row=i, column=c, value=v)
    ws.freeze_panes = "B2"
    widths(ws, {"A": 10, "B": 9, "C": 10, "D": 10, "E": 16, "F": 16, "G": 16, "H": 16, "I": 14, "AA": 22})
    return n


def build_intervals(wb, rows):
    ws = wb.create_sheet("Intervals")
    header(ws, 1, ["ticket_id", "agent_id", "start (epoch s)", "end (epoch s)", "handling in window (min)"])
    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(r, start=1):
            ws.cell(row=i, column=c, value=v).font = BLUE
        ws.cell(row=i, column=5, value=f"=MAX(0,MIN(D{i},Scenario!$B$12)-MAX(C{i},Scenario!$B$11))/60")
    ws.freeze_panes = "A2"
    widths(ws, {"A": 10, "B": 10, "C": 16, "D": 16, "E": 22})
    return len(rows)


def build_shifts(wb, rows):
    ws = wb.create_sheet("Shifts")
    header(ws, 1, ["shift_id", "agent_id", "start (epoch s)", "end (epoch s)", "declared break (min)", "recorded break (min)",
                   "staffed available in window (min)", "whole shift net of declared break (min)"])
    w0, w1 = "Scenario!$B$11", "Scenario!$B$12"
    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(r, start=1):
            ws.cell(row=i, column=c, value=v).font = BLUE
        ws.cell(row=i, column=7, value=f"=IF(MIN(D{i},{w1})>MAX(C{i},{w0}),(MIN(D{i},{w1})-MAX(C{i},{w0}))/60-IF(AND(C{i}>={w0},D{i}<={w1}),E{i},0),0)")
        ws.cell(row=i, column=8, value=f"=(D{i}-C{i})/60-E{i}")
    ws.freeze_panes = "A2"
    widths(ws, {"A": 10, "B": 10, "C": 16, "D": 16, "E": 14, "F": 14, "G": 18, "H": 22})
    return len(rows)


def build_surveys(wb, rows, n_tickets):
    ws = wb.create_sheet("Surveys")
    header(ws, 1, ["survey_id", "ticket_id", "invited", "score (blank: no valid response)", "ticket completed in window", "invited in window",
                   "valid response in window", "positive in window"])
    last = n_tickets + 1
    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(r, start=1):
            if v is not None:
                ws.cell(row=i, column=c, value=v).font = BLUE
        ws.cell(row=i, column=5, value=f"=INDEX(Tickets!$O$2:$O${last},MATCH(B{i},Tickets!$A$2:$A${last},0))")
        ws.cell(row=i, column=6, value=f"=C{i}*E{i}")
        ws.cell(row=i, column=7, value=f"=IF(AND(C{i}=1,ISNUMBER(D{i})),1,0)*E{i}")
        ws.cell(row=i, column=8, value=f"=IF(AND(G{i}=1,D{i}>=Scenario!$B$24),1,0)")
    ws.freeze_panes = "A2"
    widths(ws, {"A": 10, "B": 10, "C": 9, "D": 18, "E": 16, "F": 14, "G": 14, "H": 14})
    return len(rows)


def build_domain_reference(wb, ref):
    ws = wb.create_sheet("Domain reference")
    title(ws, "PASTED VALUES from the Python domain (compute_metrics), not formulas",
          "One column per window: the domain's figure computed over web/public/data with exact rationals. Metrics reads these only for its check column.")
    header(ws, 4, ["metric", "full", "training", "holdout"])
    keys = list(ref["full"])
    for r, k in enumerate(keys, start=5):
        ws.cell(row=r, column=1, value=k)
        for c, w in enumerate(("full", "training", "holdout"), start=2):
            ws.cell(row=r, column=c, value=ref[w][k])
    widths(ws, {"A": 36, "B": 16, "C": 16, "D": 16})
    return {k: r for r, k in enumerate(keys, start=5)}


METRICS = [
    # key, label, unit, formula template ({T} ticket rows last, {I} interval last, {S} shift last, {V} survey last, rows by key), places, definition
    ("arrival_count", "Arrivals", "tickets", "=SUM(Tickets!$L$2:$L${T})", 0, "tickets created in the window [start, end)"),
    ("completion_count", "Completions", "tickets", "=SUM(Tickets!$O$2:$O${T})", 0, "tickets whose FINAL resolution is in the window (a reopened ticket is one completion)"),
    ("arrival_rate_per_hour", "Arrival rate", "tickets / hour", "=ROUND({arrival_count}/Scenario!$B$13,6)", 6, "arrivals / window hours"),
    ("completion_rate_per_hour", "Completion rate", "tickets / hour", "=ROUND({completion_count}/Scenario!$B$13,6)", 6, "completions / window hours"),
    ("mean_first_response_minutes", "Mean first response", "minutes", "=ROUND(AVERAGE(Tickets!$M$2:$M${T}),4)", 4, "first human reply minus arrival, tickets with a human reply (a bot never counts)"),
    ("p90_first_response_minutes", "P90 first response", "minutes", "=ROUND(SMALL(Tickets!$M$2:$M${T},INT((9*COUNT(Tickets!$M$2:$M${T})+9)/10)),4)", 4, "nearest rank: the value at index ceil(0.9 n)"),
    ("mean_wait_minutes", "Mean wait to first service", "minutes", "=ROUND(AVERAGE(Tickets!$N$2:$N${T}),4)", 4, "first service_started minus the original arrival, tickets that were served"),
    ("p90_wait_minutes", "P90 wait to first service", "minutes", "=ROUND(SMALL(Tickets!$N$2:$N${T},INT((9*COUNT(Tickets!$N$2:$N${T})+9)/10)),4)", 4, "an addition beyond the 28 published keys; nearest rank"),
    ("mean_resolution_minutes", "Mean resolution", "minutes", "=ROUND(AVERAGE(Tickets!$P$2:$P${T}),4)", 4, "final resolution minus the ORIGINAL arrival (a reopened ticket is charged its dormant gap)"),
    ("p90_resolution_minutes", "P90 resolution", "minutes", "=ROUND(SMALL(Tickets!$P$2:$P${T},INT((9*COUNT(Tickets!$P$2:$P${T})+9)/10)),4)", 4, "an addition beyond the 28 published keys; nearest rank"),
    ("handling_minutes_total", "Handling", "minutes", "=ROUND(SUM(Intervals!$E$2:$E${I}),4)", 4, "service intervals clipped to the window; waiting_customer and idle time excluded"),
    ("sla_first_response_within", "First response within target", "tickets", "=SUM(Tickets!$S$2:$S${T})", 0, "first response <= the priority's target"),
    ("sla_first_response_eligible", "First response eligible", "tickets", "=COUNT(Tickets!$S$2:$S${T})", 0, "arrivals with a human reply (open overdue tickets are never in a denominator)"),
    ("sla_first_response_compliance", "SLA first response compliance", "ratio", "=ROUND({sla_first_response_within}/{sla_first_response_eligible},6)", 6, "within / eligible"),
    ("sla_resolution_within", "Resolution within target", "tickets", "=SUM(Tickets!$T$2:$T${T})", 0, "resolution <= the priority's target"),
    ("sla_resolution_eligible", "Resolution eligible", "tickets", "=COUNT(Tickets!$T$2:$T${T})", 0, "completions in the window"),
    ("sla_resolution_compliance", "SLA resolution compliance", "ratio", "=ROUND({sla_resolution_within}/{sla_resolution_eligible},6)", 6, "within / eligible"),
    ("backlog_at_window_end", "Backlog at the window end", "tickets", "=SUM(Tickets!$U$2:$U${T})", 0, "arrived by the end and not resolved or abandoned by it (a reopened ticket is read by its final resolution: exact for these three windows, W-03)"),
    ("open_overdue_at_cutoff", "Open overdue at the window end", "tickets", "=SUM(Tickets!$V$2:$V${T})", 0, "open and older than its resolution target; kept apart, never dropped (P04-TEST-06)"),
    ("staffed_available_minutes", "Staffed available time", "minutes", "=ROUND(SUM(Shifts!$G$2:$G${S}),4)", 4, "shift overlap with the window; the declared break is deducted only for a shift the window covers whole (D-05)"),
    ("staffed_shift_minutes_whole", "Staffed time, whole shifts", "minutes", "=ROUND(SUM(Shifts!$H$2:$H${S}),4)", 4, "every shift net of its declared break, clipped by nothing"),
    ("agent_occupancy", "Agent occupancy", "ratio", "=ROUND({handling_minutes_total}/{staffed_available_minutes},6)", 6, "handling / staffed available time"),
    ("abandonment_rate", "Abandonment rate", "ratio", "=ROUND(COUNTIFS(Tickets!$L$2:$L${T},1,Tickets!$W$2:$W${T},1,Tickets!$X$2:$X${T},1)/COUNTIFS(Tickets!$L$2:$L${T},1,Tickets!$W$2:$W${T},1),6)", 6, "abandoned / arrivals, only on channels that can abandon (chat, phone)"),
    ("survey_invitation_rate", "Survey invitation rate", "ratio", "=ROUND(SUM(Surveys!$F$2:$F${V})/{completion_count},6)", 6, "invited / completions"),
    ("survey_response_rate", "Survey response rate", "ratio", "=ROUND(SUM(Surveys!$G$2:$G${V})/SUM(Surveys!$F$2:$F${V}),6)", 6, "valid responses / invited"),
    ("csat", "CSAT", "ratio", "=ROUND(SUM(Surveys!$H$2:$H${V})/SUM(Surveys!$G$2:$G${V}),6)", 6, "positive (score >= 4) / valid responses; nonresponse is not random (docs/sql-analysis.md section 07)"),
]


def build_metrics(wb, ref_rows, n_t, n_i, n_s, n_v):
    ws = wb.create_sheet("Metrics", 1)
    title(ws, "Headline metrics for the selected window (live formulas)", "Window: Scenario!B3. 'check' is TRUE when the formula equals the Python domain's figure for the same window within the tolerance (half a unit in the last place).")
    ws["A3"], ws["B3"] = "selected scenario", "=Scenario!B3"
    ws["D3"], ws["E3"] = "all checks TRUE", None
    header(ws, 5, ["metric", "value (formula)", "unit", "definition", "Python domain (pasted)", "tolerance", "check"])
    rows = {k: 6 + i for i, (k, *_rest) in enumerate(METRICS)}
    last = 5 + len(METRICS)
    for k, label, unit, tpl, places, definition in METRICS:
        r = rows[k]
        f = tpl.format(T=n_t + 1, I=n_i + 1, S=n_s + 1, V=n_v + 1, **{kk: f"B{rr}" for kk, rr in rows.items()})
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f)
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=4, value=definition).alignment = Alignment(wrap_text=True, vertical="top")
        dr = ref_rows[k]
        ws.cell(row=r, column=5, value=f"=INDEX('Domain reference'!$B${dr}:$D${dr},MATCH(Scenario!$B$3,'Domain reference'!$B$4:$D$4,0))")
        ws.cell(row=r, column=6, value=0 if places == 0 else 0.5 * 10 ** -places).font = BLUE
        ws.cell(row=r, column=7, value=f"=ABS(B{r}-E{r})<=F{r}+0.0000000001")
        ws.cell(row=r, column=2).number_format = "0" if places == 0 else "0." + "0" * places
    ws["E3"] = f"=AND(G6:G{last})"
    ws["E3"].font = BOLD
    widths(ws, {"A": 34, "B": 18, "C": 14, "D": 90, "E": 20, "F": 12, "G": 9})
    return rows


def build_hour(wb, n_t, metric_rows):
    ws = wb.create_sheet("Hour of week")
    title(ws, "Arrivals by weekday and UTC hour for the selected window (COUNTIFS over Tickets)", "Weekday 0 = Monday. The Monday 09-11 UTC block is the capacity bind of the generator's design (docs/sql-analysis.md section 01).")
    header(ws, 4, ["weekday \\ hour (UTC)"] + list(range(24)) + ["total"])
    names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    T = n_t + 1
    for d in range(7):
        r = 5 + d
        ws.cell(row=r, column=1, value=d)
        for h in range(24):
            ws.cell(row=r, column=2 + h, value=f"=COUNTIFS(Tickets!$L$2:$L${T},1,Tickets!$Y$2:$Y${T},$A{r},Tickets!$Z$2:$Z${T},{h})")
        ws.cell(row=r, column=26, value=f"=SUM(B{r}:Y{r})")
        ws.cell(row=r, column=27, value=names[d])
    ws["A13"], ws["Z13"] = "all", "=SUM(Z5:Z11)"
    for h in range(24):
        col = 2 + h
        L = ws.cell(row=5, column=col).column_letter
        ws.cell(row=13, column=col, value=f"=SUM({L}5:{L}11)")
    ws["A15"], ws["B15"] = "Monday 09:00-11:00 UTC arrivals", "=K5+L5"
    ws["A16"], ws["B16"] = "grid total equals Metrics arrivals", f"=Z13=Metrics!B{metric_rows['arrival_count']}"
    ws["A17"], ws["B17"] = "expected Monday 09-11 arrivals (query 01, typed)", None
    full = next(r for r in RESULTS["01_arrivals_by_hour_of_week"] if r["block_utc"] == "Mon 09-11")
    ws["C17"], ws["D17"], ws["E17"] = full["arrivals"], full["arrivals_training"], full["arrivals_holdout"]
    ws["F17"] = "full / training / holdout"
    for c in ("C17", "D17", "E17"):
        ws[c].font = BLUE
    ws["A18"], ws["B18"] = "check (matches the selected window)", '=B15=INDEX(C17:E17,MATCH(Scenario!B3,{"full","training","holdout"},0))'
    widths(ws, {"A": 44})


def build_queue(wb, n_t):
    ws = wb.create_sheet("Queue by priority")
    title(ws, "Queue x priority for the selected window (final queue; COUNTIFS / AVERAGEIFS over Tickets)", "Expected columns are typed from query 02 and apply to the FULL window only; the check reads n/a for another scenario.")
    header(ws, 4, ["queue", "priority", "arrivals", "mean first response (min)", "mean wait to first service (min)", "first response compliance",
                   "query 02 arrivals (full)", "query 02 mean first response (full)", "query 02 mean wait (full)", "check (full only)"])
    T = n_t + 1
    q02 = {(r["queue_id"], r["priority"]): r for r in RESULTS["02_wait_and_response_distributions__first_response_and_wait"]}
    r = 5
    for q in ("q-001", "q-002", "q-003"):
        for p in ("P1", "P2", "P3"):
            e = q02[(q, p)]
            ws.cell(row=r, column=1, value=q)
            ws.cell(row=r, column=2, value=p)
            crit = f"Tickets!$D$2:$D${T},$A{r},Tickets!$B$2:$B${T},$B{r}"
            ws.cell(row=r, column=3, value=f"=COUNTIFS(Tickets!$L$2:$L${T},1,{crit})")
            ws.cell(row=r, column=4, value=f'=IFERROR(ROUND(AVERAGEIFS(Tickets!$M$2:$M${T},{crit}),4),"")')
            ws.cell(row=r, column=5, value=f'=IFERROR(ROUND(AVERAGEIFS(Tickets!$N$2:$N${T},{crit}),4),"")')
            ws.cell(row=r, column=6, value=f'=IFERROR(ROUND(SUMIFS(Tickets!$S$2:$S${T},{crit})/COUNTIFS({crit},Tickets!$S$2:$S${T},">=0"),6),"")')
            for c, v in ((7, e["arrivals"]), (8, float(e["mean_first_response_min"])), (9, float(e["mean_wait_min"]))):
                ws.cell(row=r, column=c, value=v).font = BLUE
            ws.cell(row=r, column=10, value=f'=IF(Scenario!$B$3="full",AND(C{r}=G{r},ABS(D{r}-H{r})<0.00005001,ABS(E{r}-I{r})<0.00005001),"n/a")')
            r += 1
    widths(ws, {"A": 10, "B": 10, "C": 12, "D": 20, "E": 22, "F": 20, "G": 16, "H": 20, "I": 18, "J": 16})


def build_golden(wb):
    ws = wb.create_sheet("Golden case")
    title(ws, "The known answer (P04-GOLD-01) as spreadsheet formulas", "One agent, three tickets at minute 0, deterministic 10-minute service, FIFO. Independent of every other sheet.")
    header(ws, 4, ["ticket", "arrival (min)", "service (min)", "service start (min)", "wait (min)", "completion (min)", "first response (min)"])
    ws["I4"], ws["J4"] = "service minutes", 10
    ws["J4"].font = BLUE
    for i, tid in enumerate(("t-001", "t-002", "t-003"), start=5):
        ws.cell(row=i, column=1, value=tid)
        ws.cell(row=i, column=2, value=0).font = BLUE
        ws.cell(row=i, column=3, value="=$J$4")
        ws.cell(row=i, column=4, value=f"=MAX(B{i},{'0' if i == 5 else 'F' + str(i - 1)})")
        ws.cell(row=i, column=5, value=f"=D{i}-B{i}")
        ws.cell(row=i, column=6, value=f"=D{i}+C{i}")
        ws.cell(row=i, column=7, value=f"=D{i}-B{i}")
    ws["A9"], ws["B9"] = "mean wait (min)", "=AVERAGE(E5:E7)"
    ws["A10"], ws["B10"] = "last completion (min)", "=MAX(F5:F7)"
    ws["A11"], ws["B11"] = "handling total (min)", "=SUM(C5:C7)"
    ws["A12"], ws["B12"] = "mean first response (min)", "=AVERAGE(G5:G7)"
    ws["D9"], ws["E9"] = "expected (expected.json)", float(EXPECTED["metrics"]["mean_wait_minutes"])
    ws["D10"], ws["E10"] = "expected", float(EXPECTED["simulation"]["final_completion_minute"])
    ws["D11"], ws["E11"] = "expected", float(EXPECTED["metrics"]["handling_minutes_total"])
    ws["D12"], ws["E12"] = "expected", float(EXPECTED["metrics"]["mean_first_response_minutes"])
    for r in (9, 10, 11, 12):
        ws.cell(row=r, column=5).font = BLUE
        ws.cell(row=r, column=6, value=f"=ABS(B{r}-E{r})<0.00005")
    ws["A14"], ws["B14"] = "starts 0 / 10 / 20 and waits 0 / 10 / 20", '=AND(D5=0,D6=10,D7=20,E5=0,E6=10,E7=20)'
    ws["A15"], ws["B15"] = "all golden checks TRUE", "=AND(F9:F12,B14)"
    widths(ws, {"A": 40, "B": 14, "C": 14, "D": 22, "E": 20, "F": 16, "G": 20})


ASSUMPTIONS = [
    ("D-01", "Events are an ordered log (occurred_at, event_id); FIFO within priority; a transfer keeps the arrival instant; a reopen is one customer contact."),
    ("D-02", "UTC instants; first response = first HUMAN reply (a bot never counts); resolution = final resolution minus the ORIGINAL arrival; P90 is nearest rank (index ceil(0.9 n)); handling = service intervals only."),
    ("D-03", "Staffed available time = shift minutes minus the declared break; occupancy = handling / staffed; abandonment only on channels that can abandon; CSAT = positive / valid responses."),
    ("D-05", "The break is deducted only when the window covers the whole shift; the backlog counts arrived tickets not resolved or abandoned; the survey invitation rate is over resolved tickets."),
    ("D-08", "The site window is [2026-07-06, as_of + 1 s); training = weeks 1-8, holdout = weeks 9-12; timestamps are integer epoch seconds."),
    ("D-09", "Wait means wait to first service; a reopened ticket is charged its dormant gap; calendars are ignored (every SLA policy uses the 24x7 cal-001)."),
    ("W-01", "Synthetic data only. The holdout's +20 % arrival intensity is a chosen realisation of a seeded series (docs/contracts.md D-09 R4-13)."),
    ("W-02", "The per-ticket instants and the service intervals were derived from the event log ONCE in Python (the domain's replay); the replay is not a spreadsheet formula. Everything after those instants is."),
    ("W-03", "Backlog and open overdue at the window end read a reopened ticket by its final resolution; checked equal to the domain for all three windows (1 open ticket at the end of the training window, none otherwise). It is not a general as-of reconstruction."),
    ("W-04", "Recalculation: openpyxl writes formulas only. The delivered file was recalculated by LibreOffice headless and carries its cached values; Excel was NOT run (docs/workbook.md). fullCalcOnLoad is set so Excel recalculates on open."),
    ("W-05", "Instants are integer epoch seconds (UTC), so a minute is exact to the second; the check tolerance is half a unit in the last place of the stated precision."),
    ("W-06", "Queue is the ticket's FINAL queue (D-09 R4-10). Per-group P90, the occupancy by agent and the hazard tables are in the SQL investigations, not in this workbook."),
    ("W-07", "Simulation results (docs/executive-memo.md) are estimates and are not part of this workbook."),
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions and limits", "D-nn are contract decisions (docs/contracts.md); W-nn are this workbook's own limits.")
    header(ws, 4, ["id", "statement"])
    for i, (a, b) in enumerate(ASSUMPTIONS, start=5):
        ws.cell(row=i, column=1, value=a).font = BOLD
        ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 8, "B": 150})


def put_table(ws, r, name, rows):
    ws.cell(row=r, column=1, value=name).font = BOLD
    ws.cell(row=r, column=1).fill = SEC
    r += 1
    if not rows:
        return r + 1
    cols = list(rows[0])
    header(ws, r, cols)
    for i, row in enumerate(rows, start=r + 1):
        for c, k in enumerate(cols, start=1):
            v = row[k]
            if isinstance(v, bool):
                v = "true" if v else "false"          # LibreOffice would save a boolean cell as =TRUE(): keep pasted values plain
            if isinstance(v, str) and v.replace(".", "", 1).lstrip("-").isdigit() and not v.startswith("q-") and k not in ("block_utc", "ticket_group", "window_name"):
                try:
                    v = float(v) if "." in v else int(v)
                except ValueError:
                    pass
            ws.cell(row=i, column=c, value=v)
    return r + len(rows) + 3


def build_sql_results(wb):
    ws = wb.create_sheet("SQL results (pasted)")
    meta = RESULTS["meta"][0]
    ws["A1"] = "PASTED VALUES from data/sql/analysis/run_all.py (artifacts/analysis-results.json); not formulas. They go stale if the data or the queries change."
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"dataset_hash {meta['dataset_hash']}; code revision {meta['code_revision']}; window {meta['analysis_window_utc']} UTC"
    ws["A3"] = "reconciliation (run_all.py): rows, rows equal"
    n = len(RESULTS["reconciliation"])
    ws["B3"], ws["C3"] = n, sum(1 for r in RESULTS["reconciliation"] if r["equal"] is True)
    r = 5
    for key, name in (("01_arrivals_by_hour_of_week", "01_arrivals_by_hour_of_week.sql"),
                      ("02_wait_and_response_distributions__first_response_and_wait", "02_wait_and_response_distributions.sql (first_response_and_wait)"),
                      ("03_sla_compliance_and_open_overdue__compliance_by_priority_and_window", "03_sla_compliance_and_open_overdue.sql (compliance)"),
                      ("04_occupancy_by_agent_and_shift__occupancy_by_agent", "04_occupancy_by_agent_and_shift.sql (by agent)"),
                      ("05_transfer_and_reopen_effects__by_transfer_and_reopen", "05_transfer_and_reopen_effects.sql (groups)"),
                      ("06_abandonment_by_channel_and_patience__abandonment_by_channel", "06_abandonment_by_channel_and_patience.sql (by channel)"),
                      ("07_survey_response_skew__response_by_experience", "07_survey_response_skew.sql (by experience)"),
                      ("08_rejections_and_defects_by_reason__rejected_by_reason", "08_rejections_and_defects_by_reason.sql (rejections)")):
        r = put_table(ws, r, name, RESULTS[key])
    widths(ws, {"A": 22, "B": 22, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18, "I": 18})


# ---- recalculation --------------------------------------------------------------------------------------------------------
def recalc(path: Path) -> None:
    """Open with LibreOffice headless and save a recalculated copy over ``path`` (cached values included)."""
    if shutil.which("soffice") is None:
        raise SystemExit("soffice not installed; the file has formulas but no cached values")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in" / path.name
        src.parent.mkdir()
        shutil.copy(path, src)
        out = Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=900)
        shutil.copy(out / path.name, path)
    set_full_calc_on_load(path)


def set_full_calc_on_load(path: Path) -> None:
    """LibreOffice drops fullCalcOnLoad; put it back (workbook.xml only) so Excel recalculates when it opens the file."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode("utf-8")
                assert "<calcPr " in text and "fullCalcOnLoad" not in text
                data = text.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ', 1).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build(path: Path) -> None:
    ctx = ps.load()
    _POLICY[:] = [[p, ctx.policy[p][0], ctx.policy[p][1]] for p in ("P1", "P2", "P3")]
    tickets, intervals, shifts, surveys = snapshot_rows(ctx)
    ref = domain_reference(ctx)
    wb = Workbook()
    build_readme(wb)
    build_scenario(wb)
    n_t = build_tickets(wb, tickets)
    n_i = build_intervals(wb, intervals)
    n_s = build_shifts(wb, shifts)
    n_v = build_surveys(wb, surveys, n_t)
    ref_rows = build_domain_reference(wb, ref)
    mrows = build_metrics(wb, ref_rows, n_t, n_i, n_s, n_v)
    build_hour(wb, n_t, mrows)
    build_queue(wb, n_t)
    build_golden(wb)
    build_assumptions(wb)
    build_sql_results(wb)
    order = ["Read me", "Scenario", "Metrics", "Domain reference", "Hour of week", "Queue by priority", "Golden case", "Assumptions",
             "Tickets", "Intervals", "Shifts", "Surveys", "SQL results (pasted)"]
    wb._sheets = [wb[n] for n in order]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def main(argv: list[str]) -> int:
    build(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    if "--recalc" in argv:
        recalc(OUT)
        print("recalculated with LibreOffice headless; cached values written")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
