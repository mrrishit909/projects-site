"""Build artifacts/supplyscope-analyst-workbook.xlsx with openpyxl (P03-ART-05).

Usage:  ../../.venv/bin/python artifacts/build_workbook.py [--recalc]

Inputs: tests/fixtures/golden/inputs.json and expected.json (golden fixture), artifacts/analysis-results.json (written by
data/sql/analysis/run_all.py), web/public/data/manifest.json. openpyxl cannot calculate, so the file it writes holds formulas
only. With --recalc, LibreOffice (soffice) opens it headless and saves a recalculated copy over the same path, which then also
carries cached values. Ledger, Position, Replenishment and Figures never hard-code a computed quantity: every figure is a formula
over the typed blocks of Inputs. Excel is never run.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "supplyscope-analyst-workbook.xlsx"
GOLD = json.loads((ROOT / "tests/fixtures/golden/inputs.json").read_text())
EXPECTED = json.loads((ROOT / "tests/fixtures/golden/expected.json").read_text())["states"]
RESULTS = json.loads((ROOT / "artifacts/analysis-results.json").read_text())
MANIFEST = json.loads((ROOT / "web/public/data/manifest.json").read_text())

HDR = PatternFill("solid", fgColor="1F3A5F")
SEC = PatternFill("solid", fgColor="DCE6F1")
BLUE = Font(color="0000FF")            # typed inputs
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
TS = "yyyy-mm-dd hh:mm:ss"
NUM = re.compile(r"^-?\d+(\.\d+)?$")
TS_RE = re.compile(r"^\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ$")

# Rows of each typed block on Inputs (filled by build_inputs, read by the formulas).
POS: dict[str, dict] = {}


def header(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill = WHITE, HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if note:
        ws["A2"] = note


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def parse(v):
    """Fixture value -> cell value: ISO instants become naive-UTC datetimes, null becomes blank, booleans text."""
    if isinstance(v, str) and TS_RE.match(v):
        return datetime.strptime(v, "%Y-%m-%dT%H:%M:%SZ")
    if isinstance(v, bool):
        return str(v).lower()
    if isinstance(v, str) and NUM.match(v) and "." in v:
        return float(v)                 # money such as "12.50"
    return v


# ---- Inputs ---------------------------------------------------------------------------------------------------------

PARAMS = [  # (name, value, source)
    ("as_of", GOLD["as_of"], "inputs.json as_of (editable: the as-of filter of the Ledger follows it)"),
    ("window_start", GOLD["window_start"], "inputs.json window_start"),
    ("reorder_point", GOLD["replenishment_params"][0]["reorder_point"], "replenishment_params.reorder_point"),
    ("target_position", GOLD["replenishment_params"][0]["target_position"], "replenishment_params.target_position"),
    ("minimum_order_quantity", GOLD["products"][0]["minimum_order_quantity"], "products.minimum_order_quantity"),
    ("pack_size", GOLD["products"][0]["pack_size"], "products.pack_size"),
    ("unit_cost", GOLD["products"][0]["unit_cost"], "products.unit_cost (no cost_versions in the fixture)"),
    ("reserved", 0, "fixture_declaration: no-reservation fixture, reservations is empty"),
]


def build_inputs(wb):
    ws = wb.active
    ws.title = "Inputs"
    title(ws, "Inputs: the golden fixture (tests/fixtures/golden/inputs.json, fixture supplyscope-golden-v1)",
          "SYNTHETIC. Every cell below is a typed input copied from the fixture (blue). Instants are UTC (yyyy-mm-dd hh:mm:ss). Contract v1, docs/contracts.md.")
    ws.cell(4, 1, "parameters").font = BOLD
    header(ws, 5, ["name", "value", "source"])
    for i, (name, val, src) in enumerate(PARAMS):
        r = 6 + i
        ws.cell(r, 1, name)
        v = parse(val)
        c = ws.cell(r, 2, v)
        c.font = BLUE
        if isinstance(v, datetime):
            c.number_format = TS
        ws.cell(r, 3, src)
        POS.setdefault("param", {})[name] = f"Inputs!$B${r}"
    r = 6 + len(PARAMS) + 1

    def block(ent: str, cols: list[str], rows: list[list], note: str | None = None):
        nonlocal r
        ws.cell(r, 1, ent).font = BOLD
        if note:
            ws.cell(r, 4, note)
        header(ws, r + 1, cols)
        first = r + 2
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = ws.cell(first + i, j + 1, parse(v) if v is not None else None)
                c.font = BLUE
                if isinstance(c.value, datetime):
                    c.number_format = TS
        n = max(len(rows), 1)
        POS[ent] = {"first": first, "last": first + n - 1, "col": {c: chr(65 + i) for i, c in enumerate(cols)}}
        r = first + n + 1

    mcols = ["movement_id", "product_id", "location_id", "quantity_delta", "event_at", "movement_type"]
    block("inventory_movements", mcols, [[m[c] for c in mcols] for m in GOLD["inventory_movements"]])
    ocols = ["purchase_order_id", "supplier_id", "ordered_at", "status", "currency"]
    block("purchase_orders", ocols, [[o[c] for c in ocols] for o in GOLD["purchase_orders"]])
    lcols = ["po_line_id", "purchase_order_id", "product_id", "ordered_quantity", "unit_cost"]
    added = {}
    for state, spec in GOLD["states"].items():
        for ln in spec["add_purchase_order_lines"]:
            added[ln["po_line_id"]] = (state, ln)
    lines = [[ln[c] for c in lcols] + [""] for ln in GOLD["purchase_order_lines"]]
    lines += [[ln[c] for c in lcols] + [state] for state, ln in added.values()]
    block("purchase_order_lines", lcols + ["state_added"], lines,
          "state_added is not a fixture field: it names the fixture state (inputs.json states) that adds the line; blank = in every state")
    rcols = ["receipt_id", "po_line_id", "location_id", "received_quantity", "received_at"]
    block("receipts", rcols, [["(none in the golden fixture)"]], "the fixture has no receipt rows: on order is ordered - 0")
    bcols = ["backorder_id", "product_id", "location_id", "quantity", "opened_at", "closed_at"]
    block("backorders", bcols, [[b[c] for c in bcols] for b in GOLD["backorders"]], "closed_at blank = still open")
    widths(ws, {"A": 30, "B": 24, "C": 20, "D": 18, "E": 22, "F": 20, "G": 14})


def rng(ent: str, col: str) -> str:
    p = POS[ent]
    c = p["col"][col]
    return f"Inputs!${c}${p['first']}:${c}${p['last']}"


def cell(ent: str, col: str, idx: int) -> str:
    p = POS[ent]
    return f"Inputs!${p['col'][col]}${p['first'] + idx}"


# ---- Ledger ---------------------------------------------------------------------------------------------------------

LEDGER_FIRST = 5


def build_ledger(wb) -> tuple[int, int]:
    ws = wb.create_sheet("Ledger")
    n = len(GOLD["inventory_movements"])
    last = LEDGER_FIRST + n - 1
    title(ws, "Ledger: running stock with the day order and the as-of filter (contract D-01)",
          "Live formulas over Inputs. Sort key = date x 10 + type rank (receipt 0, sale 1, other 2) + time-of-day / 10, so inside a date receipts come first, then sales, then the rest. "
          "A movement counts only when event_at <= as_of (Inputs).")
    header(ws, 4, ["movement_id", "event_at", "movement_type", "quantity_delta", "date", "type_rank", "sort_key", "in_scope (<= as_of)",
                   "day_order_rank", "running stock, day order", "running stock, timestamp order"])
    A = POS["param"]["as_of"]
    for i in range(n):
        r = LEDGER_FIRST + i
        ws.cell(r, 1, f"={cell('inventory_movements', 'movement_id', i)}")
        ws.cell(r, 2, f"={cell('inventory_movements', 'event_at', i)}").number_format = TS
        ws.cell(r, 3, f"={cell('inventory_movements', 'movement_type', i)}")
        ws.cell(r, 4, f"={cell('inventory_movements', 'quantity_delta', i)}")
        ws.cell(r, 5, f"=INT(B{r})").number_format = "yyyy-mm-dd"
        ws.cell(r, 6, f'=IF(C{r}="receipt",0,IF(C{r}="sale",1,2))')
        ws.cell(r, 7, f"=E{r}*10+F{r}+MOD(B{r},1)/10").number_format = "0.000000"
        ws.cell(r, 8, f"=IF(B{r}<={A},1,0)")
        ws.cell(r, 9, f'=IF(H{r}=1,SUMPRODUCT(($H${LEDGER_FIRST}:$H${last}=1)*($G${LEDGER_FIRST}:$G${last}<=G{r})),"")')
        ws.cell(r, 10, f'=IF(H{r}=1,SUMPRODUCT(($H${LEDGER_FIRST}:$H${last}=1)*($G${LEDGER_FIRST}:$G${last}<=G{r})*$D${LEDGER_FIRST}:$D${last}),"")')
        ws.cell(r, 11, f'=IF(H{r}=1,SUMPRODUCT(($H${LEDGER_FIRST}:$H${last}=1)*($B${LEDGER_FIRST}:$B${last}<=B{r})*$D${LEDGER_FIRST}:$D${last}),"")')
    t = last + 2
    labels = [
        ("on hand, raw (sum of in-scope deltas)", f"=SUMPRODUCT($H${LEDGER_FIRST}:$H${last}*$D${LEDGER_FIRST}:$D${last})"),
        ("on hand (a negative reconstruction counts as 0)", f"=MAX(0,B{t})"),
        ("lowest running stock, day order (incl. the opening 0)", f"=MIN(0,MIN($J${LEDGER_FIRST}:$J${last}))"),
        ("lowest running stock, timestamp order", f"=MIN(0,MIN($K${LEDGER_FIRST}:$K${last}))"),
        ("flag", f'=IF(OR(B{t}<0,B{t + 2}<0),"NEGATIVE_STOCK","ok")'),
        ("movements counted / excluded by the as-of filter", f'=SUM($H${LEDGER_FIRST}:$H${last})&" / "&({n}-SUM($H${LEDGER_FIRST}:$H${last}))'),
    ]
    for i, (lab, f) in enumerate(labels):
        ws.cell(t + i, 1, lab).font = BOLD
        ws.cell(t + i, 2, f)
    widths(ws, {"A": 46, "B": 22, "C": 14, "D": 14, "E": 12, "F": 10, "G": 16, "H": 16, "I": 14, "J": 22, "K": 24})
    return t, last


# ---- Position -------------------------------------------------------------------------------------------------------

def build_position(wb, ledger_t: int):
    ws = wb.create_sheet("Position")
    title(ws, "Position: on hand + on order - backorders (contract D-01, P03-TEST-10)",
          "Live formulas. Reserved is NOT subtracted from the position (declared no-reservation fixture: reserved = 0). Column C is the golden state, column D "
          "the position_80 state (adds the fixture's second line); column B follows the switch Figures!B3.")
    A = POS["param"]["as_of"]
    nl = len(GOLD["purchase_order_lines"]) + sum(len(s["add_purchase_order_lines"]) for s in GOLD["states"].values())
    ws.cell(4, 1, "purchase order lines (open quantity = ordered - received by as_of, confirmed orders placed by as_of)").font = BOLD
    header(ws, 5, ["po_line_id", "ordered_quantity", "received by as_of", "open quantity", "order confirmed and placed (1/0)",
                   "in state golden (1/0)", "in state position_80 (1/0)"])
    first = 6
    for i in range(nl):
        r = first + i
        ws.cell(r, 1, f"={cell('purchase_order_lines', 'po_line_id', i)}")
        ws.cell(r, 2, f"={cell('purchase_order_lines', 'ordered_quantity', i)}")
        ws.cell(r, 3, f"=SUMPRODUCT(({rng('receipts', 'po_line_id')}=A{r})*({rng('receipts', 'received_at')}<={A})*{rng('receipts', 'received_quantity')})")
        ws.cell(r, 4, f"=MAX(B{r}-C{r},0)")
        ws.cell(r, 5, f"=IF(AND(INDEX({rng('purchase_orders', 'status')},MATCH({cell('purchase_order_lines', 'purchase_order_id', i)},"
                      f"{rng('purchase_orders', 'purchase_order_id')},0))=\"confirmed\","
                      f"INDEX({rng('purchase_orders', 'ordered_at')},MATCH({cell('purchase_order_lines', 'purchase_order_id', i)},"
                      f"{rng('purchase_orders', 'purchase_order_id')},0))<={A}),1,0)")
        sa = cell("purchase_order_lines", "state_added", i)
        ws.cell(r, 6, f'=IF({sa}="",1,0)')
        ws.cell(r, 7, f'=IF(OR({sa}="",{sa}="position_80"),1,0)')
    last = first + nl - 1
    t = last + 3
    header(ws, t - 1, ["quantity", "selected (Figures!B3)", "state golden", "state position_80", "rule"])
    bo_open = (f"SUMPRODUCT(({rng('backorders', 'opened_at')}<={A})*((({rng('backorders', 'closed_at')}=\"\")+({rng('backorders', 'closed_at')}>{A}))>0)"
               f"*{rng('backorders', 'quantity')})")
    rows = [
        ("on hand", f"=Ledger!$B${ledger_t + 1}", f"=Ledger!$B${ledger_t + 1}", "Ledger: sum of in-scope posted deltas, floored at 0 (D-01)"),
        ("on order", f"=SUMPRODUCT($D${first}:$D${last}*$E${first}:$E${last}*$F${first}:$F${last})",
         f"=SUMPRODUCT($D${first}:$D${last}*$E${first}:$E${last}*$G${first}:$G${last})", "open quantity of confirmed lines (D-01, D-04)"),
        ("backorders open at as_of", f"={bo_open}", f"={bo_open}", "opened by as_of and not closed by as_of"),
        ("reserved", f"={POS['param']['reserved']}", f"={POS['param']['reserved']}", "declared 0 for the golden fixture"),
        ("available = on hand - reserved", None, None, "not the position"),
        ("inventory position", None, None, "on hand + on order - backorders; reserved not subtracted"),
    ]
    for i, (lab, c_f, d_f, rule) in enumerate(rows):
        r = t + i
        ws.cell(r, 1, lab).font = BOLD
        if lab.startswith("available"):
            c_f, d_f = f"=C{t}-C{t + 3}", f"=D{t}-D{t + 3}"
        if lab == "inventory position":
            c_f, d_f = f"=C{t}+C{t + 1}-C{t + 2}", f"=D{t}+D{t + 1}-D{t + 2}"
        ws.cell(r, 3, c_f)
        ws.cell(r, 4, d_f)
        ws.cell(r, 2, f"=IF(Figures!$B$3=1,D{r},C{r})")
        ws.cell(r, 5, rule)
    widths(ws, {"A": 34, "B": 22, "C": 18, "D": 18, "E": 62, "F": 18, "G": 20})
    return t


# ---- Replenishment --------------------------------------------------------------------------------------------------

def build_replenishment(wb, pos_t: int) -> int:
    ws = wb.create_sheet("Replenishment")
    title(ws, "Replenishment: ROP, target, MOQ and pack rounding (contract D-02, D-06)",
          "Live formulas. raw = target - position; after MOQ = MAX(raw, MOQ); order = CEILING(after MOQ to the pack size); all three are 0 when the reorder is not triggered (position >= ROP).")
    P = POS["param"]
    header(ws, 4, ["quantity", "selected (Figures!B3)", "state golden", "state position_80", "rule"])
    t = 5
    pos_row = pos_t + 5
    spec = [
        ("inventory position", lambda c: f"=Position!{c}{pos_row}", "from Position"),
        ("reorder point", lambda c: f"={P['reorder_point']}", "Inputs replenishment_params"),
        ("target position", lambda c: f"={P['target_position']}", "Inputs replenishment_params"),
        ("triggered (1 = position < reorder point)", lambda c: f"=IF({c}{t}<{c}{t + 1},1,0)", "D-02, D-06"),
        ("raw order = target - position", lambda c: f"=IF({c}{t + 3}=1,MAX({c}{t + 2}-{c}{t},0),0)", "0 when not triggered"),
        ("minimum order quantity (MOQ)", lambda c: f"={P['minimum_order_quantity']}", "Inputs products"),
        ("after MOQ = MAX(raw, MOQ)", lambda c: f"=IF({c}{t + 3}=1,MAX({c}{t + 4},{c}{t + 5}),0)", "MOQ is applied before pack rounding"),
        ("pack size", lambda c: f"={P['pack_size']}", "Inputs products"),
        ("order quantity = CEILING(after MOQ, pack size)", lambda c: f"=IF({c}{t + 3}=1,CEILING({c}{t + 6},{c}{t + 7}),0)", "whole packs"),
        ("order quantity as packs", lambda c: f"={c}{t + 8}/{c}{t + 7}", "check: a whole number"),
    ]
    for i, (lab, f, rule) in enumerate(spec):
        r = t + i
        ws.cell(r, 1, lab).font = BOLD
        ws.cell(r, 3, f("C"))
        ws.cell(r, 4, f("D"))
        ws.cell(r, 2, f"=IF(Figures!$B$3=1,D{r},C{r})")
        ws.cell(r, 5, rule)
    widths(ws, {"A": 46, "B": 22, "C": 16, "D": 18, "E": 50})
    return t


# ---- Figures --------------------------------------------------------------------------------------------------------

def build_figures(wb, pos_t: int, rep_t: int):
    ws = wb.create_sheet("Figures")
    title(ws, "Figures: the golden position and order (P03-GOLD-01) and the position-80 switch",
          "SYNTHETIC. B3 is the state switch: 0 = golden (on order 30, position 60 -> order 72), 1 = position_80 (on order 50, position 80 -> no order). "
          "Column B follows the switch, C is always golden, D always position_80. E and F are the hand-derived expectations typed from expected.json; G and H compare.")
    ws["A3"] = "state switch (0 = golden, 1 = position_80)"
    ws["A3"].font = BOLD
    ws["B3"] = 0
    ws["B3"].font = BLUE
    ws["A4"] = "inputs read"
    ws["B4"] = (f'=IF(ISNUMBER(SEARCH("supplyscope-golden",Inputs!$A$1)),"golden fixture supplyscope-golden-v1","unexpected inputs")')
    header(ws, 7, ["figure", "selected", "golden", "position_80", "expected golden (expected.json)", "expected position_80 (expected.json)",
                   "golden matches", "position_80 matches"])
    g, p = EXPECTED["golden"]["metrics"], EXPECTED["position_80"]["metrics"]
    P = POS["param"]
    figs = [
        ("On hand", f"Position!{{c}}{pos_t}", "on_hand"),
        ("On order", f"Position!{{c}}{pos_t + 1}", "on_order"),
        ("Backorders open", f"Position!{{c}}{pos_t + 2}", "backorders"),
        ("Inventory position", f"Position!{{c}}{pos_t + 5}", "inventory_position"),
        ("Reorder point", f"Replenishment!{{c}}{rep_t + 1}", "reorder_point"),
        ("Target position", f"Replenishment!{{c}}{rep_t + 2}", "target_position"),
        ("Raw order (target - position)", f"Replenishment!{{c}}{rep_t + 4}", "raw_order"),
        ("Order after MOQ", f"Replenishment!{{c}}{rep_t + 6}", "order_after_moq"),
        ("Order quantity (pack rounded)", f"Replenishment!{{c}}{rep_t + 8}", "order_quantity"),
        ("Working capital (on hand x unit cost)", f"Position!{{c}}{pos_t}*{P['unit_cost']}", "working_capital"),
    ]
    first = 8
    for i, (lab, ref, key) in enumerate(figs):
        r = first + i
        ws.cell(r, 1, lab).font = BOLD
        for col, c in (("B", "B"), ("C", "C"), ("D", "D")):
            ws[f"{col}{r}"] = "=" + ref.format(c=c)
        ev, fv = g[key], p[key]
        if key == "working_capital":
            ev, fv = float(ev), float(fv)
        ws[f"E{r}"], ws[f"F{r}"] = ev, fv
        ws[f"E{r}"].font = ws[f"F{r}"].font = BLUE
        ws[f"G{r}"] = f"=ROUND(C{r}-E{r},6)=0"
        ws[f"H{r}"] = f"=ROUND(D{r}-F{r},6)=0"
    last = first + len(figs) - 1
    r = last + 1
    ws.cell(r, 1, "Reorder triggered (1/0)").font = BOLD
    for col in "BCD":
        ws[f"{col}{r}"] = f"=Replenishment!{col}{rep_t + 3}"
    ws[f"E{r}"], ws[f"F{r}"] = 1 if g["reorder_triggered"] else 0, 1 if p["reorder_triggered"] else 0
    ws[f"E{r}"].font = ws[f"F{r}"].font = BLUE
    ws[f"G{r}"] = f"=C{r}=E{r}"
    ws[f"H{r}"] = f"=D{r}=F{r}"
    r += 2
    ws.cell(r, 1, "every golden and position_80 figure matches").font = BOLD
    ws.cell(r, 2, f"=AND(G{first}:G{last + 1},H{first}:H{last + 1})")
    r += 1
    ws.cell(r, 1, "chain 40 / 30 / 10 / 60 / 60 / 60 / 72 (golden) read off column C").font = BOLD
    ws.cell(r, 2, f'=C{first}&" / "&C{first + 1}&" / "&C{first + 2}&" / "&C{first + 3}&" / "&C{first + 6}&" / "&C{first + 7}&" / "&C{first + 8}')
    widths(ws, {"A": 46, "B": 14, "C": 14, "D": 16, "E": 30, "F": 34, "G": 18, "H": 22})
    return first, last + 1, r - 1


# ---- Assumptions ----------------------------------------------------------------------------------------------------

ASSUMPTIONS = [
    ("D-01", "Ledger, day order and as-of stock: stock = sum of posted deltas up to the cutoff; inside a date receipts, then demand, then review; negative reconstruction flagged NEGATIVE_STOCK and valued at 0; "
             "position = on hand + confirmed on order - backorders, reserved not subtracted (docs/contracts.md)."),
    ("D-02", "Replenishment: order = pack_round_up(max(target - position, MOQ)) when position < reorder point, else 0. Golden: 40 / 30 / 10 -> 60; ROP 70, target 120 -> raw 60; MOQ 50, pack 24 -> 72; position 80 -> 0."),
    ("D-03", "Demand and forecasts: observed sales are censored during stockouts, latent demand is published separately; baselines first; a forecast may use only information known at its cutoff. Not computed in this workbook (queries 05 and 07)."),
    ("D-04", "Receipts, suppliers, unmet demand: a receipt before its order is rejected; a partial receipt leaves the line open; on-time rate uses receipts dated up to the cutoff; a backorder or a lost sale, never both."),
    ("D-05", "Simulation and valuation: working capital = sum on hand x unit cost at the cost in force. The golden fixture has one cost and no cost_versions, so on hand x 12.50 = 500.00."),
    ("D-06", "Fixture-driven clarifications: requested = fulfilled + backorders opened in the window; a stockout day is a date on which a backorder opened; raw, MOQ and pack quantities are 0 when not triggered."),
    ("W-01", "Scope: one SKU, one location, one supplier, the golden fixture. Fill rate, stockout days, supplier reliability and forecast accuracy are done by the SQL and Python implementations (docs/sql-analysis.md), not by these formulas."),
    ("W-02", "Instants are Excel date serials (UTC, no zone). Sort keys use SUMPRODUCT over serials, not string-built criteria, so a time such as 08:00 compares exactly."),
    ("W-03", "State switch: Figures!B3 = 1 selects the position_80 state (adds the fixture's second line pol-002, 20 units). The column 'state_added' on Inputs is not a fixture field, it records which state adds the line."),
    ("W-04", "Recalculation: openpyxl writes formulas only. The delivered file was recalculated by LibreOffice headless and carries its cached values; Excel was NOT run (docs/workbook.md). fullCalcOnLoad is set so Excel recalculates on open."),
    ("W-05", "Snapshot summary is a paste of the stored results of data/sql/analysis/run_all.py and goes stale if the data or the queries change; its dataset hash is in the header. The synthetic snapshot is not evidence about a real business."),
]


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions: contract decisions D-01..D-06 and the workbook's own limits")
    header(ws, 3, ["ID", "Statement"])
    for i, (k, v) in enumerate(ASSUMPTIONS):
        ws.cell(4 + i, 1, k).font = BOLD
        c = ws.cell(4 + i, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 8, "B": 150})


# ---- Snapshot summary -----------------------------------------------------------------------------------------------

def put_table(ws, r, name, rows, cols=None):
    ws.cell(r, 1, name).font = BOLD
    for c in range(1, 19):
        ws.cell(r, c).fill = SEC
    cols = cols or list(rows[0])
    header(ws, r + 1, cols)
    for i, row in enumerate(rows):
        for j, c in enumerate(cols):
            v = row.get(c)
            if isinstance(v, bool):
                v = str(v).lower()       # LibreOffice would rewrite a real boolean as =TRUE()
            elif isinstance(v, str) and NUM.match(v):
                v = float(v)
            ws.cell(r + 2 + i, j + 1, v)
    return r + len(rows) + 4


def build_snapshot(wb):
    ws = wb.create_sheet("Snapshot summary")
    ws["A1"] = "PASTED VALUES (not formulas) from data/sql/analysis/run_all.py; each section is titled with its query file"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (f"SYNTHETIC. Dataset {MANIFEST['dataset_version']}, generator dataset_hash {MANIFEST['dataset_hash']}, contract {MANIFEST['contract_version']}; "
                "database seeded through the domain's validate_batch (4 rows rejected and kept in rejected_records).")
    r = 4
    for key in ("01_stock_position_month_end_by_location", "02_draft_orders_arithmetic", "03_fill_rate_stockout_by_location_abc",
                "04_supplier_on_time_by_month", "05_censored_vs_latent_class_a", "06_working_capital_by_location",
                "07_forecast_wape_by_sku", "08_rejections_and_defects_by_reason", "09_golden_position_and_order"):
        r = put_table(ws, r, f"{key}.sql", RESULTS[key])
    r = put_table(ws, r, "golden position and order: analysis SQL 09 vs data/sql/reference_metrics.sql vs Python domain vs expected.json", RESULTS["golden_three_sources"])
    r = put_table(ws, r, "options.py: draft plan under the parameters of record and under forecast-derived parameters at 0.90 and 0.95 (Python domain)", RESULTS["10_options_service_target"])
    r = put_table(ws, r, "options.py: lost sale against backlog on common random numbers (simulated direction)", RESULTS["10_options_unmet_policy"])
    r = put_table(ws, r, "bi-package.md expected measure values (computed outside Power BI)", RESULTS["bi_expected_measures"])
    r = put_table(ws, r, "reference_metrics.sql over the whole snapshot (cross-check of queries 01, 02, 03 and 06)", RESULTS["reference_metrics_snapshot"])
    ok = sum(1 for x in RESULTS["reconciliation"] if x["equal"] is True)
    r = put_table(ws, r, "reconciliation (run_all.py): checks equal / checks run", [{"checks_equal": ok, "checks_run": len(RESULTS["reconciliation"])}])
    widths(ws, {"A": 26, "B": 22, "C": 22, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16, "J": 16, "K": 16})


# ---- recalculation --------------------------------------------------------------------------------------------------

def recalc(path: Path) -> None:
    """Open with LibreOffice headless and save a recalculated copy over ``path`` (cached values included)."""
    if shutil.which("soffice") is None:
        raise SystemExit("soffice not installed; the file has formulas but no cached values")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "in" / path.name
        src.parent.mkdir()
        shutil.copy(path, src)
        out = Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=240)
        shutil.copy(out / path.name, path)
    set_full_calc_on_load(path)


def set_full_calc_on_load(path: Path) -> None:
    """LibreOffice drops fullCalcOnLoad; put it back (workbook.xml only) so Excel recalculates when it opens the file."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode("utf-8")
                assert "<calcPr " in text and "fullCalcOnLoad" not in text
                data = text.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ', 1).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def main(argv: list[str]) -> int:
    wb = Workbook()
    build_inputs(wb)
    ledger_t, _ = build_ledger(wb)
    pos_t = build_position(wb, ledger_t)
    rep_t = build_replenishment(wb, pos_t)
    build_figures(wb, pos_t, rep_t)
    build_assumptions(wb)
    build_snapshot(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(wb.sheetnames)} sheets: {', '.join(wb.sheetnames)})")
    if "--recalc" in argv:
        recalc(OUT)
        print("recalculated with LibreOffice headless; cached values written")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
